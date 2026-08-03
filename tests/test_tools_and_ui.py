import json
import math
import os
import tempfile
import threading
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import Mock, patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from PyQt5.QtCore import (
    QCoreApplication,
    QDate,
    QEvent,
    QObject,
    QPoint,
    QSize,
    Qt,
    pyqtSignal,
)
from PyQt5.QtGui import QMouseEvent, QPalette
from PyQt5.QtTest import QTest
from PyQt5.QtWidgets import (
    QApplication,
    QComboBox,
    QFrame,
    QHeaderView,
    QLabel,
    QMessageBox,
    QPushButton,
    QTableView,
    QToolButton,
)

from integrated_client.app_controller import ApplicationController
from integrated_client.config import (
    APP_NAME,
    DEFAULT_ADMIN_PASSWORD,
    DEFAULT_ADMIN_USERNAME,
)
from integrated_client.database import (
    DEFAULT_STATION_PASSWORD,
    DEFAULT_STATION_USERS,
    WORKFLOW_EMPTY_METRIC,
    WORKFLOW_HAS_PHONE_METRIC,
    WORKFLOW_INDIVIDUAL_METRIC,
    WORKFLOW_NO_OPERATION_METRIC,
    WORKFLOW_NO_PHONE_METRIC,
    WORKFLOW_NO_TRANSPORT_METRIC,
    WORKFLOW_TOTAL_METRIC,
    Database,
)
from integrated_client.online.api import ApiResponseError
from integrated_client.online.coordinator import SyncCoordinator
from integrated_client.online.sync import SyncStatus
from integrated_client.online.update import UpdateInfo
from integrated_client.preferences import ClientPreferences
from integrated_client.tencent_docs import TencentDocsImportResult
from integrated_client.timing import WorkflowTimingService
from integrated_client.tools.aiqicha_tool import (
    ADDR_COL_NAME,
    AIQICHA_RESULT_WAIT_SECONDS,
    LEGAL_COL_NAME,
    PHONE_COL_NAME,
    TARGET_COLUMNS,
    QueryWorker,
    create_browser,
    has_meaningful_value,
)
from integrated_client.tools.aiqicha_tool import MainWindow as AiqichaToolWidget
from integrated_client.tools.transport_tool import (
    CONFIG,
    BusinessBackfillWorker,
    TargetedWorkbookWriter,
    Worker,
)
from integrated_client.ui.auth_dialogs import LoginDialog, PasswordDialog
from integrated_client.ui.dashboard_page import (
    DashboardMetricIcon,
    DashboardPage,
    StationShareChart,
)
from integrated_client.ui.frameless import (
    HTBOTTOMRIGHT,
    HTCAPTION,
    HTCLIENT,
    HTTOPLEFT,
    MINMAXINFO,
    WVR_REDRAW,
    FramelessMessageBox,
)
from integrated_client.ui.main_window import MainWindow
from integrated_client.ui.online_account_page import (
    OnlineAccountPage,
    _AccountSettingsDialog,
)
from integrated_client.ui.statistics_page import (
    AnimatedDonutChart,
    StationDistributionChart,
    StatisticsPage,
    ViolationReasonChart,
    WorkflowDistributionChart,
)
from integrated_client.ui.tencent_docs_dialog import TencentDocsProgressDialog
from integrated_client.ui.theme import APP_STYLESHEET, _control_asset_path
from integrated_client.ui.update_dialog import UpdatePromptDialog
from integrated_client.ui.workflow_page import DataFrameTableModel, WorkflowPage


class ToolAndUiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    @classmethod
    def tearDownClass(cls):
        # PyQt 在解释器关闭阶段回收仍存活的主窗口和全局计时器时，
        # Windows 下可能返回非零退出码；显式释放可让测试结果稳定传给 CI。
        for widget in list(cls.app.topLevelWidgets()):
            workflow = getattr(widget, "workflow_page", None)
            if workflow is not None:
                workflow.shutdown()
            if hasattr(widget, "_prepared_to_close"):
                widget._prepared_to_close = True
            widget.close()
            widget.deleteLater()
        cls.app.processEvents()
        cursor_filter = getattr(cls.app, "_intdemo_disabled_cursor_filter", None)
        if cursor_filter is not None:
            cursor_filter.stop()
            cls.app.removeEventFilter(cursor_filter)
        cls.app.processEvents()

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db = Database(Path(self.temp_dir.name) / "test.db")
        self.db.ensure_default_admin()
        self.admin = self.db.authenticate(DEFAULT_ADMIN_USERNAME, DEFAULT_ADMIN_PASSWORD)

    def tearDown(self):
        for widget in list(self.app.topLevelWidgets()):
            workflow = getattr(widget, "workflow_page", None)
            if workflow is not None:
                workflow.shutdown()
            else:
                shutdown = getattr(widget, "shutdown", None)
                if callable(shutdown):
                    shutdown()
            if hasattr(widget, "_prepared_to_close"):
                widget._prepared_to_close = True
            widget.close()
            widget.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
        self.app.processEvents()
        self.temp_dir.cleanup()

    def test_value_detection_and_target_header_creation(self):
        self.assertTrue(has_meaningful_value("张三"))
        self.assertFalse(has_meaningful_value("nan"))
        self.assertFalse(has_meaningful_value("未查询到公司"))

        workbook = Workbook()
        sheet = workbook.active
        sheet.cell(1, 1, "车辆所有人/企业")
        header = AiqichaToolWidget._ensure_target_headers(sheet)
        for name in TARGET_COLUMNS:
            self.assertIn(name, header)
            self.assertEqual(sheet.cell(1, header[name]).value, name)

    def test_targeted_writer_preserves_rows_below_a_blank_row_and_workbook_layout(self):
        file_path = Path(self.temp_dir.name) / "targeted-write.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "业务数据"
        sheet.append(["车辆标识", "已协助补缴", "查询状态", "金额"])
        sheet.append(["粤A10001(黄色)", "", "已完成", 10])
        sheet.append([None, None, None, None])
        sheet.append(
            [
                "粤A10002(黄色)",
                "保留下方数据",
                "待处理",
                "=1+2",
                "无表头列也要保留",
            ]
        )
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 28
        sheet["A4"].fill = PatternFill("solid", fgColor="FFF2CC")
        notes = workbook.create_sheet("说明")
        notes["A1"] = "这个工作表必须保留"
        notes.merge_cells("A1:B1")
        workbook.save(file_path)
        workbook.close()

        writer = TargetedWorkbookWriter(
            file_path,
            ("运输证号_纯数字", "查询状态"),
        )
        writer.write_row(
            1,
            {
                "运输证号_纯数字": "",
                "查询状态": "中间空行停止",
            },
        )
        writer.save()
        writer.close()

        saved = load_workbook(file_path, data_only=False)
        sheet = saved["业务数据"]
        headers = {
            cell.value: cell.column
            for cell in sheet[1]
            if cell.value is not None
        }
        self.assertEqual(sheet.cell(3, headers["查询状态"]).value, "中间空行停止")
        self.assertEqual(sheet["A4"].value, "粤A10002(黄色)")
        self.assertEqual(sheet["B4"].value, "保留下方数据")
        self.assertEqual(sheet["C4"].value, "待处理")
        self.assertEqual(sheet["D4"].value, "=1+2")
        self.assertEqual(sheet["E4"].value, "无表头列也要保留")
        self.assertEqual(headers["运输证号_纯数字"], 6)
        self.assertEqual(sheet["A4"].fill.fgColor.rgb, "00FFF2CC")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.column_dimensions["A"].width, 28)
        self.assertEqual(saved["说明"]["A1"].value, "这个工作表必须保留")
        self.assertIn("A1:B1", {str(cell_range) for cell_range in saved["说明"].merged_cells.ranges})
        saved.close()

    def test_targeted_writer_reuses_formatted_blank_columns_after_anchor(self):
        file_path = Path(self.temp_dir.name) / "formatted-empty-columns.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet["H1"] = "车辆所有人/企业"
        sheet["H2"] = "示例公司"
        sheet["O1"] = "站场"
        sheet["P1"] = "已协助补缴"
        for row in range(1, 4):
            for column in range(17, 20):
                sheet.cell(row, column).fill = PatternFill("solid", fgColor="FFF2CC")
        sheet["T1"] = "运输证号_纯数字"
        sheet["T2"] = "131102061"
        sheet["U1"] = "查询状态"
        sheet["U2"] = "查询成功"
        sheet["V1"] = "回填状态"
        sheet["V2"] = "回填成功"
        workbook.save(file_path)
        workbook.close()

        transport_writer = TargetedWorkbookWriter(
            file_path,
            ("运输证号_纯数字", "查询状态"),
        )
        transport_writer.save()
        transport_writer.close()
        backfill_writer = TargetedWorkbookWriter(
            file_path,
            ("车辆所有人/企业", "回填状态"),
        )
        backfill_writer.save()
        backfill_writer.close()

        saved = load_workbook(file_path)
        sheet = saved.active
        self.assertEqual(sheet["Q1"].value, "运输证号_纯数字")
        self.assertEqual(sheet["Q2"].value, "131102061")
        self.assertEqual(sheet["R1"].value, "查询状态")
        self.assertEqual(sheet["R2"].value, "查询成功")
        self.assertEqual(sheet["S1"].value, "回填状态")
        self.assertEqual(sheet["S2"].value, "回填成功")
        self.assertEqual(sheet["H1"].value, "车辆所有人/企业")
        self.assertEqual(sheet["H2"].value, "示例公司")
        for column in ("T", "U", "V"):
            self.assertIsNone(sheet[f"{column}1"].value)
            self.assertIsNone(sheet[f"{column}2"].value)
        saved.close()
        self.assertEqual(
            list(file_path.parent.glob(f".{file_path.stem}.*.tmp{file_path.suffix}")),
            [],
        )

    def test_transport_worker_skips_rows_with_existing_company_data(self):
        file_path = Path(self.temp_dir.name) / "existing-company.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["车辆标识", "已协助补缴", "车辆所有人/企业"])
        sheet.append(["冀T2892_黄色", "", "无运输证号"])
        workbook.save(file_path)
        workbook.close()

        class FakePage:
            def __init__(self):
                self.goto_calls = 0

            def goto(self, *_args, **_kwargs):
                self.goto_calls += 1
                raise AssertionError("已有企业信息时不应访问运输证查询网站")

        class FakeContext:
            def __init__(self, page):
                self.page = page

            def new_page(self):
                return self.page

        class FakeBrowser:
            def __init__(self, page):
                self.page = page

            def new_context(self, **_kwargs):
                return FakeContext(self.page)

            def close(self):
                return None

        class FakeChromium:
            def __init__(self, page):
                self.page = page

            def launch(self, **_kwargs):
                return FakeBrowser(self.page)

        class FakeRuntime:
            def __init__(self, page):
                self.chromium = FakeChromium(page)

            def stop(self):
                return None

        class FakePlaywrightManager:
            def __init__(self, runtime):
                self.runtime = runtime

            def start(self):
                return self.runtime

        fake_page = FakePage()
        runtime = FakeRuntime(fake_page)
        worker = Worker(
            str(file_path),
            True,
            False,
            2,
            True,
            2,
            False,
        )
        finished_results = []
        worker.finished.connect(finished_results.append)
        with patch(
            "integrated_client.tools.transport_tool.sync_playwright",
            return_value=FakePlaywrightManager(runtime),
        ), patch(
            "integrated_client.tools.transport_tool.get_builtin_chromium_path",
            return_value=r"C:\browser\chrome.exe",
        ):
            worker.run()

        self.assertEqual(fake_page.goto_calls, 0)
        self.assertEqual(finished_results, ["完成"])
        saved = load_workbook(file_path)
        sheet = saved.active
        headers = {
            cell.value: cell.column
            for cell in sheet[1]
            if cell.value is not None
        }
        self.assertEqual(sheet.cell(2, headers["车辆所有人/企业"]).value, "无运输证号")
        self.assertEqual(sheet.cell(2, headers["查询状态"]).value, "已有企业信息（跳过）")
        saved.close()

    def test_transport_worker_restarts_closed_browser_without_limit(self):
        file_path = Path(self.temp_dir.name) / "transport-browser-restart.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["车辆标识", "已协助补缴", "车辆所有人/企业", "查询状态"])
        sheet.append(
            [
                "冀TJ2892_黄色",
                "",
                "无运输证号",
                "Page.goto: Target page, context or browser has bee",
            ]
        )
        workbook.save(file_path)
        workbook.close()

        class Response:
            status = 200

        class Locator:
            def __init__(self, selector):
                self.selector = selector

            def click(self, *_args, **_kwargs):
                return None

            def select_option(self, *_args, **_kwargs):
                return None

            def fill(self, *_args, **_kwargs):
                return None

            def count(self):
                return 0

        class Page:
            def __init__(self, closed=False):
                self.closed = closed
                self.goto_calls = 0

            def goto(self, *_args, **_kwargs):
                self.goto_calls += 1
                if self.closed:
                    raise RuntimeError(
                        "Page.goto: Target page, context or browser has been closed"
                    )
                return Response()

            def locator(self, selector):
                return Locator(selector)

            def wait_for_timeout(self, *_args, **_kwargs):
                return None

            def evaluate(self, *_args, **_kwargs):
                return None

            def content(self):
                return "查询不到信息"

            def reload(self):
                return None

        class Context:
            def __init__(self, page):
                self.page = page

            def new_page(self):
                return self.page

            def close(self):
                return None

        class Browser:
            def __init__(self, page):
                self.page = page

            def new_context(self, **_kwargs):
                return Context(self.page)

            def close(self):
                return None

        class Chromium:
            def __init__(self, page):
                self.page = page

            def launch(self, **_kwargs):
                return Browser(self.page)

        class Runtime:
            def __init__(self, page):
                self.chromium = Chromium(page)

            def stop(self):
                return None

        class Manager:
            def __init__(self, runtime):
                self.runtime = runtime

            def start(self):
                return self.runtime

        closed_pages = [Page(closed=True) for _ in range(4)]
        recovered_page = Page()
        managers = iter(
            [
                *(Manager(Runtime(page)) for page in closed_pages),
                Manager(Runtime(recovered_page)),
            ]
        )
        worker = Worker(str(file_path), True, False, 2, True, 2, False)
        logs = []
        retries = []
        results = []
        worker.log.connect(logs.append)
        worker.retry_signal.connect(lambda retry_type, reason: retries.append((retry_type, reason)))
        worker.finished.connect(results.append)
        with patch(
            "integrated_client.tools.transport_tool.sync_playwright",
            side_effect=lambda: next(managers),
        ), patch(
            "integrated_client.tools.transport_tool.get_builtin_chromium_path",
            return_value=r"C:\browser\chrome.exe",
        ), patch(
            "integrated_client.tools.transport_tool.ocr_code",
            return_value="1234",
        ):
            worker.run()

        self.assertTrue(all(page.goto_calls == 1 for page in closed_pages))
        self.assertGreaterEqual(recovered_page.goto_calls, 1)
        self.assertEqual(results, ["完成"])
        self.assertGreaterEqual(
            sum(item[0] == "transport_browser" for item in retries),
            4,
        )
        self.assertTrue(any("重试当前记录" in line for line in logs))
        self.assertTrue(any("无次数上限" in line for line in logs))
        saved = load_workbook(file_path, data_only=True)
        headers = {
            cell.value: cell.column
            for cell in saved.active[1]
            if cell.value is not None
        }
        self.assertEqual(
            saved.active.cell(2, headers["查询状态"]).value,
            "查询无结果",
        )
        saved.close()

    def test_business_worker_restarts_after_repeated_http_404(self):
        file_path = Path(self.temp_dir.name) / "business-browser-restart.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            [
                "车辆标识",
                "已协助补缴",
                "运输证号_纯数字",
                "车辆所有人/企业",
            ]
        )
        sheet.append(["冀TJ2892_黄色", "", "130000001", "无运输证号"])
        workbook.save(file_path)
        workbook.close()

        class Response:
            def __init__(self, status):
                self.status = status

        class Page:
            def __init__(self, status):
                self.status = status
                self.goto_calls = 0

            def goto(self, *_args, **_kwargs):
                self.goto_calls += 1
                return Response(self.status)

            def wait_for_selector(self, *_args, **_kwargs):
                return None

            def click(self, *_args, **_kwargs):
                return None

            def fill(self, *_args, **_kwargs):
                return None

        class Context:
            def __init__(self, page):
                self.page = page

            def new_page(self):
                return self.page

        class Browser:
            def __init__(self, page):
                self.page = page

            def new_context(self, **_kwargs):
                return Context(self.page)

            def close(self):
                return None

        class Chromium:
            def __init__(self, page):
                self.page = page

            def launch(self, **_kwargs):
                return Browser(self.page)

        class Runtime:
            def __init__(self, page):
                self.chromium = Chromium(page)

            def stop(self):
                return None

        class Manager:
            def __init__(self, runtime):
                self.runtime = runtime

            def start(self):
                return self.runtime

        not_found_pages = [Page(404) for _ in range(4)]
        recovered_page = Page(200)
        managers = iter(
            [
                *(Manager(Runtime(page)) for page in not_found_pages),
                Manager(Runtime(recovered_page)),
            ]
        )
        worker = BusinessBackfillWorker(str(file_path), True, True, 2, False)
        logs = []
        retries = []
        results = []
        worker.log.connect(logs.append)
        worker.retry_signal.connect(lambda retry_type, reason: retries.append((retry_type, reason)))
        worker.finished.connect(results.append)
        with patch(
            "integrated_client.tools.transport_tool.sync_playwright",
            side_effect=lambda: next(managers),
        ), patch(
            "integrated_client.tools.transport_tool.get_builtin_chromium_path",
            return_value=r"C:\browser\chrome.exe",
        ), patch.object(
            BusinessBackfillWorker,
            "_wait_for_business_response",
            return_value="captcha",
        ), patch.object(BusinessBackfillWorker, "solve_captcha", return_value=False):
            worker.run()

        self.assertTrue(all(page.goto_calls == 1 for page in not_found_pages))
        self.assertEqual(recovered_page.goto_calls, 1)
        self.assertEqual(results, ["结束"])
        self.assertGreaterEqual(
            sum(item[0] == "business_browser" for item in retries),
            4,
        )
        self.assertTrue(any("重新启动并重试当前记录" in line for line in logs))
        self.assertTrue(any("无次数上限" in line for line in logs))
        saved = load_workbook(file_path, data_only=True)
        headers = {
            cell.value: cell.column
            for cell in saved.active[1]
            if cell.value is not None
        }
        self.assertEqual(
            saved.active.cell(2, headers["回填状态"]).value,
            "验证码识别失败",
        )
        saved.close()

    def test_business_waits_for_stable_captcha_image_and_result(self):
        worker = BusinessBackfillWorker("unused.xlsx", True, True, 2, False)

        class Page:
            def __init__(self):
                self.waits = []

            def wait_for_timeout(self, milliseconds):
                self.waits.append(milliseconds)

        page = Page()
        worker.page = page
        logs = []
        worker.log.connect(logs.append)

        with patch.object(
            worker,
            "_business_result_is_ready",
            return_value=False,
        ), patch.object(
            worker,
            "_captcha_image_is_ready",
            side_effect=(False, True, False, True, True),
        ):
            response_state = worker._wait_for_business_response()

        self.assertEqual(worker.web_timeout, 120_000)
        self.assertEqual(CONFIG["BROWSER_WAIT_TIMEOUT_MS"], 120_000)
        self.assertEqual(response_state, "captcha")
        self.assertEqual(page.waits, [250, 800, 250, 800])
        self.assertTrue(any("等待营运查询响应" in line for line in logs))
        self.assertTrue(any("验证码图片加载完成" in line for line in logs))

        page.waits.clear()
        with patch.object(
            worker,
            "_business_result_is_ready",
            side_effect=(False, False, True),
        ):
            self.assertTrue(worker._wait_for_business_result())
        self.assertEqual(page.waits, [250, 250])

    def test_aiqicha_worker_restarts_browser_without_limit_and_retries_same_row(self):
        dataframe = pd.DataFrame(
            {
                "车辆所有人/企业": ["测试运输有限公司"],
                "负责人/法人代表": [""],
                "地址": [""],
                "电话": [""],
            }
        )

        class Page:
            def __init__(self, fail_search=False):
                self.fail_search = fail_search
                self.home_calls = 0
                self.search_calls = 0

            def get(self, url):
                if url == "https://aiqicha.baidu.com":
                    self.home_calls += 1
                    return None
                self.search_calls += 1
                if self.fail_search:
                    raise RuntimeError("connection disconnected: browser closed")
                return None

            def quit(self):
                return None

        failed_pages = [Page(fail_search=True) for _ in range(4)]
        recovered_page = Page()
        pages = iter([*failed_pages, recovered_page])
        worker = QueryWorker(dataframe, "车辆所有人/企业")
        worker.login_required_signal.connect(worker.confirm_login)
        logs = []
        retries = []
        rows = []
        results = []
        worker.log_signal.connect(logs.append)
        worker.retry_signal.connect(
            lambda retry_type, reason: retries.append((retry_type, reason))
        )
        worker.row_done_signal.connect(
            lambda row_index, values: rows.append((row_index, values))
        )
        worker.finished_signal.connect(results.append)

        with patch(
            "integrated_client.tools.aiqicha_tool.create_browser",
            side_effect=lambda: next(pages),
        ), patch.object(
            worker,
            "_interruptible_sleep",
            return_value=True,
        ), patch.object(
            worker,
            "_wait_for_results",
            return_value=True,
        ) as wait_for_results, patch.object(
            worker,
            "_check_captcha",
            return_value=False,
        ), patch(
            "integrated_client.tools.aiqicha_tool.check_no_results",
            return_value=False,
        ), patch(
            "integrated_client.tools.aiqicha_tool.extract_info",
            return_value={
                "法定代表人": "张三",
                "地址": "测试地址",
                "电话": "13800138000",
            },
        ):
            worker.run()

        self.assertTrue(all(page.search_calls == 1 for page in failed_pages))
        self.assertEqual(recovered_page.search_calls, 1)
        self.assertEqual(results, [True])
        self.assertEqual(
            rows,
            [
                (
                    0,
                    {
                        "负责人/法人代表": "张三",
                        "地址": "测试地址",
                        "电话": "13800138000",
                    },
                )
            ],
        )
        self.assertGreaterEqual(
            sum(item[0] == "aiqicha_browser" for item in retries),
            4,
        )
        self.assertTrue(any("无次数上限" in line for line in logs))
        wait_for_results.assert_called_once_with(
            timeout=AIQICHA_RESULT_WAIT_SECONDS
        )
        self.assertEqual(AIQICHA_RESULT_WAIT_SECONDS, 120)

    def test_aiqicha_browser_profile_and_persisted_login_are_reused(self):
        profile_directory = Path(self.temp_dir.name) / "aiqicha-profile"

        class Options:
            def __init__(self):
                self.user_data_path = None
                self.local_port = None

            def set_argument(self, _argument):
                return self

            def set_user_data_path(self, path):
                self.user_data_path = path
                return self

            def set_local_port(self, port):
                self.local_port = port
                return self

            def set_user_agent(self, _user_agent):
                return self

            def set_browser_path(self, _path):
                return self

        class Timeouts:
            def timeouts(self, **_kwargs):
                return None

        options = Options()
        browser_page = type("BrowserPage", (), {"set": Timeouts()})()
        with patch(
            "integrated_client.tools.aiqicha_tool.ChromiumOptions",
            return_value=options,
        ), patch(
            "integrated_client.tools.aiqicha_tool.ChromiumPage",
            return_value=browser_page,
        ), patch(
            "integrated_client.tools.aiqicha_tool.get_builtin_chromium_path",
            return_value=r"C:\browser\chrome.exe",
        ), patch(
            "integrated_client.tools.aiqicha_tool._available_local_port",
            return_value=19321,
        ):
            self.assertIs(
                create_browser(profile_directory),
                browser_page,
            )

        self.assertTrue(profile_directory.is_dir())
        self.assertEqual(
            options.user_data_path,
            str(profile_directory.resolve()),
        )
        self.assertEqual(options.local_port, 19321)

        dataframe = pd.DataFrame({"车辆所有人/企业": []})

        class Page:
            def __init__(self):
                self.home_calls = 0

            def get(self, url):
                self.assertEqualUrl(url)
                self.home_calls += 1

            @staticmethod
            def assertEqualUrl(url):
                if url != "https://aiqicha.baidu.com":
                    raise AssertionError(url)

            @staticmethod
            def cookies(all_domains=False):
                if not all_domains:
                    raise AssertionError("expected all-domain cookies")
                return [
                    {
                        "name": "BDUSS",
                        "value": "persisted-session",
                        "domain": ".baidu.com",
                    }
                ]

            def quit(self):
                return None

        page = Page()
        worker = QueryWorker(
            dataframe,
            "车辆所有人/企业",
            browser_profile_directory=profile_directory,
        )
        login_requests = []
        logs = []
        worker.login_required_signal.connect(lambda: login_requests.append(True))
        worker.log_signal.connect(logs.append)
        with patch(
            "integrated_client.tools.aiqicha_tool.create_browser",
            return_value=page,
        ) as create_browser_mock, patch.object(
            worker,
            "_interruptible_sleep",
            return_value=True,
        ):
            self.assertTrue(worker._start_browser_session())

        create_browser_mock.assert_called_once_with(profile_directory.resolve())
        self.assertEqual(page.home_calls, 1)
        self.assertEqual(login_requests, [])
        self.assertTrue(any("无需重复登录" in message for message in logs))

    def test_aiqicha_persisted_login_requires_baidu_auth_cookie(self):
        class Page:
            def __init__(self, cookies):
                self._cookies = cookies

            def cookies(self, all_domains=False):
                self.assert_all_domains = all_domains
                return self._cookies

        self.assertFalse(
            QueryWorker._has_persisted_login(
                Page(
                    [
                        {
                            "name": "BAIDUID",
                            "value": "anonymous",
                            "domain": ".baidu.com",
                        }
                    ]
                )
            )
        )
        self.assertFalse(
            QueryWorker._has_persisted_login(
                Page(
                    [
                        {
                            "name": "BDUSS",
                            "value": "other-site",
                            "domain": ".example.com",
                        }
                    ]
                )
            )
        )

    def test_aiqicha_worker_reopens_browser_without_limit_while_waiting_for_login(self):
        dataframe = pd.DataFrame(
            {
                "车辆所有人/企业": ["测试运输有限公司"],
                "负责人/法人代表": [""],
                "地址": [""],
                "电话": [""],
            }
        )

        class States:
            def __init__(self):
                self.is_alive = True

        class Page:
            def __init__(self):
                self.states = States()
                self.home_calls = 0
                self.quit_calls = 0

            def get(self, url):
                self.assert_home_url(url)
                self.home_calls += 1

            @staticmethod
            def assert_home_url(url):
                if url != "https://aiqicha.baidu.com":
                    raise AssertionError(f"unexpected URL: {url}")

            def quit(self):
                self.quit_calls += 1
                self.states.is_alive = False

        closed_pages = [Page() for _ in range(4)]
        reopened_page = Page()
        worker = QueryWorker(dataframe, "车辆所有人/企业")
        login_requests = []
        logs = []
        retries = []

        def handle_login_request():
            login_requests.append(worker.page)
            if len(login_requests) <= len(closed_pages):
                worker.page.states.is_alive = False
            else:
                worker.confirm_login()

        worker.login_required_signal.connect(handle_login_request)
        worker.log_signal.connect(logs.append)
        worker.retry_signal.connect(
            lambda retry_type, reason: retries.append((retry_type, reason))
        )

        with patch(
            "integrated_client.tools.aiqicha_tool.create_browser",
            side_effect=[*closed_pages, reopened_page],
        ) as create_browser_mock, patch.object(
            worker,
            "_interruptible_sleep",
            return_value=True,
        ):
            self.assertTrue(worker._start_browser_session())

        self.assertEqual(create_browser_mock.call_count, 5)
        self.assertEqual(login_requests, [*closed_pages, reopened_page])
        self.assertTrue(all(page.home_calls == 1 for page in closed_pages))
        self.assertTrue(all(page.quit_calls == 1 for page in closed_pages))
        self.assertEqual(reopened_page.home_calls, 1)
        self.assertIs(worker.page, reopened_page)
        self.assertGreaterEqual(
            sum(item[0] == "aiqicha_browser_start" for item in retries),
            4,
        )
        self.assertTrue(any("自动重新打开" in line for line in logs))
        self.assertTrue(any("正在重新启动爱企查浏览器" in line for line in logs))

    def test_aiqicha_stop_interrupts_home_navigation_after_browser_closed(self):
        dataframe = pd.DataFrame(
            {
                "车辆所有人/企业": ["测试运输有限公司"],
                "负责人/法人代表": [""],
                "地址": [""],
                "电话": [""],
            }
        )

        class BlockingPage:
            def __init__(self):
                self.navigation_started = threading.Event()
                self.navigation_released = threading.Event()
                self.quit_finished = threading.Event()
                self.stop_loading_calls = 0
                self.quit_calls = 0
                self.manually_closed = False

            def get(self, url):
                if url != "https://aiqicha.baidu.com":
                    raise AssertionError(f"unexpected URL: {url}")
                self.navigation_started.set()
                if not self.navigation_released.wait(3):
                    raise AssertionError("home navigation was not interrupted")
                return False

            def simulate_manual_close(self):
                self.manually_closed = True

            def stop_loading(self):
                self.stop_loading_calls += 1
                self.navigation_released.set()

            def quit(self):
                self.quit_calls += 1
                self.navigation_released.set()
                self.quit_finished.set()

        page = BlockingPage()
        worker = QueryWorker(dataframe, "车辆所有人/企业")
        results = []
        retries = []
        worker.finished_signal.connect(results.append)
        worker.retry_signal.connect(
            lambda retry_type, reason: retries.append((retry_type, reason))
        )

        with patch(
            "integrated_client.tools.aiqicha_tool.create_browser",
            return_value=page,
        ) as create_browser_mock:
            worker.start()
            self.assertTrue(page.navigation_started.wait(1))
            page.simulate_manual_close()
            worker.stop()
            self.assertTrue(worker.wait(2_000))
            self.assertTrue(page.quit_finished.wait(1))

        self.app.processEvents()
        self.assertFalse(worker.isRunning())
        self.assertTrue(page.manually_closed)
        self.assertEqual(create_browser_mock.call_count, 1)
        self.assertEqual(page.stop_loading_calls, 1)
        self.assertEqual(page.quit_calls, 1)
        self.assertIsNone(worker.page)
        self.assertEqual(results, [False])
        self.assertEqual(retries, [])

    def test_workflow_force_shutdown_atomically_saves_completed_aiqicha_data(self):
        file_path = Path(self.temp_dir.name) / "force-exit-save.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            ["车辆所有人/企业", "负责人/法人代表", "地址", "电话"]
        )
        sheet.append(["测试运输有限公司", "", "", ""])
        workbook.save(file_path)
        workbook.close()

        class StuckThread:
            def __init__(self):
                self.running = True
                self.stop_called = False
                self.interruption_requested = False
                self.terminated = False
                self.deleted = False

            def isRunning(self):
                return self.running

            def stop(self):
                self.stop_called = True

            def requestInterruption(self):
                self.interruption_requested = True

            def terminate(self):
                self.terminated = True
                self.running = False

            def wait(self, _timeout):
                return not self.running

            def deleteLater(self):
                self.deleted = True

        worker = StuckThread()
        page = WorkflowPage()
        page.file_path = str(file_path)
        page.df = pd.DataFrame(
            {
                "车辆所有人/企业": ["测试运输有限公司"],
                "负责人/法人代表": ["张三"],
                "地址": ["测试地址"],
                "电话": ["13800138000"],
            }
        )
        page.current_step = 3
        page.pipeline_running = True
        page.current_worker = worker

        self.assertTrue(page.force_shutdown(terminate_wait_ms=0))
        self.assertTrue(worker.stop_called)
        self.assertTrue(worker.interruption_requested)
        self.assertTrue(worker.terminated)
        self.assertTrue(worker.deleted)
        self.assertIsNone(page.current_worker)
        self.assertFalse(page.pipeline_running)

        saved = load_workbook(file_path)
        saved_sheet = saved.active
        self.assertEqual(saved_sheet.cell(2, 2).value, "张三")
        self.assertEqual(saved_sheet.cell(2, 3).value, "测试地址")
        self.assertEqual(saved_sheet.cell(2, 4).value, "13800138000")
        saved.close()
        self.assertEqual(
            list(file_path.parent.glob(f".{file_path.stem}.*.tmp{file_path.suffix}")),
            [],
        )
        page.close()

    def test_workflow_force_shutdown_does_not_terminate_when_save_fails(self):
        file_path = Path(self.temp_dir.name) / "force-exit-save-failure.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            ["车辆所有人/企业", "负责人/法人代表", "地址", "电话"]
        )
        sheet.append(["测试运输有限公司", "", "", ""])
        workbook.save(file_path)
        workbook.close()

        class StuckThread:
            def __init__(self):
                self.running = True
                self.stop_called = False
                self.terminated = False

            def isRunning(self):
                return self.running

            def stop(self):
                self.stop_called = True

            def requestInterruption(self):
                return None

            def terminate(self):
                self.terminated = True
                self.running = False

            def wait(self, _timeout):
                return not self.running

            def deleteLater(self):
                return None

        worker = StuckThread()
        page = WorkflowPage()
        page.file_path = str(file_path)
        page.df = pd.DataFrame(
            {
                "车辆所有人/企业": ["测试运输有限公司"],
                "负责人/法人代表": ["张三"],
                "地址": ["测试地址"],
                "电话": ["13800138000"],
            }
        )
        page.current_step = 3
        page.pipeline_running = True
        page.current_worker = worker

        with patch(
            "integrated_client.ui.workflow_page.os.replace",
            side_effect=PermissionError("file is locked"),
        ):
            self.assertFalse(page.force_shutdown(terminate_wait_ms=0))
        self.assertTrue(worker.stop_called)
        self.assertFalse(worker.terminated)
        self.assertTrue(worker.running)
        self.assertIn("尚未强制退出", page.last_force_shutdown_error)
        unchanged = load_workbook(file_path)
        unchanged_sheet = unchanged.active
        self.assertIsNone(unchanged_sheet.cell(2, 2).value)
        self.assertIsNone(unchanged_sheet.cell(2, 3).value)
        self.assertIsNone(unchanged_sheet.cell(2, 4).value)
        unchanged.close()
        self.assertEqual(
            list(file_path.parent.glob(f".{file_path.stem}.*.tmp{file_path.suffix}")),
            [],
        )

        worker.running = False
        page.current_worker = None
        page.current_step = 0
        page.df = pd.DataFrame()
        self.assertTrue(page.shutdown())
        page.close()

    def test_main_window_offers_save_and_force_exit_for_stuck_task(self):
        window = MainWindow(self.db, self.admin)
        captured = {}

        def execute_dialog(dialog):
            captured["title"] = dialog.text()
            captured["message"] = dialog.informativeText()
            captured["force_text"] = dialog.button(
                FramelessMessageBox.Save
            ).text()
            captured["force_object"] = dialog.button(
                FramelessMessageBox.Save
            ).objectName()
            captured["wait_text"] = dialog.button(
                FramelessMessageBox.Cancel
            ).text()
            captured["wait_default"] = dialog.button(
                FramelessMessageBox.Cancel
            ).isDefault()
            return FramelessMessageBox.Save

        with patch.object(
            window.workflow_page,
            "shutdown",
            return_value=False,
        ), patch.object(
            window.workflow_page,
            "force_shutdown",
            return_value=True,
        ) as force_shutdown, patch.object(
            window.workflow_page,
            "has_running_shutdown_threads",
            return_value=True,
        ), patch.object(
            window,
            "_schedule_hard_exit_fallback",
        ) as hard_exit, patch.object(
            FramelessMessageBox,
            "exec_",
            new=execute_dialog,
        ):
            self.assertTrue(window._shutdown_tools())

        force_shutdown.assert_called_once_with()
        hard_exit.assert_called_once_with()
        self.assertEqual(captured["title"], "任务仍在结束")
        self.assertIn("保存所有已完成记录和计时状态", captured["message"])
        self.assertIn("当前正在处理的一条记录", captured["message"])
        self.assertEqual(captured["force_text"], "保存并强制退出")
        self.assertEqual(captured["force_object"], "DangerButton")
        self.assertEqual(captured["wait_text"], "继续等待")
        self.assertTrue(captured["wait_default"])

        window.workflow_page.shutdown()
        window._prepared_to_close = True
        window.close()

    def test_main_window_contains_integrated_pages(self):
        window = MainWindow(self.db, self.admin)
        self.assertEqual(
            window.windowTitle(),
            f"{APP_NAME} - {self.admin.name_label}",
        )
        self.assertIs(window._pages["home"], window.dashboard_page)
        self.assertIs(window._pages["statistics"], window.statistics_page)
        self.assertIn("workflow", window._pages)
        self.assertIs(window._pages["workflow"], window.workflow_page)
        self.assertIs(
            window.workflow_page._timing_service,
            window.workflow_timing,
        )
        self.assertIn("本批次累计", window.workflow_page.timing_label.text())
        self.assertIs(window._pages["personal"], window.personal_center_page)
        self.assertNotIn("transport", window._pages)
        self.assertNotIn("aiqicha", window._pages)
        self.assertIn("statistics", window._pages)
        self.assertIn("accounts", window._pages)
        self.assertNotIn("statistics", window._nav_buttons)
        for key in (
            "data_station",
            "data_timing",
            "data_completion",
            "data_violation",
        ):
            self.assertIn(key, window._nav_buttons)
        self.assertNotIn("data_anomaly", window._nav_buttons)
        self.assertIn("workflow", window._nav_buttons)
        self.assertIn("personal", window._nav_buttons)
        self.assertEqual(window.sidebar.width(), 230)
        self.assertEqual(window.sidebar_brand_badge.text(), "逃")
        self.assertEqual(window.sidebar_brand_badge.objectName(), "BrandBadge")
        self.assertEqual(
            window.sidebar.findChild(QLabel, "BrandTitle").text(),
            "逃费车辆智能\n查询平台",
        )
        self.assertEqual(window.sidebar_role.text(), "管理员  ·  admin")
        self.assertEqual(window.sidebar_avatar.text(), "系")

        expected_nav = {
            "home": ("数据仪表盘", "nav-dashboard.svg"),
            "workflow": ("一键业务处理", "nav-workflow.svg"),
            "accounts": ("账号管理", "nav-accounts.svg"),
            "personal": ("系统设置", "nav-user.svg"),
        }
        for key, (label, icon_name) in expected_nav.items():
            button = window._nav_buttons[key]
            self.assertEqual(button.text(), label)
            self.assertEqual(button.iconSize(), QSize(19, 19))
            self.assertFalse(button.icon().isNull())
            self.assertTrue(Path(_control_asset_path(icon_name)).is_file())
        spec_text = (
            Path(__file__).resolve().parents[1] / "integrated_client.spec"
        ).read_text(encoding="utf-8")
        for _, icon_name in expected_nav.values():
            self.assertIn(icon_name, spec_text)
        self.assertIn("nav-data.svg", spec_text)
        self.assertEqual(window.data_nav_toggle.objectName(), "NavGroupButton")
        self.assertFalse(window.data_nav_toggle.isChecked())
        self.assertTrue(window.data_nav_container.isHidden())
        self.assertFalse(window.data_nav_toggle.icon().isNull())
        expected_data_nav = {
            "data_station": "全站分布",
            "data_timing": "用时效率",
            "data_completion": "完成类型",
            "data_violation": "违规原因",
        }
        for key, label in expected_data_nav.items():
            button = window._nav_buttons[key]
            self.assertEqual(button.text(), label)
            self.assertEqual(button.objectName(), "NavSubButton")
        for selector in (
            "QLabel#BrandBadge",
            "QFrame#SidebarProfile",
            "QLabel#SidebarAvatar",
            "QLabel#SidebarRole",
            "QPushButton#NavGroupButton",
            "QPushButton#NavSubButton",
            "QFrame#DashboardMetricCard",
            "border-left: 4px solid #8ce3d1",
        ):
            self.assertIn(selector, APP_STYLESHEET)
        self.assertIs(window.page_scroll_area.widget(), window.stack)
        self.assertTrue(window.page_scroll_area.widgetResizable())
        self.assertEqual(
            window.page_scroll_area.horizontalScrollBarPolicy(),
            Qt.ScrollBarAsNeeded,
        )
        self.assertEqual(
            window.page_scroll_area.verticalScrollBarPolicy(),
            Qt.ScrollBarAsNeeded,
        )
        self.assertEqual(window.stack.minimumSize(), window.PAGE_CANVAS_SIZE)
        self.assertIs(window.stack.currentWidget(), window.dashboard_page)
        self.assertTrue(window._nav_buttons["home"].isChecked())
        self.assertTrue(window.page_title.isHidden())
        self.assertEqual(window.page_title.text(), "")
        self.assertEqual(
            set(window.dashboard_page.metric_cards),
            {"total", "today", "phone", "no_phone", "time"},
        )
        self.assertEqual(
            [
                window.statistics_page.category_combo.itemData(index)
                for index in range(window.statistics_page.category_combo.count())
            ],
            ["station_distribution", "timing", "completion", "violation"],
        )
        self.assertEqual(
            window.statistics_page.category_combo.currentData(),
            "station_distribution",
        )
        self.assertFalse(window.statistics_page.station_combo.isEnabled())
        self.assertFalse(window.statistics_page.station_distribution_chart.isHidden())
        self.assertTrue(window.statistics_page.category_filter_group.isHidden())
        window.show_page("data_station")
        self.assertIs(window.stack.currentWidget(), window.statistics_page)
        self.assertTrue(window.data_nav_toggle.isChecked())
        self.assertFalse(window.data_nav_container.isHidden())
        self.assertTrue(window._nav_buttons["data_station"].isChecked())
        self.assertEqual(window.page_title.text(), "")
        window.show_page("data_timing")
        self.assertEqual(window.statistics_page.navigation_view, "timing")
        self.assertFalse(window.statistics_page.timing_section.isHidden())
        self.assertIs(
            window.statistics_page.detail_tabs.currentWidget(),
            window.statistics_page.chart_tab,
        )
        self.assertFalse(
            window.statistics_page.station_distribution_chart.isHidden()
        )
        self.assertEqual(
            window.statistics_page.station_distribution_chart._series_mode,
            "timing",
        )
        window.show_page("data_completion")
        self.assertEqual(window.statistics_page.navigation_view, "completion")
        self.assertTrue(window.statistics_page.timing_section.isHidden())
        self.assertFalse(window.statistics_page.anomaly_button.isHidden())
        window.statistics_page.anomaly_button.click()
        self.assertEqual(window.statistics_page.navigation_view, "anomaly")
        self.assertEqual(window.page_title.text(), "")
        window.statistics_page.anomaly_button.click()
        self.assertEqual(window.statistics_page.navigation_view, "completion")
        window.show_page("data_violation")
        self.assertEqual(window.statistics_page.navigation_view, "violation")
        window.show_page("home")
        self.assertIs(window.stack.currentWidget(), window.dashboard_page)
        self.assertFalse(hasattr(window, "top_identity"))
        self.assertEqual(window.sidebar_user.text(), "系统管理员")
        self.assertEqual(window.personal_center_page.identity_value.text(), "管理员")
        self.assertEqual(window.personal_center_page.name_value.text(), "系统管理员")
        self.assertEqual(window.personal_center_page.username_value.text(), "admin")
        password_dialog = PasswordDialog(
            self.db, self.admin.id, require_current=True
        )
        self.assertIsNotNone(password_dialog.current_password_edit)
        self.assertEqual(
            password_dialog.current_password_edit.placeholderText(),
            "请输入当前登录密码",
        )
        password_dialog.close()
        window.show_page("personal")
        self.assertIs(window.stack.currentWidget(), window.personal_center_page)
        self.assertTrue(window._nav_buttons["personal"].isChecked())
        self.assertEqual(window.page_title.text(), "")
        window.show_page("home")
        self.assertEqual(window.statistics_page.detail_tabs.count(), 2)
        self.assertFalse(window.statistics_page.detail_tabs.documentMode())
        self.assertIn(
            "border-top-left-radius: 0;",
            APP_STYLESHEET,
        )
        self.assertEqual(
            [
                window.statistics_page.detail_tabs.tabText(index)
                for index in range(window.statistics_page.detail_tabs.count())
            ],
            ["图表分析", "完整数据"],
        )
        self.assertIs(
            window.statistics_page.detail_tabs.widget(0),
            window.statistics_page.chart_tab,
        )
        self.assertFalse(hasattr(window.workflow_page, "browser_combo"))
        self.assertEqual(
            window.workflow_page.browser_info.text(),
            "Chromium（自动适配内置/系统）",
        )
        self.app.processEvents()
        self.assertEqual(
            window.workflow_page.mode_combo.view().window().frameShape(),
            QFrame.NoFrame,
        )
        workflow_popup = window.workflow_page.mode_combo.view().window()
        self.assertIs(
            workflow_popup.property("_intdemo_combo_owner"),
            window.workflow_page.mode_combo,
        )
        self.assertTrue(workflow_popup.windowFlags() & Qt.FramelessWindowHint)
        self.assertTrue(workflow_popup.windowFlags() & Qt.NoDropShadowWindowHint)
        self.assertTrue(workflow_popup.testAttribute(Qt.WA_TranslucentBackground))
        self.assertEqual(
            (
                workflow_popup.contentsMargins().left(),
                workflow_popup.contentsMargins().top(),
                workflow_popup.contentsMargins().right(),
                workflow_popup.contentsMargins().bottom(),
            ),
            (0, 0, 0, 0),
        )
        window.show()
        self.app.processEvents()
        self.assertEqual(
            [
                window.workflow_page.page_tabs.tabText(index)
                for index in range(window.workflow_page.page_tabs.count())
            ],
            ["业务处理", "运行设置"],
        )
        self.assertEqual(
            window.workflow_page.page_tabs.objectName(),
            "WorkflowTabs",
        )
        self.assertFalse(window.workflow_page.page_tabs.documentMode())
        self.assertIn("QTabWidget#WorkflowTabs::pane", APP_STYLESHEET)
        window.show_page("data_station")
        self.assertIs(
            window.statistics_page.detail_tabs.widget(1),
            window.statistics_page.data_tab,
        )
        self.assertFalse(hasattr(window.statistics_page, "recent_tab"))
        self.assertFalse(hasattr(window.statistics_page, "recent_table"))
        self.assertEqual(window.statistics_page.data_title.text(), "完整数据")
        self.assertEqual(
            window.statistics_page.data_description.text(),
            "全部站点 · 各站总计数占比与有电话数占比",
        )
        self.assertEqual(window.statistics_page.data_count_label.text(), "0 行")
        self.assertEqual(
            [
                window.statistics_page.summary_table.horizontalHeaderItem(column).text()
                for column in range(window.statistics_page.summary_table.columnCount())
            ],
            ["站点", "总计数", "总数占比", "有电话数", "有电话数占比"],
        )
        self.assertFalse(window.statistics_page.summary_table.showGrid())
        self.assertTrue(window.statistics_page.summary_table.verticalHeader().isHidden())
        self.assertEqual(
            window.statistics_page.summary_table.verticalHeader().defaultSectionSize(),
            42,
        )
        self.assertTrue(window.workflow_page.shutdown())
        window._prepared_to_close = True
        window.close()

    def test_online_admin_page_does_not_request_network_during_construction(self):
        class Api:
            def __init__(self):
                self.account_requests = 0

            def admin_accounts(self, _token):
                self.account_requests += 1
                raise AssertionError("startup must not make a blocking request")

        class State:
            is_online = True

        class Session:
            def __init__(self):
                self.api = Api()
                self.state = State()

            @staticmethod
            def access_token():
                return "test-token"

        session = Session()
        page = OnlineAccountPage(self.db, self.admin, session)
        self.assertEqual(session.api.account_requests, 0)
        self.assertFalse(page.edit_btn.isEnabled())
        page.deleteLater()

    def test_online_account_settings_exposes_device_limit(self):
        dialog = _AccountSettingsDialog()
        self.assertEqual(dialog.device_limit.minimum(), 1)
        self.assertEqual(dialog.device_limit.maximum(), 10000)
        self.assertEqual(dialog.device_limit.value(), 10000)
        dialog.device_limit.setValue(3)
        self.assertEqual(dialog.values()["device_limit"], 3)
        dialog.deleteLater()

    def test_sync_coordinator_uses_pyqt5_socket_state_without_crashing(self):
        class Engine:
            pass

        coordinator = SyncCoordinator(Engine())
        coordinator._running = True
        status = SyncStatus(
            state="online",
            pending_count=0,
            quarantined_count=0,
            last_sync_at="2026-07-27T22:00:00+08:00",
        )
        with patch.object(coordinator, "_connect_websocket") as connect:
            coordinator._on_worker_finished(status)
        connect.assert_called_once_with()
        self.assertFalse(coordinator._running)
        coordinator.stop()
        coordinator.deleteLater()

    def test_sync_coordinator_checks_immediately_when_server_blocks_connection(self):
        coordinator = SyncCoordinator(Mock())
        with patch.object(coordinator, "request_sync") as request_sync:
            coordinator._on_websocket_message(
                '{"type":"test_connection_blocked"}'
            )
        request_sync.assert_called_once_with()
        coordinator.stop()
        coordinator.deleteLater()

    def test_websocket_reconnect_keeps_test_block_in_offline_state(self):
        class Session:
            state = object()

            @staticmethod
            def access_token():
                raise ApiResponseError(
                    "test_connection_blocked",
                    "测试工具已断开连接",
                    status_code=503,
                    retryable=True,
                )

        class Engine:
            session = Session()

            @staticmethod
            def status(state, error):
                return SyncStatus(
                    state=state,
                    pending_count=0,
                    quarantined_count=0,
                    last_sync_at=None,
                    error=error,
                )

        coordinator = SyncCoordinator(Engine())
        coordinator._stopped = False
        statuses = []
        coordinator.status_changed.connect(statuses.append)

        coordinator._connect_websocket()

        self.assertEqual(statuses[-1].state, "offline")
        self.assertIn("断开连接", statuses[-1].error)
        coordinator.stop()
        coordinator.deleteLater()

    def test_sync_action_and_update_details_live_in_the_sidebar_and_settings(self):
        class FakeEngine:
            def __init__(self):
                self.current_status = SyncStatus(
                    state="online",
                    pending_count=0,
                    quarantined_count=0,
                    last_sync_at=None,
                )

            def status(self):
                return self.current_status

        class FakeSyncCoordinator(QObject):
            status_changed = pyqtSignal(object)
            data_changed = pyqtSignal()

            def __init__(self):
                super().__init__()
                self.engine = FakeEngine()
                self.retry_count = 0

            def retry_now(self):
                self.retry_count += 1

        class FakeUpdateCoordinator(QObject):
            update_available = pyqtSignal(object)
            state_changed = pyqtSignal(str, str)
            download_progress = pyqtSignal(int, int)
            download_completed = pyqtSignal(object)

            @staticmethod
            def check(manual=False):
                return manual

            @staticmethod
            def download(_update):
                return True

        sync = FakeSyncCoordinator()
        updates = FakeUpdateCoordinator()
        preferences = ClientPreferences(self.temp_dir.name)
        preferences.ignore_update("0.2.4")
        window = MainWindow(
            self.db,
            self.admin,
            sync_coordinator=sync,
            update_coordinator=updates,
            client_preferences=preferences,
        )
        status = SyncStatus(
            state="online",
            pending_count=35,
            quarantined_count=0,
            last_sync_at="2026-07-28T10:00:00+08:00",
        )
        window._update_sync_status(status)
        self.assertIn("待同步：35", window.sync_retry_button.text())
        self.assertTrue(
            window.sync_status_card.isAncestorOf(window.sync_retry_button)
        )
        window.sync_retry_button.click()
        self.assertEqual(sync.retry_count, 1)

        update = UpdateInfo(
            version="0.2.4",
            installer_url="https://example.com/updates/files/update.exe",
            installer_name="update.exe",
            sha256="0" * 64,
            size=10,
            notes="新增更新弹窗与系统设置。",
            mandatory=False,
        )
        window._update_available(update)
        self.assertEqual(
            window.personal_center_page.latest_version_value.text(),
            "v0.2.4",
        )
        self.assertIn(
            "新增更新弹窗",
            window.personal_center_page.update_notes.toPlainText(),
        )
        self.assertIn("●", window._nav_buttons["personal"].text())
        self.assertEqual(
            window._nav_buttons["personal"].toolTip(),
            "新版本发布!",
        )

        window._update_download_state("downloading", "正在下载 v0.2.4…")
        self.assertTrue(window.workflow_page.isEnabled())
        self.assertTrue(window.personal_center_page.isEnabled())
        window._update_download_progress(5, 10)
        self.assertEqual(
            window.personal_center_page.update_progress.value(),
            50,
        )
        with patch(
            "integrated_client.ui.main_window.QMessageBox.question",
            return_value=QMessageBox.No,
        ):
            window._update_downloaded(Path(self.temp_dir.name) / "update.exe")
        self.assertTrue(window.workflow_page.isEnabled())
        self.assertEqual(
            window.personal_center_page.update_action_btn.text(),
            "重启并安装",
        )

        window.workflow_page.shutdown()
        window._prepared_to_close = True
        window.close()

    def test_optional_update_cancel_releases_input_and_keeps_window_controls_available(self):
        class FakeUpdateCoordinator(QObject):
            update_available = pyqtSignal(object)
            state_changed = pyqtSignal(str, str)
            download_progress = pyqtSignal(int, int)
            download_speed = pyqtSignal(object)
            download_completed = pyqtSignal(object)

            def __init__(self):
                super().__init__()
                self.download_calls = 0
                self.cancel_calls = 0

            def download(self, _update):
                self.download_calls += 1
                return True

            def cancel_download(self):
                self.cancel_calls += 1
                return True

        updates = FakeUpdateCoordinator()
        window = MainWindow(
            self.db,
            self.admin,
            update_coordinator=updates,
        )
        window.show()
        self.app.processEvents()
        update = UpdateInfo(
            version="0.2.7",
            installer_url="https://example.com/updates/files/update.exe",
            installer_name="update.exe",
            sha256="0" * 64,
            size=100,
            notes="修复更新取消后的界面状态。",
            mandatory=False,
        )
        window.available_update = update
        window.personal_center_page.set_update_available(update)
        window._show_update_dialog(update)
        self.app.processEvents()

        dialog = window.update_dialog
        self.assertIsNotNone(dialog)
        self.assertEqual(dialog.cancel_button.text(), "稍后更新")
        self.assertFalse(dialog.background_button.isHidden())
        self.assertEqual(dialog.background_button.text(), "后台更新")
        self.assertEqual(dialog.update_button.text(), "立即更新")
        self.assertIsNone(QApplication.activeModalWidget())
        self.assertFalse(window.sidebar.isEnabled())
        self.assertFalse(window.page_scroll_area.isEnabled())
        for button in (
            window.window_controls.minimize_button,
            window.window_controls.maximize_button,
            window.window_controls.close_button,
        ):
            self.assertTrue(button.isEnabled())

        window._download_available_update()
        self.app.processEvents()
        self.assertEqual(updates.download_calls, 1)
        self.assertTrue(window._update_busy)
        self.assertTrue(window.sidebar.isEnabled())
        self.assertTrue(window.page_scroll_area.isEnabled())
        self.assertIsNone(QApplication.activeModalWidget())
        for button in (
            window.window_controls.minimize_button,
            window.window_controls.maximize_button,
            window.window_controls.close_button,
        ):
            self.assertTrue(button.isEnabled())

        with patch(
            "integrated_client.ui.main_window.QMessageBox.question",
            return_value=QMessageBox.No,
        ) as close_prompt:
            window.close()
            self.app.processEvents()
        self.assertTrue(window.isVisible())
        self.assertEqual(close_prompt.call_args.args[1], "停止更新并退出")

        window._cancel_update_download()
        self.assertEqual(updates.cancel_calls, 1)
        window._update_download_state(
            "download_cancelled",
            "更新下载已停止，可稍后继续重新下载。",
        )
        self.assertEqual(dialog.cancel_button.text(), "稍后更新")
        self.assertTrue(dialog.background_button.isEnabled())
        self.assertFalse(dialog.background_button.isHidden())
        dialog.close()
        self.app.processEvents()

        self.assertIsNone(window.update_dialog)
        self.assertIsNone(QApplication.activeModalWidget())
        self.assertFalse(window._update_busy)
        self.assertTrue(window.sidebar.isEnabled())
        self.assertTrue(window.page_scroll_area.isEnabled())
        for button in (
            window.window_controls.minimize_button,
            window.window_controls.maximize_button,
            window.window_controls.close_button,
        ):
            self.assertTrue(button.isEnabled())

        window._show_update_dialog(update)
        self.app.processEvents()
        background_dialog = window.update_dialog
        background_dialog.background_button.click()
        self.app.processEvents()
        self.assertEqual(updates.download_calls, 2)
        self.assertIsNone(window.update_dialog)
        self.assertTrue(window._update_busy)
        self.assertTrue(window.sidebar.isEnabled())
        self.assertTrue(window.page_scroll_area.isEnabled())
        self.assertIsNone(QApplication.activeModalWidget())
        window._update_download_state(
            "download_cancelled",
            "后台更新下载已停止。",
        )

        window.workflow_page.shutdown()
        window._prepared_to_close = True
        window.close()

    def test_update_prompt_enforces_mandatory_and_reports_progress(self):
        mandatory_update = UpdateInfo(
            version="0.2.4",
            installer_url="https://example.com/updates/files/update.exe",
            installer_name="update.exe",
            sha256="0" * 64,
            size=100,
            notes="必须安装的安全更新。",
            mandatory=True,
        )
        dialog = UpdatePromptDialog(mandatory_update)
        dialog.show()
        self.app.processEvents()
        self.assertTrue(dialog.ignore_button.isHidden())
        self.assertTrue(dialog.cancel_button.isHidden())
        self.assertTrue(dialog.background_button.isHidden())
        self.assertTrue(dialog.window_controls.close_button.isHidden())
        dialog.reject()
        self.app.processEvents()
        self.assertTrue(dialog.isVisible())

        dialog.begin_download()
        dialog.set_progress(40, 100)
        self.assertEqual(dialog.progress.value(), 40)
        self.assertFalse(dialog.update_button.isEnabled())
        dialog.set_downloaded()
        self.assertEqual(dialog.update_button.text(), "重启并安装")
        dialog.allow_close()
        dialog.reject()
        self.app.processEvents()
        self.assertFalse(dialog.isVisible())
        dialog.deleteLater()

    def test_login_preferences_are_saved_after_successful_login(self):
        class Store:
            is_available = True

            def __init__(self):
                self.saved = None
                self.cleared = False

            @staticmethod
            def load():
                return None

            def save(self, username, password, *, auto_login):
                self.saved = (username, password, auto_login)

            def clear(self):
                self.cleared = True

        self.db.change_password(
            self.admin.id,
            DEFAULT_ADMIN_PASSWORD,
            must_change=False,
        )
        store = Store()
        dialog = LoginDialog(self.db, credential_store=store)
        dialog.username_edit.setText(DEFAULT_ADMIN_USERNAME)
        dialog.password_edit.setText(DEFAULT_ADMIN_PASSWORD)
        dialog.auto_login_checkbox.setChecked(True)
        self.assertTrue(dialog.remember_password_checkbox.isChecked())

        dialog._login()

        self.assertEqual(dialog.result(), dialog.Accepted)
        self.assertEqual(
            store.saved,
            (
                DEFAULT_ADMIN_USERNAME,
                DEFAULT_ADMIN_PASSWORD,
                True,
            ),
        )
        dialog.deleteLater()

    def test_offline_login_enters_guest_without_credentials_or_account_access(self):
        store = Mock()
        store.is_available = True
        store.load.return_value = None
        session_manager = Mock()
        before_accounts = self.db.list_accounts()

        dialog = LoginDialog(
            self.db,
            session_manager=session_manager,
            configuration_error="服务器未配置",
            credential_store=store,
        )
        self.assertEqual(dialog.username_edit.text(), "")
        self.assertEqual(dialog.password_edit.text(), "")
        self.assertTrue(dialog.offline_login_btn.isEnabled())
        self.assertIn("游客", dialog.offline_login_btn.text())

        with patch.object(self.db, "authenticate") as authenticate:
            dialog.offline_login_btn.click()

        self.assertEqual(dialog.result(), dialog.Accepted)
        self.assertTrue(dialog.offline_business_mode)
        self.assertEqual(dialog.account.id, -1)
        self.assertEqual(dialog.account.username, "guest")
        self.assertEqual(dialog.account.name_label, "离线游客")
        authenticate.assert_not_called()
        session_manager.login.assert_not_called()
        session_manager.offline_login.assert_not_called()
        store.save.assert_not_called()
        store.clear.assert_not_called()
        self.assertEqual(self.db.list_accounts(), before_accounts)
        dialog.deleteLater()

    def test_guest_window_only_exposes_business_and_never_writes_metrics(self):
        dialog = LoginDialog(self.db)
        dialog._guest_login()
        guest = dialog.account

        tracked_tables = (
            "activity_events",
            "workflow_batches",
            "workflow_runs",
            "workflow_step_attempts",
            "workflow_timer_events",
            "sync_outbox",
        )

        def row_counts():
            with self.db._connect() as connection:
                return {
                    table: connection.execute(
                        f"SELECT COUNT(*) FROM {table}"
                    ).fetchone()[0]
                    for table in tracked_tables
                }

        before = row_counts()
        window = MainWindow(
            self.db,
            guest,
            session_manager=Mock(),
            sync_coordinator=Mock(),
            update_coordinator=Mock(),
            business_metrics_enabled=True,
            offline_business_mode=True,
        )

        self.assertFalse(window.business_metrics_enabled)
        self.assertIsNone(window.session_manager)
        self.assertIsNone(window.sync_coordinator)
        self.assertIsNone(window.update_coordinator)
        self.assertIsNone(window.workflow_timing)
        self.assertIsNone(window.dashboard_page)
        self.assertIsNone(window.statistics_page)
        self.assertIsNone(window.personal_center_page)
        self.assertEqual(set(window._pages), {"workflow"})
        self.assertEqual(set(window._nav_buttons), {"workflow"})
        self.assertIsNone(window.account_page)
        self.assertFalse(window.announcement_service_available)
        self.assertTrue(window.guest_logout_button.isVisibleTo(window.sidebar))
        self.assertTrue(
            window._record_workflow_summary(
                {WORKFLOW_TOTAL_METRIC: 9},
                source="unified_workflow",
                task_id="guest-task",
            )
        )
        window.workflow_page._record_workflow_stats()
        self.assertIn("不记录统计", window.workflow_page.log_text.toPlainText())
        self.assertEqual(row_counts(), before)
        self.assertEqual(self.db.list_accounts(), [self.admin])

        self.assertTrue(window.workflow_page.shutdown())
        window._prepared_to_close = True
        window.close()
        dialog.deleteLater()

    def test_guest_logout_does_not_clear_online_session_or_saved_profile(self):
        session_manager = Mock()
        session_manager.state = None
        controller = ApplicationController(
            self.app,
            self.db,
            session_manager=session_manager,
        )
        controller.window = Mock()
        controller.offline_business_mode = True
        controller.credential_store = Mock()

        with patch(
            "integrated_client.app_controller.QTimer.singleShot"
        ) as single_shot:
            controller._handle_logout()

        session_manager.logout.assert_not_called()
        session_manager.end_offline_session.assert_not_called()
        controller.credential_store.disable_auto_login.assert_not_called()
        single_shot.assert_called_once()

    def test_frameless_controls_are_embedded_without_an_extra_title_bar(self):
        window = MainWindow(self.db, self.admin)
        window.show()
        self.app.processEvents()

        self.assertTrue(window.windowFlags() & Qt.FramelessWindowHint)
        top_bar = window.window_controls.parentWidget()
        self.assertEqual(top_bar.objectName(), "TopBar")
        self.assertEqual(top_bar.height(), 62)
        self.assertEqual(window.centralWidget().layout().count(), 2)
        self.assertFalse(window.window_controls.minimize_button.isHidden())
        self.assertFalse(window.window_controls.maximize_button.isHidden())
        self.assertFalse(window.window_controls.close_button.isHidden())
        self.assertEqual(window.minimumSize().width(), 800)
        self.assertEqual(window.minimumSize().height(), 600)
        self.assertFalse(window.page_scroll_area.horizontalScrollBar().isVisible())
        self.assertFalse(window.page_scroll_area.verticalScrollBar().isVisible())
        window.resize(800, 600)
        self.app.processEvents()
        self.assertTrue(window.page_scroll_area.horizontalScrollBar().isVisible())
        self.assertTrue(window.page_scroll_area.verticalScrollBar().isVisible())
        self.assertIn("QScrollBar::handle:horizontal:hover", APP_STYLESHEET)
        self.assertIn("QScrollBar::handle:vertical:pressed", APP_STYLESHEET)
        native_limits = MINMAXINFO()
        window._update_minimum_track_size(native_limits)
        scale = max(1.0, float(window.devicePixelRatioF()))
        self.assertEqual(
            native_limits.ptMinTrackSize.x,
            math.ceil(window.minimumWidth() * scale),
        )
        self.assertEqual(
            native_limits.ptMinTrackSize.y,
            math.ceil(window.minimumHeight() * scale),
        )
        self.assertEqual(window._windows_nccalcsize_result(1), WVR_REDRAW)
        self.assertEqual(window._windows_nccalcsize_result(0), 0)

        caption_point = top_bar.mapTo(window, QPoint(320, top_bar.height() // 2))
        control_point = window.window_controls.mapTo(
            window, window.window_controls.rect().center()
        )
        self.assertEqual(window._window_hit_test(caption_point), HTCAPTION)
        self.assertEqual(window._window_hit_test(control_point), HTCLIENT)
        self.assertEqual(window._window_hit_test(QPoint(0, 0)), HTTOPLEFT)
        self.assertEqual(
            window._window_hit_test(QPoint(window.width() - 1, window.height() - 1)),
            HTBOTTOMRIGHT,
        )

        window.window_controls.toggle_maximized()
        self.app.processEvents()
        self.assertTrue(window.isMaximized())
        self.assertEqual(window.window_controls.maximize_button.toolTip(), "还原")
        window.window_controls.toggle_maximized()
        self.app.processEvents()
        self.assertFalse(window.isMaximized())

        login = LoginDialog(self.db)
        login.show()
        self.app.processEvents()
        self.assertTrue(login.windowFlags() & Qt.FramelessWindowHint)
        self.assertTrue(login.window_controls.minimize_button.isHidden())
        self.assertTrue(login.window_controls.maximize_button.isHidden())
        self.assertEqual(login.minimumSize(), login.maximumSize())
        self.assertEqual(login._window_hit_test(QPoint(60, 20)), HTCAPTION)
        self.assertEqual(login.username_edit.text(), "")
        self.assertEqual(login.remember_password_checkbox.text(), "记住密码")
        self.assertEqual(login.auto_login_checkbox.text(), "自动登录")
        self.assertFalse(
            any(
                "首次启动已创建管理员" in label.text()
                for label in login.findChildren(QLabel)
            )
        )
        login_title = login.findChild(QLabel, "PageTitle")
        login_button = login.findChild(QPushButton, "PrimaryButton")
        self.assertLessEqual(
            abs(
                login_title.geometry().top()
                - (login.height() - login_button.geometry().bottom())
            ),
            1,
        )
        login.close()

        message_box = FramelessMessageBox(window)
        message_box.setWindowTitle("确认操作")
        message_box.setText("确认操作")
        message_box.setInformativeText("确定继续吗？")
        message_box.setIcon(QMessageBox.Information)
        message_box.setStandardButtons(QMessageBox.Yes | QMessageBox.No)
        message_box.show()
        self.app.processEvents()
        self.assertTrue(message_box.windowFlags() & Qt.FramelessWindowHint)
        self.assertTrue(message_box.window_controls.minimize_button.isHidden())
        self.assertTrue(message_box.window_controls.maximize_button.isHidden())
        self.assertFalse(message_box.window_controls.close_button.isHidden())
        self.assertFalse(message_box.icon_label.pixmap().isNull())
        self.assertEqual(
            message_box.ICON_STYLES[QMessageBox.Information][1], "#1d8178"
        )
        self.assertEqual(message_box.button(QMessageBox.Yes).text(), "是")
        self.assertEqual(message_box.button(QMessageBox.No).text(), "否")
        message_box.close()

        window.workflow_page.shutdown()
        window._prepared_to_close = True
        window.close()

    def test_normal_user_permissions_and_stat_recording(self):
        self.assertEqual(self.db.ensure_default_station_users(), 5)
        user = self.db.create_account(
            "normal01", "Normal@123", "user", self.admin.id,
            display_name="测试站点",
        )
        window = MainWindow(self.db, user)
        self.assertNotIn("accounts", window._pages)
        self.assertNotIn("statistics", window._nav_buttons)
        self.assertIn("statistics", window._pages)
        for key in (
            "data_station",
            "data_timing",
            "data_completion",
            "data_violation",
        ):
            self.assertIn(key, window._nav_buttons)
        self.assertNotIn("data_anomaly", window._nav_buttons)
        self.assertIsNone(window.account_page)
        self.assertIs(window.stack.currentWidget(), window.workflow_page)
        self.assertTrue(window._nav_buttons["workflow"].isChecked())
        self.assertFalse(window._nav_buttons["home"].isChecked())
        self.assertTrue(window.page_title.isHidden())
        self.assertEqual(window.page_title.text(), "")
        window.show_page("data_anomaly")
        self.assertIs(window.stack.currentWidget(), window.workflow_page)
        self.assertTrue(window._nav_buttons["workflow"].isChecked())
        self.assertTrue(window.statistics_page.category_filter_group.isHidden())
        self.assertFalse(hasattr(window, "top_identity"))
        self.assertEqual(window.sidebar_user.text(), "测试站点")
        self.assertEqual(window.sidebar_role.text(), "用户  ·  normal01")
        self.assertEqual(window.sidebar_avatar.text(), "测")
        self.assertEqual(window.personal_center_page.identity_value.text(), "用户")
        self.assertEqual(window.personal_center_page.name_value.text(), "测试站点")
        self.assertEqual(window.personal_center_page.username_value.text(), "normal01")
        self.assertEqual(
            window.statistics_page.category_combo.currentData(),
            "station_distribution",
        )
        self.assertFalse(window.statistics_page.station_combo.isEnabled())
        self.assertIsNone(window.statistics_page.station_combo.currentData())
        station_names = [
            window.statistics_page.station_combo.itemText(index)
            for index in range(window.statistics_page.station_combo.count())
        ]
        self.assertEqual(station_names[0], "全部站点")
        self.assertIn("萝岗中心站", station_names)
        self.assertIn(user.name_label, station_names)
        window.statistics_page.category_combo.setCurrentIndex(
            window.statistics_page.category_combo.findData("completion")
        )
        self.app.processEvents()
        self.assertTrue(window.statistics_page.station_combo.isEnabled())
        self.assertEqual(window.statistics_page.station_combo.currentData(), user.id)

        counts = {
            WORKFLOW_TOTAL_METRIC: 6,
            WORKFLOW_EMPTY_METRIC: 1,
            WORKFLOW_NO_TRANSPORT_METRIC: 1,
            WORKFLOW_NO_OPERATION_METRIC: 1,
            WORKFLOW_INDIVIDUAL_METRIC: 1,
            WORKFLOW_NO_PHONE_METRIC: 1,
            WORKFLOW_HAS_PHONE_METRIC: 1,
        }
        self.assertTrue(
            window._record_workflow_summary(
                counts,
                "unified_workflow",
                {"file_name": "test.xlsx"},
                "ui-workflow-1",
            )
        )
        totals = {
            row["metric_key"]: row["total"]
            for row in self.db.get_user_totals(user.id)
        }
        self.assertEqual(totals[WORKFLOW_TOTAL_METRIC], 6)
        self.assertEqual(totals[WORKFLOW_EMPTY_METRIC], 1)
        self.assertEqual(totals[WORKFLOW_NO_TRANSPORT_METRIC], 1)
        self.assertEqual(totals[WORKFLOW_NO_OPERATION_METRIC], 1)
        self.assertEqual(totals[WORKFLOW_INDIVIDUAL_METRIC], 1)
        self.assertEqual(totals[WORKFLOW_NO_PHONE_METRIC], 1)
        self.assertEqual(totals[WORKFLOW_HAS_PHONE_METRIC], 1)
        window.statistics_page.refresh()
        self.assertEqual(
            window.statistics_page.distribution_chart._values[WORKFLOW_TOTAL_METRIC],
            6,
        )
        self.assertEqual(
            [
                window.statistics_page.summary_table.item(row, 0).text()
                for row in range(window.statistics_page.summary_table.rowCount())
            ],
            [
                "总计数",
                "有电话（有公司名）",
                "无电话（有公司名）",
                "个体经营",
                "无营运信息",
                "无运输证号",
            ],
        )
        self.assertTrue(
            all(
                window.statistics_page.kpi_layout.itemAtPosition(
                    index // 3,
                    index % 3,
                )
                is not None
                for index in range(6)
            )
        )
        self.assertIsNone(
            window.statistics_page.kpi_layout.itemAtPosition(0, 3)
        )

        accounts = {account.username: account for account in self.db.list_accounts()}
        luogang = accounts["luogang"]
        other_counts = dict(counts)
        other_counts[WORKFLOW_TOTAL_METRIC] = 4
        other_counts[WORKFLOW_EMPTY_METRIC] = 0
        other_counts[WORKFLOW_NO_TRANSPORT_METRIC] = 0
        self.db.record_activity_batch(
            luogang.id,
            other_counts,
            "unified_workflow",
            details={
                "violation_counts": {
                    "超限": {"total": 2, "has_phone": 1, "other": 1}
                }
            },
            task_id="normal-user-other-station",
        )
        station_index = window.statistics_page.station_combo.findData(luogang.id)
        window.statistics_page.station_combo.setCurrentIndex(station_index)
        self.app.processEvents()
        self.assertEqual(
            window.statistics_page.distribution_chart._values[WORKFLOW_TOTAL_METRIC],
            4,
        )
        window.statistics_page.category_combo.setCurrentIndex(
            window.statistics_page.category_combo.findData("violation")
        )
        self.app.processEvents()
        self.assertEqual(window.statistics_page.violation_chart._rows[0]["reason"], "超限")

        window.statistics_page.station_combo.setCurrentIndex(0)
        window.statistics_page.category_combo.setCurrentIndex(
            window.statistics_page.category_combo.findData("completion")
        )
        self.app.processEvents()
        self.assertEqual(
            window.statistics_page.distribution_chart._values[WORKFLOW_TOTAL_METRIC],
            10,
        )
        self.assertEqual(
            [
                window.statistics_page.summary_table.horizontalHeaderItem(column).text()
                for column in range(window.statistics_page.summary_table.columnCount())
            ],
            ["完成类型", "累计数量", "单位"],
        )
        self.assertEqual(window.statistics_page.summary_table.item(0, 0).text(), "总计数")
        self.assertEqual(window.statistics_page.summary_table.item(0, 1).text(), "10")
        self.assertTrue(window.workflow_page.shutdown())
        window._prepared_to_close = True
        window.close()

    def test_admin_can_filter_station_and_violation_category(self):
        self.assertEqual(self.db.ensure_default_station_users(), 5)
        accounts = {account.username: account for account in self.db.list_accounts()}
        luogang = accounts["luogang"]
        counts = {
            WORKFLOW_TOTAL_METRIC: 4,
            WORKFLOW_EMPTY_METRIC: 0,
            WORKFLOW_NO_TRANSPORT_METRIC: 0,
            WORKFLOW_NO_OPERATION_METRIC: 0,
            WORKFLOW_INDIVIDUAL_METRIC: 0,
            WORKFLOW_NO_PHONE_METRIC: 2,
            WORKFLOW_HAS_PHONE_METRIC: 2,
        }
        self.db.record_activity_batch(
            luogang.id,
            counts,
            "unified_workflow",
            details={
                "violation_counts": {
                    "超限": {"total": 3, "has_phone": 2, "other": 1}
                }
            },
            task_id="luogang-filter-test",
        )

        window = MainWindow(self.db, self.admin)
        station_names = [
            window.statistics_page.station_combo.itemText(index)
            for index in range(window.statistics_page.station_combo.count())
        ]
        self.assertEqual(
            station_names,
            ["全部站点"] + [display_name for display_name, _ in DEFAULT_STATION_USERS],
        )
        window.statistics_page.category_combo.setCurrentIndex(
            window.statistics_page.category_combo.findData("violation")
        )
        station_index = window.statistics_page.station_combo.findData(luogang.id)
        window.statistics_page.station_combo.setCurrentIndex(station_index)
        self.app.processEvents()

        violation_rows = window.statistics_page.violation_chart._rows
        self.assertEqual(violation_rows[0]["reason"], "超限")
        self.assertEqual(violation_rows[0]["has_phone"], 2)
        self.assertEqual(window.statistics_page.summary_table.rowCount(), 1)
        self.assertEqual(window.statistics_page.summary_table.item(0, 3).text(), "1")
        violation_headers = [
            window.statistics_page.summary_table.horizontalHeaderItem(column).text()
            for column in range(window.statistics_page.summary_table.columnCount())
        ]
        self.assertNotIn("账号", violation_headers)
        self.assertNotIn("状态", violation_headers)
        self.assertEqual(
            window.statistics_page.data_description.text(),
            "萝岗中心站 · 违规原因明细",
        )
        self.assertEqual(window.statistics_page.data_count_label.text(), "1 行")
        mode_button = window.statistics_page.violation_chart.mode_button
        self.assertIs(mode_button.parentWidget(), window.statistics_page.violation_chart)
        self.assertEqual(window.statistics_page.violation_chart._bar_mode, "share")
        self.assertEqual(mode_button.text(), "切换为电话拆分")
        mode_button.click()
        self.assertEqual(window.statistics_page.violation_chart._bar_mode, "split")
        self.assertEqual(mode_button.text(), "切换为原因占比")
        mode_button.click()
        self.assertEqual(window.statistics_page.violation_chart._bar_mode, "share")
        self.assertTrue(window.workflow_page.shutdown())
        window._prepared_to_close = True
        window.close()

    def test_only_admin_can_view_empty_anomalies_grouped_by_station(self):
        self.assertEqual(self.db.ensure_default_station_users(), 5)
        accounts = {account.username: account for account in self.db.list_accounts()}
        luogang = accounts["luogang"]
        taiping = accounts["taiping"]
        for station, total, empty, task_id in (
            (luogang, 5, 2, "admin-anomaly-luogang"),
            (taiping, 4, 1, "admin-anomaly-taiping"),
        ):
            self.db.record_activity_batch(
                station.id,
                {
                    WORKFLOW_TOTAL_METRIC: total,
                    WORKFLOW_EMPTY_METRIC: empty,
                    WORKFLOW_HAS_PHONE_METRIC: total - empty,
                },
                "unified_workflow",
                task_id=task_id,
            )

        page = StatisticsPage(self.db, self.admin)
        self.assertTrue(page.anomaly_button.isHidden())
        page.category_combo.setCurrentIndex(
            page.category_combo.findData("completion")
        )
        self.app.processEvents()

        self.assertFalse(page.anomaly_button.isHidden())
        self.assertEqual(page.anomaly_button.objectName(), "AnomalyButton")
        self.assertEqual(page.anomaly_button.text(), "异常数据（3）")
        page.anomaly_button.click()
        self.app.processEvents()

        self.assertTrue(page.anomaly_button.isChecked())
        self.assertEqual(page.anomaly_button.text(), "返回完成类型（3）")
        self.assertEqual(page.data_title.text(), "异常数据")
        self.assertIs(page.detail_tabs.currentWidget(), page.data_tab)
        self.assertFalse(page.detail_tabs.isTabEnabled(0))
        self.assertEqual(
            [
                page.summary_table.horizontalHeaderItem(column).text()
                for column in range(page.summary_table.columnCount())
            ],
            ["用户（站）", "登录账号", "异常条数", "本站总计数", "异常占比"],
        )
        anomaly_rows = {
            page.summary_table.item(row, 1).text(): [
                page.summary_table.item(row, column).text()
                for column in range(page.summary_table.columnCount())
            ]
            for row in range(page.summary_table.rowCount())
        }
        self.assertEqual(
            anomaly_rows["luogang"],
            ["萝岗中心站", "luogang", "2", "5", "40.0%"],
        )
        self.assertEqual(
            anomaly_rows["taiping"],
            ["太平中心站", "taiping", "1", "4", "25.0%"],
        )
        self.assertEqual(page.data_count_label.text(), "2 行")

        anomaly_export = Path(self.temp_dir.name) / "anomaly-dashboard.xlsx"
        export_result = page._save_dashboard_excel(anomaly_export)
        self.assertEqual(export_result["category_name"], "异常数据")
        export_workbook = load_workbook(anomaly_export)
        self.assertEqual(export_workbook.active["D2"].value, "异常数据")
        export_workbook.close()

        page.station_combo.setCurrentIndex(page.station_combo.findData(luogang.id))
        self.app.processEvents()
        self.assertEqual(page.anomaly_button.text(), "返回完成类型（2）")
        self.assertEqual(page.summary_table.rowCount(), 1)
        self.assertEqual(page.summary_table.item(0, 1).text(), "luogang")

        page.anomaly_button.click()
        self.app.processEvents()
        self.assertFalse(page.anomaly_button.isChecked())
        self.assertTrue(page.detail_tabs.isTabEnabled(0))
        self.assertIs(page.detail_tabs.currentWidget(), page.chart_tab)
        self.assertEqual(page.data_title.text(), "完整数据")
        self.assertNotIn(
            "空",
            [
                page.summary_table.item(row, 0).text()
                for row in range(page.summary_table.rowCount())
            ],
        )

        user_page = StatisticsPage(self.db, luogang)
        user_page.category_combo.setCurrentIndex(
            user_page.category_combo.findData("completion")
        )
        self.app.processEvents()
        self.assertFalse(hasattr(user_page, "anomaly_button"))

        for widget in (user_page, page):
            widget.close()
            widget.deleteLater()
        self.app.processEvents()

    def test_dashboard_date_range_filters_all_chart_categories(self):
        self.assertEqual(self.db.ensure_default_station_users(), 5)
        luogang = next(
            account for account in self.db.list_accounts()
            if account.username == "luogang"
        )
        for task_id, created_at, total, reason in (
            ("ui-date-early", "2025-01-01T09:00:00+08:00", 2, "早期异常"),
            ("ui-date-middle", "2025-01-15T09:00:00+08:00", 4, "中期异常"),
        ):
            self.db.record_activity_batch(
                luogang.id,
                {
                    WORKFLOW_TOTAL_METRIC: total,
                    WORKFLOW_HAS_PHONE_METRIC: total - 1,
                },
                "unified_workflow",
                details={
                    "violation_counts": {
                        reason: {
                            "total": total,
                            "has_phone": total - 1,
                            "other": 1,
                        }
                    }
                },
                task_id=task_id,
            )
            with self.db._connect() as conn:
                conn.execute(
                    "UPDATE activity_events SET created_at=? WHERE task_id=?",
                    (created_at, task_id),
                )

        window = MainWindow(self.db, self.admin)
        window.show()
        window.resize(800, 600)
        window.show_page("data_station")
        self.app.processEvents()
        page = window.statistics_page
        selector = page.date_range_selector
        self.assertTrue(selector.all_dates_check.isChecked())
        filter_groups = (
            page.station_filter_group,
            page.category_filter_group,
            page.date_filter_group,
        )
        self.assertEqual(
            [group.objectName() for group in filter_groups],
            ["DashboardFilterGroup"] * 3,
        )
        self.assertTrue(page.category_filter_group.isHidden())
        self.assertEqual(
            page.station_filter_group.y(),
            page.date_filter_group.y(),
        )
        self.assertLess(
            page.station_filter_group.x(),
            page.date_filter_group.x(),
        )
        self.assertEqual(page.date_filter_label.text(), "日期范围")
        self.assertEqual(
            page.date_filter_label.objectName(),
            "DashboardFilterLabel",
        )
        self.assertGreaterEqual(selector.width(), selector.minimumWidth())
        self.assertGreaterEqual(
            selector.all_dates_panel.width(),
            selector.all_dates_panel.minimumWidth(),
        )
        self.assertGreaterEqual(
            selector.all_dates_check.width(),
            selector.all_dates_check.sizeHint().width(),
        )
        for editor in (selector.start_edit, selector.end_edit):
            self.assertGreaterEqual(editor.minimumWidth(), 160)
            self.assertGreaterEqual(editor.width(), editor.minimumWidth())
            calendar = editor.calendarWidget()
            self.assertGreaterEqual(calendar.minimumWidth(), 332)
            self.assertEqual(calendar.firstDayOfWeek(), Qt.Monday)
            self.assertEqual(
                calendar.findChild(QToolButton, "qt_calendar_prevmonth").text(),
                "‹",
            )
            self.assertEqual(
                calendar.findChild(QToolButton, "qt_calendar_nextmonth").text(),
                "›",
            )
            calendar_view = calendar.findChild(
                QTableView, "qt_calendar_calendarview"
            )
            self.assertTrue(calendar_view.hasMouseTracking())
            self.assertTrue(calendar_view.viewport().hasMouseTracking())
            self.assertEqual(
                calendar_view.itemDelegate().__class__.__name__,
                "CalendarHoverDelegate",
            )

        selector.all_dates_check.setChecked(False)
        click_editor = selector.end_edit
        stable_date = QDate(2031, 7, 19)
        click_editor.setDate(stable_date)
        line_edit = click_editor.lineEdit()
        self.assertTrue(line_edit.isReadOnly())
        QTest.mouseClick(
            line_edit,
            Qt.LeftButton,
            pos=line_edit.rect().center(),
        )
        self.app.processEvents()
        self.assertTrue(click_editor.calendarWidget().isVisible())
        self.assertEqual(click_editor.date(), stable_date)
        QTest.keyClick(click_editor.calendarWidget(), Qt.Key_Escape)
        QTest.keyClicks(click_editor, "2024-01-01")
        QTest.keyClick(click_editor, Qt.Key_Up)
        self.assertEqual(click_editor.date(), stable_date)
        self.assertLess(
            selector.layout().indexOf(selector.end_edit),
            selector.layout().indexOf(selector.all_dates_panel),
        )
        self.assertLess(
            selector.layout().indexOf(selector.range_separator),
            selector.layout().indexOf(selector.all_dates_panel),
        )
        selector.start_edit.setDate(QDate(2025, 1, 15))
        selector.end_edit.setDate(QDate(2025, 1, 15))
        self.app.processEvents()

        station_rows = {
            row["username"]: row for row in page.station_distribution_chart._rows
        }
        self.assertEqual(station_rows["luogang"]["total"], 4)
        self.assertIn("2025-01-15 至 2025-01-15", page.data_description.text())

        page.category_combo.setCurrentIndex(
            page.category_combo.findData("completion")
        )
        page.station_combo.setCurrentIndex(page.station_combo.findData(luogang.id))
        self.app.processEvents()
        self.assertEqual(
            page.distribution_chart._values[WORKFLOW_TOTAL_METRIC],
            4,
        )

        page.category_combo.setCurrentIndex(page.category_combo.findData("violation"))
        self.app.processEvents()
        self.assertEqual(page.violation_chart._rows[0]["reason"], "中期异常")
        self.assertEqual(page.violation_chart._rows[0]["total"], 4)

        self.assertTrue(window.workflow_page.shutdown())
        window._prepared_to_close = True
        window.close()

    def test_dashboard_exports_current_station_and_date_range_to_excel(self):
        self.assertEqual(self.db.ensure_default_station_users(), 5)
        accounts = {account.username: account for account in self.db.list_accounts()}
        station = accounts["luogang"]
        self.db.record_activity_batch(
            station.id,
            {
                WORKFLOW_TOTAL_METRIC: 7,
                WORKFLOW_HAS_PHONE_METRIC: 4,
            },
            "unified_workflow",
            task_id="dashboard-excel-export",
        )
        with self.db._connect() as conn:
            conn.execute(
                "UPDATE activity_events SET created_at=? WHERE task_id=?",
                ("2025-03-15 10:00:00", "dashboard-excel-export"),
            )

        window = MainWindow(self.db, self.admin)
        page = window.statistics_page
        page.category_combo.setCurrentIndex(
            page.category_combo.findData("completion")
        )
        page.station_combo.setCurrentIndex(
            page.station_combo.findData(station.id)
        )
        selector = page.date_range_selector
        selector.all_dates_check.setChecked(False)
        selector.start_edit.setDate(QDate(2025, 3, 1))
        selector.end_edit.setDate(QDate(2025, 3, 31))
        self.app.processEvents()

        self.assertEqual(page.export_btn.text(), "导出 Excel")
        self.assertEqual(page.export_btn.objectName(), "PrimaryButton")
        export_base = Path(self.temp_dir.name) / "dashboard-export"
        with patch(
            "integrated_client.ui.statistics_page.QFileDialog.getSaveFileName",
            return_value=(str(export_base), "Excel 工作簿 (*.xlsx)"),
        ), patch(
            "integrated_client.ui.statistics_page.QMessageBox.information"
        ) as information:
            page._export_dashboard_excel()

        export_path = export_base.with_suffix(".xlsx")
        self.assertTrue(export_path.exists())
        workbook = load_workbook(export_path, data_only=True)
        self.assertEqual(workbook.sheetnames, ["仪表盘数据"])
        sheet = workbook["仪表盘数据"]
        self.assertEqual(sheet["A1"].value, "数据仪表盘导出")
        self.assertEqual(sheet["B2"].value, station.name_label)
        self.assertEqual(sheet["D2"].value, "按完成类型")
        self.assertEqual(sheet["F2"].value, "2025-03-01 至 2025-03-31")
        self.assertEqual(
            [sheet.cell(5, column).value for column in range(1, 6)],
            ["站名", "时间范围", "完成类型", "累计数量", "单位"],
        )
        self.assertEqual(sheet["A6"].value, station.name_label)
        self.assertEqual(sheet["B6"].value, "2025-03-01 至 2025-03-31")
        self.assertEqual(sheet["C6"].value, "总计数")
        self.assertEqual(sheet["D6"].value, 7)
        self.assertEqual(sheet["E6"].value, "条")
        self.assertEqual(sheet.freeze_panes, "A6")
        self.assertGreaterEqual(sheet.column_dimensions["F"].width, 25)
        workbook.close()
        information.assert_called_once()

        self.assertTrue(window.workflow_page.shutdown())
        window._prepared_to_close = True
        window.close()

    def test_dashboard_exports_all_stations_with_station_sheets(self):
        self.assertEqual(self.db.ensure_default_station_users(), 5)
        accounts = {
            account.username: account
            for account in self.db.list_accounts()
        }
        luogang = accounts["luogang"]
        taiping = accounts["taiping"]
        self.db.record_activity_batch(
            luogang.id,
            {
                WORKFLOW_TOTAL_METRIC: 7,
                WORKFLOW_HAS_PHONE_METRIC: 4,
            },
            "unified_workflow",
            task_id="all-stations-export-luogang",
        )
        self.db.record_activity_batch(
            taiping.id,
            {
                WORKFLOW_TOTAL_METRIC: 3,
                WORKFLOW_HAS_PHONE_METRIC: 2,
            },
            "unified_workflow",
            task_id="all-stations-export-taiping",
        )

        window = MainWindow(self.db, self.admin)
        page = window.statistics_page
        page.category_combo.setCurrentIndex(
            page.category_combo.findData("completion")
        )
        page.station_combo.setCurrentIndex(0)
        self.app.processEvents()

        export_path = Path(self.temp_dir.name) / "all-stations-export.xlsx"
        result = page._save_dashboard_excel(export_path)
        self.assertEqual(result["station_name"], "全部站点")
        self.assertEqual(result["station_sheet_count"], 5)

        workbook = load_workbook(export_path, data_only=True)
        station_names = [
            account.name_label
            for account in page._export_station_accounts()
        ]
        self.assertEqual(
            workbook.sheetnames,
            ["仪表盘数据", *station_names],
        )
        summary = workbook["仪表盘数据"]
        self.assertEqual(summary["C6"].value, "总计数")
        self.assertEqual(summary["D6"].value, 10)

        luogang_sheet = workbook[luogang.name_label]
        self.assertIn(luogang.name_label, luogang_sheet["A1"].value)
        self.assertEqual(
            [
                luogang_sheet.cell(4, column).value
                for column in range(1, 4)
            ],
            ["完成类型", "累计数量", "单位"],
        )
        self.assertEqual(luogang_sheet["A5"].value, "总计数")
        self.assertEqual(luogang_sheet["B5"].value, 7)
        self.assertEqual(luogang_sheet.freeze_panes, "A5")

        taiping_sheet = workbook[taiping.name_label]
        self.assertEqual(taiping_sheet["A5"].value, "总计数")
        self.assertEqual(taiping_sheet["B5"].value, 3)
        workbook.close()

        category_headers = {
            "timing": ["指标", "数值", "统计说明"],
            "violation": ["违规原因", "总计", "有电话", "其他数据", "有电话占比"],
            "station_distribution": ["指标", "数值", "单位 / 统计说明"],
        }
        for category, expected_headers in category_headers.items():
            page.category_combo.setCurrentIndex(
                page.category_combo.findData(category)
            )
            if page.station_combo.isEnabled():
                page.station_combo.setCurrentIndex(0)
            self.app.processEvents()
            category_export = (
                Path(self.temp_dir.name)
                / f"all-stations-{category}.xlsx"
            )
            category_result = page._save_dashboard_excel(category_export)
            self.assertEqual(category_result["station_sheet_count"], 5)
            category_workbook = load_workbook(category_export, data_only=True)
            first_station_sheet = category_workbook[station_names[0]]
            self.assertEqual(
                [
                    first_station_sheet.cell(4, column).value
                    for column in range(1, len(expected_headers) + 1)
                ],
                expected_headers,
            )
            category_workbook.close()

        page.category_combo.setCurrentIndex(
            page.category_combo.findData("completion")
        )
        page.station_combo.setCurrentIndex(0)
        page._toggle_anomaly_view(True)
        self.app.processEvents()
        anomaly_export = (
            Path(self.temp_dir.name)
            / "all-stations-anomaly.xlsx"
        )
        anomaly_result = page._save_dashboard_excel(anomaly_export)
        self.assertEqual(anomaly_result["station_sheet_count"], 5)
        anomaly_workbook = load_workbook(anomaly_export, data_only=True)
        self.assertEqual(
            [
                anomaly_workbook[station_names[0]].cell(4, column).value
                for column in range(1, 6)
            ],
            ["用户（站）", "登录账号", "异常条数", "本站总计数", "异常占比"],
        )
        anomaly_workbook.close()

        self.assertTrue(window.workflow_page.shutdown())
        window._prepared_to_close = True
        window.close()

    def test_dashboard_shows_timing_efficiency_kpis(self):
        self.assertEqual(self.db.ensure_default_station_users(), 5)
        station = next(
            account
            for account in self.db.list_accounts()
            if account.username == "luogang"
        )
        now = 1000.0
        dataframe = pd.DataFrame(
            {
                "车辆标识": ["粤A12345_黄色"] * 4,
                "已协助补缴": [""] * 4,
                "原因": [""] * 4,
            }
        )
        service = WorkflowTimingService(
            self.db,
            station.id,
            clock=lambda: now,
        )
        source_path = Path(self.temp_dir.name) / "dashboard-timing.xlsx"
        service.start_run(source_path, dataframe, run_id="dashboard-timing-run")
        service.start_step(1)
        now += 120.125
        service.pause("等待登录")
        now += 60.250
        service.resume("登录完成")
        now += 60.375
        service.finish_run("succeeded")
        self.db.record_activity_batch(
            station.id,
            {WORKFLOW_TOTAL_METRIC: 4},
            source="unified_workflow",
            task_id="dashboard-timing-run",
        )

        page = StatisticsPage(self.db, self.admin)
        page.category_combo.setCurrentIndex(
            page.category_combo.findData("completion")
        )
        page.station_combo.setCurrentIndex(page.station_combo.findData(station.id))
        self.app.processEvents()

        def card_value(label):
            labels = page.timing_kpi_cards[label].findChildren(QLabel)
            return labels[-1].text()

        self.assertEqual(card_value("总用时"), "0.1 小时")
        self.assertEqual(card_value("有效用时"), "0.1 小时")
        self.assertEqual(card_value("平均处理效率"), "60 条/小时")
        self.assertEqual(card_value("较纯人工效率提升"), "27.6 %")
        self.assertEqual(page.timing_scope_label.text(), "完成数据：4 条")

        page.resize(1220, 760)
        page.show()
        self.app.processEvents()
        total_card = page.timing_kpi_cards["总用时"]
        total_point = total_card.rect().center()
        total_event = QMouseEvent(
            QEvent.MouseMove,
            total_point,
            total_card.mapToGlobal(total_point),
            Qt.NoButton,
            Qt.NoButton,
            Qt.NoModifier,
        )
        total_card.mouseMoveEvent(total_event)
        self.assertTrue(total_card._hover_card.isVisible())
        self.assertEqual(
            total_card._hover_card.details,
            [
                ("精确用时", "00:04:00.750"),
                ("有效用时", "00:03:00.500"),
                ("暂停等待", "00:01:00.250"),
            ],
        )
        self.assertEqual(total_card._hover_card._detail_row_count, 3)
        hover_keys = total_card._hover_card.findChildren(QLabel, "HoverCardKey")
        self.assertTrue(hover_keys)
        self.assertTrue(all(label.text().endswith("：") for label in hover_keys))
        self.assertTrue(
            all(label.contentsMargins().right() >= 4 for label in hover_keys)
        )
        self.assertEqual(total_card.toolTip(), "")

        average_card = page.timing_kpi_cards["平均处理效率"]
        average_point = average_card.rect().center()
        average_event = QMouseEvent(
            QEvent.MouseMove,
            average_point,
            average_card.mapToGlobal(average_point),
            Qt.NoButton,
            Qt.NoButton,
            Qt.NoModifier,
        )
        average_card.mouseMoveEvent(average_event)
        self.assertEqual(
            average_card._hover_card.details,
            [
                ("平均耗时", "60.2 秒/条"),
                ("精确平均", "00:01:00.188"),
                ("完成数据", "4 条"),
            ],
        )

        efficiency_card = page.timing_kpi_cards["较纯人工效率提升"]
        efficiency_point = efficiency_card.rect().center()
        efficiency_event = QMouseEvent(
            QEvent.MouseMove,
            efficiency_point,
            efficiency_card.mapToGlobal(efficiency_point),
            Qt.NoButton,
            Qt.NoButton,
            Qt.NoModifier,
        )
        efficiency_card.mouseMoveEvent(efficiency_event)
        self.assertEqual(
            efficiency_card._hover_card.details,
            [
                ("总用时", "00:04:00.750"),
                ("人工预计用时", "00:05:32.308"),
                ("完成数据", "4 条"),
            ],
        )

        page.category_combo.setCurrentIndex(
            page.category_combo.findData("timing")
        )
        self.app.processEvents()
        timing_metrics = [
            page.summary_table.item(row, 0).text()
            for row in range(page.summary_table.rowCount())
        ]
        self.assertNotIn("计时运行", timing_metrics)
        efficiency_row = timing_metrics.index("平均处理效率")
        self.assertEqual(
            page.summary_table.item(efficiency_row, 1).text(),
            "60 条/小时",
        )
        self.assertEqual(
            page.summary_table.item(efficiency_row, 2).text(),
            "完成数据÷总用时",
        )
        timing_text = " ".join(
            page.summary_table.item(row, column).text()
            for row in range(page.summary_table.rowCount())
            for column in range(page.summary_table.columnCount())
        )
        self.assertNotIn("260 条", timing_text)
        self.assertNotIn("6 小时", timing_text)

        page.close()
        page.deleteLater()
        self.app.processEvents()

    def test_workflow_settings_browser_check_and_collapsible_panels(self):
        page = WorkflowPage()
        page.resize(1400, 900)
        page.show()
        self.app.processEvents()
        self.assertIs(page.page_tabs.widget(0), page.workflow_tab)
        self.assertIs(page.page_tabs.widget(1), page.settings_tab)
        self.assertEqual(page.browser_check_state, "unchecked")
        self.assertIn("尚未检测", page.browser_status_label.text())
        self.assertTrue(page.browser_detail_label.isHidden())
        self.assertEqual(page.browser_detail_toggle_btn.text(), "查看详情")
        for selector in (
            "QComboBox::down-arrow",
            "QComboBox::drop-down:on",
            "QCheckBox::indicator:checked",
            "QSpinBox::up-button",
        ):
            self.assertIn(selector, APP_STYLESHEET)
        for asset_name in (
            "check.svg",
            "minus.svg",
            "chevron-down.svg",
            "chevron-up.svg",
        ):
            self.assertTrue(Path(_control_asset_path(asset_name)).is_file())
        self.assertEqual(
            [
                page.mode_card.objectName(),
                page.captcha_card.objectName(),
                page.query_card.objectName(),
            ],
            ["SettingCard", "SettingCard", "SettingCard"],
        )
        self.assertIn("人工接管", page.mode_hint.text())
        page.mode_combo.setCurrentIndex(page.mode_combo.findData(True))
        self.assertIn("全自动运行", page.mode_hint.text())
        self.assertFalse(page.manual_captcha.isEnabled())
        page.mode_combo.setCurrentIndex(page.mode_combo.findData(False))

        with patch.object(page, "check_browser") as automatic_check:
            page.page_tabs.setCurrentWidget(page.settings_tab)
            page.page_tabs.setCurrentWidget(page.workflow_tab)
            page.page_tabs.setCurrentWidget(page.settings_tab)
        automatic_check.assert_called_once_with()

        with patch(
            "integrated_client.ui.workflow_page.check_builtin_chromium",
            return_value=(r"C:\\browser\\chrome.exe", "123.0"),
        ) as browser_check:
            page.check_browser()
            worker = page.browser_check_worker
            self.assertIsNotNone(worker)
            self.assertTrue(worker.wait(3000))
            self.app.processEvents()

        browser_check.assert_called_once_with()
        self.assertEqual(page.browser_check_state, "ready")
        self.assertIn("运行正常", page.browser_status_label.text())
        self.assertIn("Chromium 123.0", page.browser_status_label.text())
        self.assertIn("chrome.exe", page.browser_detail_label.text())
        self.assertTrue(page.browser_detail_label.isHidden())
        page.browser_detail_toggle_btn.click()
        self.assertFalse(page.browser_detail_label.isHidden())
        self.assertEqual(page.browser_detail_toggle_btn.text(), "收起详情")
        page.browser_detail_toggle_btn.click()
        self.assertTrue(page.browser_detail_label.isHidden())
        self.assertEqual(page.browser_detail_toggle_btn.text(), "查看详情")

        preview_header = page.preview_panel.header_layout
        self.assertLess(
            preview_header.indexOf(page.open_workbook_btn),
            preview_header.indexOf(page.open_workbook_folder_btn),
        )
        self.assertLess(
            preview_header.indexOf(page.open_workbook_folder_btn),
            preview_header.indexOf(page.preview_toggle_btn),
        )
        self.assertFalse(page.open_workbook_btn.isEnabled())
        self.assertFalse(page.open_workbook_folder_btn.isEnabled())

        workbook_path = Path(self.temp_dir.name) / "preview-actions.xlsx"
        workbook = Workbook()
        workbook.active.append(["车辆标识", "已协助补缴"])
        workbook.active.append(["粤A12345", ""])
        workbook.save(workbook_path)
        workbook.close()
        page.file_path = str(workbook_path)
        page.file_edit.setText(str(workbook_path))
        self.assertTrue(page._reload_preview(force=True))
        self.assertTrue(page.open_workbook_btn.isEnabled())
        self.assertTrue(page.open_workbook_folder_btn.isEnabled())
        with patch(
            "integrated_client.ui.workflow_page.QDesktopServices.openUrl",
            return_value=True,
        ) as open_url:
            page.open_workbook_btn.click()
            page.open_workbook_folder_btn.click()
        opened_paths = [
            Path(invocation.args[0].toLocalFile()).resolve()
            for invocation in open_url.call_args_list
        ]
        self.assertEqual(
            opened_paths,
            [workbook_path.resolve(), workbook_path.parent.resolve()],
        )

        self.assertTrue(page.preview_panel.is_expanded())
        self.assertTrue(page.log_panel.is_expanded())
        file_group_height = page.file_group.height()
        step_card_heights = [card.height() for card in page.step_cards]
        page.preview_toggle_btn.click()
        self.assertFalse(page.preview_panel.is_expanded())
        self.assertTrue(page.table.isHidden())
        self.assertIn("展开", page.preview_toggle_btn.text())
        page.preview_toggle_btn.click()
        self.assertTrue(page.preview_panel.is_expanded())
        self.assertFalse(page.table.isHidden())

        page.preview_toggle_btn.click()
        page.log_toggle_btn.click()
        self.assertFalse(page.log_panel.is_expanded())
        self.assertTrue(page.log_text.isHidden())
        self.assertIn("展开", page.log_toggle_btn.text())
        self.app.processEvents()
        self.assertFalse(page.collapsed_content_spacer.isHidden())
        self.assertEqual(page.file_group.height(), file_group_height)
        self.assertEqual(
            [card.height() for card in page.step_cards],
            step_card_heights,
        )
        page.log_toggle_btn.click()
        self.assertTrue(page.log_panel.is_expanded())
        self.assertFalse(page.log_text.isHidden())
        self.assertTrue(page.collapsed_content_spacer.isHidden())
        self.assertTrue(page.shutdown())
        page.close()

    def test_tencent_docs_button_reuses_account_link_without_starting_timing(self):
        preferences = ClientPreferences(self.temp_dir.name)
        old_url = "https://docs.qq.com/sheet/old"
        new_url = "https://docs.qq.com/sheet/new"
        preferences.set_tencent_document_url("admin", old_url)
        timing = WorkflowTimingService(self.db, self.admin.id)
        page = WorkflowPage(
            timing_service=timing,
            client_preferences=preferences,
            account_key="admin",
        )

        self.assertEqual(page.tencent_docs_btn.text(), "导入腾讯文档")
        file_layout = page.file_group.layout()
        self.assertEqual(
            file_layout.indexOf(page.tencent_docs_btn),
            file_layout.indexOf(page.choose_btn) + 1,
        )
        self.assertFalse(timing.is_active)

        with patch(
            "integrated_client.ui.workflow_page.TencentDocsLinkDialog",
        ) as dialog_type, patch(
            "integrated_client.ui.workflow_page.get_builtin_chromium_path",
            return_value=r"C:\browser\chrome.exe",
        ), patch.object(
            page,
            "_start_tencent_docs_import",
        ) as start_import:
            dialog = dialog_type.return_value
            dialog.Accepted = 1
            dialog.exec_.return_value = 1
            dialog.value.return_value = new_url
            page.tencent_docs_btn.click()

        dialog_type.assert_called_once_with(old_url, page)
        start_import.assert_called_once_with(new_url)
        self.assertEqual(
            ClientPreferences(self.temp_dir.name).tencent_document_url(
                "ADMIN"
            ),
            new_url,
        )
        self.assertFalse(timing.is_active)
        self.assertTrue(page.shutdown())
        page.close()

    def test_tencent_docs_dialog_displays_percentage_progress(self):
        dialog = TencentDocsProgressDialog()
        self.assertEqual(dialog.progress.minimum(), 0)
        self.assertEqual(dialog.progress.maximum(), 100)
        self.assertTrue(dialog.progress.isTextVisible())
        self.assertEqual(dialog.progress.value(), 0)
        dialog.set_progress(63)
        self.assertEqual(dialog.progress.value(), 63)
        dialog.set_progress(150)
        self.assertEqual(dialog.progress.value(), 100)
        dialog.finish()
        dialog.close()

    def test_workflow_run_settings_are_persisted_per_program_account(self):
        preferences = ClientPreferences(self.temp_dir.name)
        page = WorkflowPage(
            client_preferences=preferences,
            account_key="Admin",
        )
        page.mode_combo.setCurrentIndex(page.mode_combo.findData(True))
        page.manual_captcha.setChecked(False)
        page.auto_continue.setChecked(False)
        page.only_yellow.setChecked(False)
        page.infinite_captcha.setChecked(True)
        page.page_retry.setValue(17)
        page.captcha_retry.setValue(56)

        saved = ClientPreferences(
            self.temp_dir.name
        ).workflow_run_settings("ADMIN")
        self.assertEqual(
            saved,
            {
                "auto_mode": True,
                "manual_captcha": False,
                "auto_continue": False,
                "only_yellow": False,
                "infinite_captcha": True,
                "page_retry": 17,
                "captcha_retry": 56,
            },
        )
        self.assertTrue(page.shutdown())
        page.close()

        restored = WorkflowPage(
            client_preferences=ClientPreferences(self.temp_dir.name),
            account_key="admin",
        )
        self.assertTrue(bool(restored.mode_combo.currentData()))
        self.assertFalse(restored.manual_captcha.isChecked())
        self.assertFalse(restored.auto_continue.isChecked())
        self.assertFalse(restored.only_yellow.isChecked())
        self.assertTrue(restored.infinite_captcha.isChecked())
        self.assertEqual(restored.page_retry.value(), 17)
        self.assertEqual(restored.captcha_retry.value(), 56)

        other_account = WorkflowPage(
            client_preferences=ClientPreferences(self.temp_dir.name),
            account_key="station",
        )
        self.assertFalse(bool(other_account.mode_combo.currentData()))
        self.assertTrue(other_account.manual_captcha.isChecked())
        self.assertTrue(other_account.only_yellow.isChecked())
        self.assertEqual(other_account.page_retry.value(), 5)
        self.assertTrue(restored.shutdown())
        self.assertTrue(other_account.shutdown())
        restored.close()
        other_account.close()

    def test_tencent_docs_result_is_automatically_loaded_without_timing(self):
        workbook_path = Path(self.temp_dir.name) / "tencent-pending.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["车辆标识", "公司名称", "已协助补缴"])
        worksheet.append(["粤A12345", "", "否"])
        workbook.save(workbook_path)
        workbook.close()

        timing = WorkflowTimingService(self.db, self.admin.id)
        page = WorkflowPage(timing_service=timing)
        worker = object()
        page.tencent_import_worker = worker

        class Dialog:
            finished = False

            def finish(self):
                self.finished = True

        dialog = Dialog()
        result = TencentDocsImportResult(
            path=workbook_path,
            copied_rows=1,
            source_start_row=10,
            source_row_count=9,
            company_header="公司名称",
        )
        with patch(
            "integrated_client.ui.workflow_page.QMessageBox.information"
        ) as information:
            page._tencent_docs_import_succeeded(
                worker,
                dialog,
                result,
            )

        self.assertTrue(dialog.finished)
        self.assertEqual(page.file_path, str(workbook_path))
        self.assertEqual(len(page.df), 1)
        self.assertEqual(page.df.iloc[0]["车辆标识"], "粤A12345")
        self.assertFalse(timing.is_active)
        self.assertIn("第 10 行", information.call_args.args[2])
        page.tencent_import_worker = None
        self.assertTrue(page.shutdown())
        page.close()

    def test_workflow_timing_refreshes_milliseconds_independently(self):
        class TimingService:
            is_active = True

            def __init__(self):
                self.heartbeat_count = 0

            @staticmethod
            def format_duration(milliseconds):
                return WorkflowTimingService.format_duration(milliseconds)

            @staticmethod
            def snapshot():
                return {
                    "state": "running",
                    "run_active_ms": 57,
                    "run_paused_ms": 0,
                    "batch_active_ms": 1_057,
                }

            def heartbeat(self):
                self.heartbeat_count += 1

        service = TimingService()
        page = WorkflowPage(timing_service=service)

        self.assertEqual(
            page.timing_timer.interval(),
            page.TIMING_DISPLAY_INTERVAL_MS,
        )
        self.assertEqual(page.TIMING_DISPLAY_INTERVAL_MS, 16)
        self.assertEqual(page.timing_timer.timerType(), Qt.PreciseTimer)
        displayed_last_digits = {
            service.format_duration(
                tick * page.TIMING_DISPLAY_INTERVAL_MS
            )[-1]
            for tick in range(10)
        }
        self.assertGreater(len(displayed_last_digits), 2)
        self.assertEqual(
            page.timing_heartbeat_timer.interval(),
            page.TIMING_HEARTBEAT_INTERVAL_MS,
        )
        self.assertEqual(page.TIMING_HEARTBEAT_INTERVAL_MS, 5_000)

        page._timing_tick()
        self.assertIn("本次有效用时 00:00:00.057", page.timing_label.text())
        self.assertIn("本批次累计 00:00:01.057", page.timing_label.text())
        self.assertEqual(service.heartbeat_count, 0)

        page._timing_heartbeat()
        self.assertEqual(service.heartbeat_count, 1)

        service.is_active = False
        self.assertTrue(page.shutdown())
        page.close()

    def test_aiqicha_login_wait_pauses_and_resumes_business_timing(self):
        class TimingService:
            is_active = False

            def __init__(self):
                self.pause_reasons = []
                self.resume_reasons = []

            def pause(self, reason):
                self.pause_reasons.append(reason)

            def resume(self, reason):
                self.resume_reasons.append(reason)

        timing = TimingService()
        page = WorkflowPage(timing_service=timing)
        worker = QueryWorker(
            pd.DataFrame({"车辆所有人/企业": []}),
            "车辆所有人/企业",
        )
        page.current_worker = worker
        page.current_step = 3
        with patch.object(
            page,
            "_refresh_timing_label",
        ), patch.object(
            worker,
            "isRunning",
            return_value=True,
        ):
            page._on_login_required()
            self.assertTrue(page.awaiting_login)
            self.assertEqual(
                timing.pause_reasons,
                ["等待用户登录爱企查"],
            )
            page.continue_pipeline()

        self.assertTrue(worker._login_wait.is_set())
        self.assertFalse(page.awaiting_login)
        self.assertEqual(timing.resume_reasons, ["用户继续执行"])
        page.current_worker = None
        self.assertTrue(page.shutdown())
        page.close()

    def test_workflow_start_keeps_previous_log_and_adds_run_separator(self):
        file_path = Path(self.temp_dir.name) / "log-history.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["车辆标识", "已协助补缴"])
        sheet.append(["粤A12345_黄色", ""])
        workbook.save(file_path)
        workbook.close()

        page = WorkflowPage()
        page.file_path = str(file_path)
        page.df = pd.DataFrame(
            {"车辆标识": ["粤A12345_黄色"], "已协助补缴": [""]}
        )
        page.browser_check_state = "ready"
        page.log_text.setPlainText("[19:00:00] 上一次执行发生浏览器异常")
        with patch.object(page, "_reload_preview", return_value=True), patch.object(
            page,
            "_start_transport_worker",
        ), patch(
            "integrated_client.ui.workflow_page.get_builtin_chromium_path",
            return_value=r"C:\browser\chrome.exe",
        ):
            page.start_pipeline()

        log_text = page.log_text.toPlainText()
        self.assertIn("上一次执行发生浏览器异常", log_text)
        self.assertIn("第 1 次流水线执行", log_text)
        self.assertIn("一键三步流水线已启动", log_text)
        self.assertLess(
            log_text.index("上一次执行发生浏览器异常"),
            log_text.index("第 1 次流水线执行"),
        )
        page._finish_pipeline(False, "测试结束", outcome="failed")
        page.close()
        page.deleteLater()
        self.app.processEvents()

    def test_workflow_start_browser_check_dialog_actions(self):
        self.assertFalse(
            WorkflowPage._browser_start_result_allows_pipeline(
                QMessageBox.Cancel
            )
        )
        self.assertTrue(
            WorkflowPage._browser_start_result_allows_pipeline(
                QMessageBox.Ignore
            )
        )
        self.assertTrue(
            WorkflowPage._browser_start_result_allows_pipeline(QMessageBox.Ok)
        )
        self.assertEqual(
            WorkflowPage.BROWSER_CHECK_ANIMATION_FRAMES,
            (
                "正在检测内置浏览器",
                "正在检测内置浏览器 ·",
                "正在检测内置浏览器 ··",
                "正在检测内置浏览器 ···",
            ),
        )
        self.assertIn("QProgressBar#BrowserCheckProgress", APP_STYLESHEET)

    def test_disabled_widgets_use_forbidden_cursor(self):
        previous_stylesheet = self.app.styleSheet()
        self.addCleanup(self.app.setStyleSheet, previous_stylesheet)
        self.app.setStyleSheet(APP_STYLESHEET)
        window = MainWindow(self.db, self.admin)
        self.assertFalse(window.statistics_page.station_combo.isEnabled())
        self.assertEqual(
            window.statistics_page.station_combo.cursor().shape(),
            Qt.ForbiddenCursor,
        )
        self.assertFalse(window.workflow_page.pause_btn.isEnabled())
        self.assertEqual(
            window.workflow_page.pause_btn.cursor().shape(),
            Qt.ForbiddenCursor,
        )
        window.show()
        self.app.processEvents()

        disabled_button = QPushButton("禁用按钮")
        disabled_button.setEnabled(False)
        disabled_button.show()
        disabled_combo = QComboBox()
        disabled_combo.addItem("禁用内容")
        disabled_combo.setEnabled(False)
        disabled_combo.show()
        disabled_label = QLabel("禁用内容")
        disabled_label.setEnabled(False)
        disabled_label.show()
        self.app.processEvents()
        self.assertEqual(
            disabled_button.palette()
            .color(QPalette.Disabled, QPalette.Button)
            .name(),
            "#f3f5f8",
        )
        self.assertEqual(
            disabled_button.palette()
            .color(QPalette.Disabled, QPalette.ButtonText)
            .name(),
            "#aab3bf",
        )
        self.assertEqual(
            disabled_combo.palette()
            .color(QPalette.Disabled, QPalette.Base)
            .name(),
            "#f5f7fa",
        )
        self.assertEqual(
            disabled_combo.palette()
            .color(QPalette.Disabled, QPalette.Text)
            .name(),
            "#a3adba",
        )
        self.assertEqual(
            disabled_label.palette()
            .color(QPalette.Disabled, QPalette.WindowText)
            .name(),
            "#a3adba",
        )
        disabled_button.close()
        disabled_combo.close()
        disabled_label.close()

        window.workflow_page.pause_btn.setEnabled(True)
        self.assertNotEqual(
            window.workflow_page.pause_btn.cursor().shape(),
            Qt.ForbiddenCursor,
        )
        window.workflow_page.pause_btn.setEnabled(False)
        self.assertEqual(
            window.workflow_page.pause_btn.cursor().shape(),
            Qt.ForbiddenCursor,
        )

        custom_cursor_button = QPushButton("测试自定义指针")
        custom_cursor_button.setCursor(Qt.PointingHandCursor)
        custom_cursor_button.setEnabled(False)
        self.assertEqual(
            custom_cursor_button.cursor().shape(),
            Qt.ForbiddenCursor,
        )
        cursor_filter = self.app._intdemo_disabled_cursor_filter
        cursor_filter.sync_hover_target(custom_cursor_button)
        self.assertIsNotNone(QApplication.overrideCursor())
        self.assertEqual(
            QApplication.overrideCursor().shape(),
            Qt.ForbiddenCursor,
        )
        custom_cursor_button.setEnabled(True)
        cursor_filter.sync_hover_target(custom_cursor_button)
        self.assertIsNone(QApplication.overrideCursor())
        self.assertEqual(
            custom_cursor_button.cursor().shape(),
            Qt.PointingHandCursor,
        )
        custom_cursor_button.close()

        self.assertTrue(window.workflow_page.shutdown())
        window._prepared_to_close = True
        window.close()

    def test_station_distribution_shows_total_and_phone_shares(self):
        self.assertEqual(self.db.ensure_default_station_users(), 5)
        accounts = {account.username: account for account in self.db.list_accounts()}
        luogang = accounts["luogang"]
        taiping = accounts["taiping"]

        def record_timing(account, run_id, active_before_pause, paused, active_after_pause):
            clock = [1000.0]
            service = WorkflowTimingService(
                self.db,
                account.id,
                clock=lambda: clock[0],
            )
            dataframe = pd.DataFrame(
                {
                    "车辆标识": [f"粤A{account.id:05d}_黄色"],
                    "已协助补缴": [""],
                    "原因": [""],
                }
            )
            service.start_run(
                Path(self.temp_dir.name) / f"{run_id}.xlsx",
                dataframe,
                run_id=run_id,
            )
            service.start_step(1)
            clock[0] += active_before_pause
            if paused:
                service.pause("等待人工")
                clock[0] += paused
                service.resume("继续执行")
            clock[0] += active_after_pause
            service.finish_run("succeeded")

        record_timing(
            luogang,
            "station-share-luogang",
            3600,
            600,
            1800,
        )
        record_timing(
            taiping,
            "station-share-taiping",
            1800,
            0,
            0,
        )
        self.db.record_activity_batch(
            luogang.id,
            {WORKFLOW_TOTAL_METRIC: 30, WORKFLOW_HAS_PHONE_METRIC: 12},
            "unified_workflow",
            task_id="station-share-luogang",
        )
        self.db.record_activity_batch(
            taiping.id,
            {WORKFLOW_TOTAL_METRIC: 10, WORKFLOW_HAS_PHONE_METRIC: 8},
            "unified_workflow",
            task_id="station-share-taiping",
        )

        window = MainWindow(self.db, self.admin)
        page = window.statistics_page
        page.category_combo.setCurrentIndex(
            page.category_combo.findData("completion")
        )
        station_index = window.statistics_page.station_combo.findData(luogang.id)
        window.statistics_page.station_combo.setCurrentIndex(station_index)
        category_index = window.statistics_page.category_combo.findData(
            "station_distribution"
        )
        window.statistics_page.category_combo.setCurrentIndex(category_index)
        self.app.processEvents()

        self.assertIsInstance(page.station_distribution_chart, StationDistributionChart)
        self.assertFalse(page.station_combo.isEnabled())
        self.assertIsNone(page.station_combo.currentData())
        self.assertIn("全部站点", page.station_combo.toolTip())
        self.assertTrue(page.distribution_chart.isHidden())
        self.assertTrue(page.violation_chart.isHidden())
        self.assertFalse(page.station_distribution_chart.isHidden())

        rows = {
            row["username"]: row for row in page.station_distribution_chart._rows
        }
        self.assertEqual(len(rows), 5)
        self.assertEqual(rows["luogang"]["total"], 30)
        self.assertEqual(rows["luogang"]["has_phone"], 12)
        self.assertAlmostEqual(rows["luogang"]["total_share"], 75.0)
        self.assertAlmostEqual(rows["luogang"]["phone_share"], 60.0)
        self.assertAlmostEqual(rows["taiping"]["total_share"], 25.0)
        self.assertAlmostEqual(rows["taiping"]["phone_share"], 40.0)
        self.assertEqual(rows["luogang"]["total_time_ms"], 6_000_000)
        self.assertEqual(rows["luogang"]["active_ms"], 5_400_000)
        self.assertAlmostEqual(rows["luogang"]["total_time_share"], 6000 / 78)
        self.assertAlmostEqual(rows["luogang"]["active_time_share"], 75.0)
        self.assertAlmostEqual(rows["taiping"]["total_time_share"], 1800 / 78)
        self.assertAlmostEqual(rows["taiping"]["active_time_share"], 25.0)
        page.station_distribution_chart.resize(1000, 280)
        page.station_distribution_chart.grab()
        self.app.processEvents()
        self.assertEqual(len(page.station_distribution_chart._slice_hitboxes), 4)
        self.assertEqual(
            {
                item["payload"]["series"]
                for item in page.station_distribution_chart._slice_hitboxes
            },
            {
                "各站总计数占比",
                "各站有电话数占比",
            },
        )
        chart = page.station_distribution_chart
        slice_item = chart._slice_hitboxes[0]
        outer = slice_item["outer"]
        inner = slice_item["inner"]
        slice_local = QPoint(
            int(outer.center().x() + (outer.width() + inner.width()) / 4),
            int(outer.center().y()),
        )
        slice_event = QMouseEvent(
            QEvent.MouseMove,
            slice_local,
            chart.mapToGlobal(slice_local),
            Qt.NoButton,
            Qt.NoButton,
            Qt.NoModifier,
        )
        chart.mouseMoveEvent(slice_event)
        slice_payload = slice_item["payload"]
        self.assertEqual(
            chart._hover_card.details,
            [
                ("数量", f'{slice_payload["value"]} 条'),
                ("占比", chart._format_share(slice_payload["share"])),
            ],
        )
        self.assertEqual(chart._hover_card.width(), 240)

        legend_rect, legend_row = chart._legend_hitboxes[0]
        legend_local = legend_rect.center().toPoint()
        legend_event = QMouseEvent(
            QEvent.MouseMove,
            legend_local,
            chart.mapToGlobal(legend_local),
            Qt.NoButton,
            Qt.NoButton,
            Qt.NoModifier,
        )
        chart.mouseMoveEvent(legend_event)
        self.app.processEvents()
        self.assertEqual(
            chart._hover_card.details,
            [
                ("总计数", f'{legend_row["total"]} 条'),
                ("总数占比", chart._format_share(legend_row["total_share"])),
                ("有电话数", f'{legend_row["has_phone"]} 条'),
                ("有电话占比", chart._format_share(legend_row["phone_share"])),
            ],
        )
        self.assertEqual(chart._hover_card.width(), 340)

        page.category_combo.setCurrentIndex(
            page.category_combo.findData("timing")
        )
        self.app.processEvents()
        self.assertTrue(page.station_combo.isEnabled())
        self.assertEqual(page.station_combo.currentData(), luogang.id)
        self.assertTrue(chart.isHidden())
        self.assertFalse(page.detail_tabs.isTabEnabled(0))
        self.assertIs(page.detail_tabs.currentWidget(), page.data_tab)

        page.station_combo.setCurrentIndex(0)
        self.app.processEvents()
        self.assertIsNone(page.station_combo.currentData())
        self.assertTrue(page.detail_tabs.isTabEnabled(0))
        self.assertIs(page.detail_tabs.currentWidget(), page.chart_tab)
        self.assertFalse(chart.isHidden())
        self.assertEqual(chart._series_mode, "timing")
        chart.resize(1000, 280)
        chart.grab()
        self.app.processEvents()
        self.assertEqual(len(chart._slice_hitboxes), 4)
        self.assertEqual(
            {
                item["payload"]["series"]
                for item in chart._slice_hitboxes
            },
            {"各站总耗时占比", "各站有效耗时占比"},
        )

        duration_item = next(
            item
            for item in chart._slice_hitboxes
            if item["payload"]["series"] == "各站总耗时占比"
        )
        duration_outer = duration_item["outer"]
        duration_inner = duration_item["inner"]
        duration_local = QPoint(
            int(
                duration_outer.center().x()
                + (duration_outer.width() + duration_inner.width()) / 4
            ),
            int(duration_outer.center().y()),
        )
        duration_event = QMouseEvent(
            QEvent.MouseMove,
            duration_local,
            chart.mapToGlobal(duration_local),
            Qt.NoButton,
            Qt.NoButton,
            Qt.NoModifier,
        )
        chart.mouseMoveEvent(duration_event)
        duration_payload = duration_item["payload"]
        self.assertEqual(
            chart._hover_card.details,
            [
                ("小时数", "1.7 小时"),
                ("精确耗时", "01:40:00.000"),
                ("占比", chart._format_share(duration_payload["share"])),
            ],
        )
        self.assertEqual(chart._hover_card.width(), 300)

        legend_rect, legend_row = chart._legend_hitboxes[0]
        legend_local = legend_rect.center().toPoint()
        legend_event = QMouseEvent(
            QEvent.MouseMove,
            legend_local,
            chart.mapToGlobal(legend_local),
            Qt.NoButton,
            Qt.NoButton,
            Qt.NoModifier,
        )
        chart.mouseMoveEvent(legend_event)
        self.app.processEvents()
        self.assertEqual(
            chart._hover_card.details,
            [
                ("总耗时", "1.7 小时"),
                ("精确总耗时", "01:40:00.000"),
                (
                    "总耗时占比",
                    chart._format_share(legend_row["total_time_share"]),
                ),
                ("有效耗时", "1.5 小时"),
                ("精确有效耗时", "01:30:00.000"),
                (
                    "有效耗时占比",
                    chart._format_share(legend_row["active_time_share"]),
                ),
            ],
        )
        self.assertEqual(chart._hover_card.width(), 520)
        for index in range(chart._hover_card._details_layout.count()):
            detail_label = chart._hover_card._details_layout.itemAt(index).widget()
            self.assertGreaterEqual(detail_label.width(), detail_label.sizeHint().width())

        page.station_combo.setCurrentIndex(
            page.station_combo.findData(luogang.id)
        )
        self.app.processEvents()
        self.assertTrue(chart.isHidden())
        self.assertFalse(page.detail_tabs.isTabEnabled(0))
        self.assertIs(page.detail_tabs.currentWidget(), page.data_tab)

        page.category_combo.setCurrentIndex(
            page.category_combo.findData("station_distribution")
        )
        self.app.processEvents()
        self.assertEqual(page.summary_table.rowCount(), 5)
        self.assertEqual(
            [
                page.summary_table.horizontalHeaderItem(column).text()
                for column in range(page.summary_table.columnCount())
            ],
            ["站点", "总计数", "总数占比", "有电话数", "有电话数占比"],
        )
        luogang_row = next(
            row
            for row in range(page.summary_table.rowCount())
            if page.summary_table.item(row, 0).text() == "萝岗中心站"
        )
        self.assertEqual(page.summary_table.item(luogang_row, 1).text(), "30")
        self.assertEqual(page.summary_table.item(luogang_row, 2).text(), "75.0%")
        self.assertEqual(page.summary_table.item(luogang_row, 3).text(), "12")
        self.assertEqual(page.summary_table.item(luogang_row, 4).text(), "60.0%")
        self.assertEqual(
            page.data_description.text(),
            "全部站点 · 各站总计数占比与有电话数占比",
        )

        page.category_combo.setCurrentIndex(
            page.category_combo.findData("completion")
        )
        self.app.processEvents()
        self.assertTrue(page.station_combo.isEnabled())
        self.assertEqual(page.station_combo.currentData(), luogang.id)
        self.assertTrue(window.workflow_page.shutdown())
        window._prepared_to_close = True
        window.close()

    def test_account_page_summary_and_selected_account_state(self):
        self.assertEqual(self.db.ensure_default_station_users(), 5)
        window = MainWindow(self.db, self.admin)
        page = window.account_page

        self.assertEqual(page.summary_values["total"].text(), "6")
        self.assertEqual(page.summary_values["admin"].text(), "1")
        self.assertEqual(page.summary_values["user"].text(), "5")
        self.assertEqual(page.summary_values["active"].text(), "6")
        self.assertEqual(page.table.rowCount(), 6)
        self.assertFalse(page.table.showGrid())
        self.assertFalse(page.table.verticalHeader().isVisible())
        table_header = page.table.horizontalHeader()
        for column in (0, 1, 4, 5):
            self.assertEqual(
                table_header.sectionResizeMode(column),
                QHeaderView.Stretch,
            )
        for column in (2, 3, 6):
            self.assertEqual(
                table_header.sectionResizeMode(column),
                QHeaderView.ResizeToContents,
            )

        luogang_row = next(
            index
            for index, account in enumerate(page.accounts)
            if account.username == "luogang"
        )
        page.table.selectRow(luogang_row)
        self.app.processEvents()
        self.assertIn("萝岗中心站", page.selection_hint.text())
        self.assertTrue(page.rename_btn.isEnabled())
        self.assertTrue(page.permission_btn.isEnabled())
        self.assertTrue(page.export_btn.isEnabled())
        self.assertTrue(page.import_btn.isEnabled())
        self.assertTrue(page.reset_stats_btn.isEnabled())
        self.assertTrue(page.data_range_selector.isEnabled())
        self.assertTrue(page.reset_btn.isEnabled())
        self.assertTrue(page.toggle_btn.isEnabled())
        self.assertTrue(page.delete_btn.isEnabled())
        self.assertEqual(page.toggle_btn.text(), "停用账号")
        self.assertEqual(page.toggle_btn.objectName(), "DangerButton")

        current_row = next(
            index
            for index, account in enumerate(page.accounts)
            if account.id == self.admin.id
        )
        page.table.selectRow(current_row)
        self.app.processEvents()
        self.assertIn("当前登录账号", page.selection_hint.text())
        self.assertTrue(page.rename_btn.isEnabled())
        self.assertFalse(page.permission_btn.isEnabled())
        self.assertFalse(page.export_btn.isEnabled())
        self.assertFalse(page.import_btn.isEnabled())
        self.assertFalse(page.reset_stats_btn.isEnabled())
        self.assertFalse(page.data_range_selector.isEnabled())
        self.assertFalse(page.reset_btn.isEnabled())
        self.assertFalse(page.toggle_btn.isEnabled())
        self.assertFalse(page.delete_btn.isEnabled())
        self.assertIn("个人中心", page.reset_btn.toolTip())

        self.assertTrue(window.workflow_page.shutdown())
        window._prepared_to_close = True
        window.close()

    def test_admin_can_delete_an_account_and_refresh_dashboard(self):
        self.assertEqual(self.db.ensure_default_station_users(), 5)
        window = MainWindow(self.db, self.admin)
        page = window.account_page
        target = next(
            account for account in page.accounts if account.username == "luogang"
        )
        self.db.record_activity(target.id, WORKFLOW_TOTAL_METRIC, 3, "manual")
        page.table.selectRow(page.accounts.index(target))

        with (
            patch(
                "integrated_client.ui.account_page.QMessageBox.question",
                return_value=QMessageBox.Yes,
            ) as question,
            patch(
                "integrated_client.ui.account_page.QMessageBox.information"
            ) as information,
        ):
            page._delete_account()

        confirmation = question.call_args.args[2]
        self.assertIn(target.name_label, confirmation)
        self.assertIn(target.username, confirmation)
        self.assertIn("无法恢复", confirmation)
        self.assertIn("清除 1 条统计事件", information.call_args.args[2])
        self.assertNotIn(target.id, [account.id for account in page.accounts])
        self.assertEqual(page.summary_values["total"].text(), "5")
        self.assertEqual(
            window.statistics_page.station_combo.findData(target.id),
            -1,
        )
        self.assertEqual(self.db.ensure_default_station_users(), 0)

        self.assertTrue(window.workflow_page.shutdown())
        window._prepared_to_close = True
        window.close()

    def test_admin_can_rename_users_and_current_identity(self):
        self.assertEqual(self.db.ensure_default_station_users(), 5)
        window = MainWindow(self.db, self.admin)
        page = window.account_page
        luogang = next(
            account for account in page.accounts if account.username == "luogang"
        )
        luogang_row = page.accounts.index(luogang)
        page.table.selectRow(luogang_row)

        with (
            patch(
                "integrated_client.ui.account_page.RenameAccountDialog"
            ) as dialog_class,
            patch(
                "integrated_client.ui.account_page.QMessageBox.information"
            ) as information,
        ):
            dialog = dialog_class.return_value
            dialog.Accepted = 1
            dialog.exec_.return_value = 1
            dialog.value.return_value = "萝岗业务站"
            page._rename_account()

            renamed_station = next(
                account
                for account in self.db.list_accounts()
                if account.id == luogang.id
            )
            self.assertEqual(renamed_station.name_label, "萝岗业务站")
            self.assertEqual(renamed_station.username, "luogang")
            station_index = window.statistics_page.station_combo.findData(luogang.id)
            self.assertEqual(
                window.statistics_page.station_combo.itemText(station_index),
                "萝岗业务站",
            )

            current_row = next(
                index
                for index, account in enumerate(page.accounts)
                if account.id == self.admin.id
            )
            page.table.selectRow(current_row)
            dialog.value.return_value = "主管理员"
            page._rename_account()

            self.assertEqual(window.account.name_label, "主管理员")
            self.assertEqual(page.current_account.name_label, "主管理员")
            self.assertEqual(window.sidebar_user.text(), "主管理员")
            self.assertEqual(window.sidebar_user.toolTip(), "主管理员")
            self.assertEqual(window.sidebar_avatar.text(), "主")
            self.assertEqual(window.personal_center_page.name_value.text(), "主管理员")
            self.assertIn("主管理员", window.windowTitle())
            self.assertEqual(information.call_count, 2)

        self.assertTrue(window.workflow_page.shutdown())
        window._prepared_to_close = True
        window.close()

    def test_admin_can_change_account_permissions(self):
        self.assertEqual(self.db.ensure_default_station_users(), 5)
        window = MainWindow(self.db, self.admin)
        page = window.account_page
        luogang = next(
            account for account in page.accounts if account.username == "luogang"
        )
        page.table.selectRow(page.accounts.index(luogang))

        with (
            patch(
                "integrated_client.ui.account_page.AccountPermissionDialog"
            ) as dialog_class,
            patch("integrated_client.ui.account_page.QMessageBox.information"),
        ):
            dialog = dialog_class.return_value
            dialog.Accepted = 1
            dialog.exec_.return_value = 1
            dialog.value.return_value = "admin"
            page._change_permission()

            promoted = next(
                account
                for account in self.db.list_accounts()
                if account.id == luogang.id
            )
            self.assertTrue(promoted.is_admin)
            self.assertEqual(page.summary_values["admin"].text(), "2")
            self.assertEqual(
                window.statistics_page.station_combo.findData(luogang.id),
                -1,
            )

            promoted_row = next(
                index
                for index, account in enumerate(page.accounts)
                if account.id == luogang.id
            )
            page.table.selectRow(promoted_row)
            dialog.value.return_value = "user"
            page._change_permission()

            demoted = next(
                account
                for account in self.db.list_accounts()
                if account.id == luogang.id
            )
            self.assertFalse(demoted.is_admin)
            self.assertEqual(page.summary_values["admin"].text(), "1")
            self.assertGreaterEqual(
                window.statistics_page.station_combo.findData(luogang.id),
                0,
            )

        self.assertTrue(window.workflow_page.shutdown())
        window._prepared_to_close = True
        window.close()

    def test_admin_can_export_and_import_selected_station_data(self):
        self.assertEqual(self.db.ensure_default_station_users(), 5)
        accounts = {account.username: account for account in self.db.list_accounts()}
        luogang = accounts["luogang"]
        taiping = accounts["taiping"]
        self.db.record_activity_batch(
            luogang.id,
            {
                WORKFLOW_TOTAL_METRIC: 5,
                WORKFLOW_HAS_PHONE_METRIC: 3,
            },
            "unified_workflow",
            task_id="ui-station-transfer",
        )
        window = MainWindow(self.db, self.admin)
        page = window.account_page
        luogang_row = next(
            index
            for index, account in enumerate(page.accounts)
            if account.id == luogang.id
        )
        page.table.selectRow(luogang_row)
        today = QDate.currentDate()
        page.data_range_selector.all_dates_check.setChecked(False)
        page.data_range_selector.start_edit.setDate(today)
        page.data_range_selector.end_edit.setDate(today)
        export_base = Path(self.temp_dir.name) / "ui-station-transfer"

        with (
            patch(
                "integrated_client.ui.account_page.QFileDialog.getSaveFileName",
                return_value=(str(export_base), "站点数据文件 (*.json)"),
            ),
            patch("integrated_client.ui.account_page.QMessageBox.information"),
        ):
            page._export_station_data()
        export_path = export_base.with_suffix(".json")
        self.assertTrue(export_path.exists())
        exported_payload = json.loads(export_path.read_text(encoding="utf-8"))
        self.assertEqual(
            exported_payload["date_range"],
            {
                "start": today.toString("yyyy-MM-dd"),
                "end": today.toString("yyyy-MM-dd"),
            },
        )

        taiping_row = next(
            index
            for index, account in enumerate(page.accounts)
            if account.id == taiping.id
        )
        page.table.selectRow(taiping_row)
        with (
            patch(
                "integrated_client.ui.account_page.QFileDialog.getOpenFileName",
                return_value=(str(export_path), "站点数据文件 (*.json)"),
            ),
            patch(
                "integrated_client.ui.account_page.QMessageBox.question",
                return_value=QMessageBox.Yes,
            ),
            patch("integrated_client.ui.account_page.QMessageBox.information"),
        ):
            page._import_station_data()

        target_totals = {
            row["metric_key"]: row["total"]
            for row in self.db.get_user_totals(taiping.id)
        }
        self.assertEqual(target_totals[WORKFLOW_TOTAL_METRIC], 5)
        self.assertEqual(target_totals[WORKFLOW_HAS_PHONE_METRIC], 3)
        station_rows = {
            row["username"]: row
            for row in window.statistics_page.station_distribution_chart._rows
        }
        self.assertEqual(station_rows["taiping"]["total"], 5)
        self.assertEqual(station_rows["taiping"]["has_phone"], 3)

        self.assertTrue(page.reset_stats_btn.isEnabled())
        with (
            patch(
                "integrated_client.ui.account_page.QMessageBox.question",
                return_value=QMessageBox.Yes,
            ) as question,
            patch(
                "integrated_client.ui.account_page.QMessageBox.information"
            ) as information,
        ):
            page._reset_station_statistics()

        confirmation_text = question.call_args.args[2]
        self.assertIn(taiping.name_label, confirmation_text)
        self.assertIn(taiping.username, confirmation_text)
        self.assertIn("永久删除", confirmation_text)
        self.assertIn("已清除 2 条统计事件", information.call_args.args[2])
        reset_totals = {
            row["metric_key"]: row["total"]
            for row in self.db.get_user_totals(taiping.id)
        }
        self.assertEqual(reset_totals[WORKFLOW_TOTAL_METRIC], 0)
        self.assertEqual(reset_totals[WORKFLOW_HAS_PHONE_METRIC], 0)
        self.assertEqual(
            self.db.authenticate("taiping", DEFAULT_STATION_PASSWORD).id,
            taiping.id,
        )
        station_rows = {
            row["username"]: row
            for row in window.statistics_page.station_distribution_chart._rows
        }
        self.assertEqual(station_rows["taiping"]["total"], 0)

        admin_row = next(
            index
            for index, account in enumerate(page.accounts)
            if account.id == self.admin.id
        )
        page.table.selectRow(admin_row)
        self.assertFalse(page.reset_stats_btn.isEnabled())

        self.assertTrue(window.workflow_page.shutdown())
        window._prepared_to_close = True
        window.close()

    def test_live_preview_model_updates_a_completed_row(self):
        model = DataFrameTableModel(pd.DataFrame([{"车辆标识": "粤A12345"}]))
        model.update_row(0, {LEGAL_COL_NAME: "张三", "查询状态": "完成"})

        legal_column = model.dataframe.columns.get_loc(LEGAL_COL_NAME)
        status_column = model.dataframe.columns.get_loc("查询状态")
        self.assertEqual(model.data(model.index(0, legal_column), Qt.DisplayRole), "张三")
        self.assertEqual(model.data(model.index(0, status_column), Qt.DisplayRole), "完成")

    def test_violation_chart_bar_hover_exposes_phone_breakdown(self):
        full_reason = "恶意大车小标出入口车型不符其它复杂描述内容"
        chart = ViolationReasonChart()
        chart.resize(900, 280)
        chart.set_rows(
            [{"reason": full_reason, "total": 5, "has_phone": 2, "other": 3}],
            "测试站点",
        )
        chart.mode_button.click()
        self.assertEqual(chart._bar_mode, "split")
        chart.show()
        pixmap = chart.grab()
        self.app.processEvents()
        bar_rect, bar_row = chart._bar_hitboxes[0]
        image = pixmap.toImage()
        bar_y = int(bar_rect.center().y())
        phone_x = int(bar_rect.left() + bar_rect.width() * 0.2)
        other_x = int(bar_rect.left() + bar_rect.width() * 0.8)
        self.assertEqual(image.pixelColor(phone_x, bar_y).name(), chart.PHONE_COLOR.name())
        self.assertEqual(image.pixelColor(other_x, bar_y).name(), chart.OTHER_COLOR.name())
        self.assertEqual(bar_rect.height(), chart.BAR_HEIGHT)
        self.assertEqual(chart._format_total_label(5), "5 条")
        self.assertNotIn("%", chart._format_total_label(5))

        slice_item = chart._slice_hitboxes[0]
        outer = slice_item["outer"]
        inner = slice_item["inner"]
        donut_local = QPoint(
            int(outer.center().x() + (outer.width() + inner.width()) / 4),
            int(outer.center().y()),
        )
        donut_event = QMouseEvent(
            QEvent.MouseMove,
            donut_local,
            chart.mapToGlobal(donut_local),
            Qt.NoButton,
            Qt.NoButton,
            Qt.NoModifier,
        )
        chart.mouseMoveEvent(donut_event)
        self.assertEqual(chart._hovered_slice, 0)
        self.assertEqual(
            chart._hover_card.title_text,
            full_reason,
        )
        self.assertEqual(
            chart._hover_card.details,
            [
                ("总计", "5 条"),
                ("占比", "100.0%"),
                ("有电话", "2 条"),
                ("其他数据", "3 条"),
            ],
        )
        self.assertTrue(chart._hover_card.isVisible())
        self.assertTrue(chart._hover_card.isWindow())
        self.assertTrue(
            chart._hover_card.windowFlags() & Qt.NoDropShadowWindowHint
        )
        self.assertTrue(
            chart._hover_card.testAttribute(Qt.WA_TranslucentBackground)
        )
        self.assertTrue(chart._hover_card.testAttribute(Qt.WA_StyledBackground))
        hover_card_image = chart._hover_card.grab().toImage()
        self.assertEqual(hover_card_image.pixelColor(0, 0).alpha(), 0)
        self.assertGreater(
            hover_card_image.pixelColor(hover_card_image.rect().center()).alpha(),
            0,
        )
        self.assertGreaterEqual(
            chart._hover_card.height(),
            50 + chart._hover_card._detail_row_count * 22,
        )
        donut_card_height = chart._hover_card.height()
        donut_card_width = chart._hover_card.width()

        bar_local = QPoint(int(bar_rect.center().x()), int(bar_rect.center().y()))
        bar_event = QMouseEvent(
            QEvent.MouseMove,
            bar_local,
            chart.mapToGlobal(bar_local),
            Qt.NoButton,
            Qt.NoButton,
            Qt.NoModifier,
        )
        chart.mouseMoveEvent(bar_event)
        self.assertEqual(bar_row["reason"], full_reason)
        self.assertEqual(chart._hover_card.title_text, "电话数据拆分")
        self.assertEqual(
            chart._hover_card.details,
            [("有电话", "2 条"), ("其他数据", "3 条")],
        )
        self.assertEqual(chart._hover_card.height(), donut_card_height)
        self.assertEqual(chart._hover_card.width(), donut_card_width)
        self.assertIsNone(chart._hovered_slice)
        phase = chart._animation_phase
        chart._advance_animation()
        self.assertNotEqual(chart._animation_phase, phase)
        self.assertTrue(chart._animation_timer.isActive())
        chart.hide()
        self.app.processEvents()
        self.assertFalse(chart._animation_timer.isActive())
        chart.close()

    def test_completion_donut_and_bar_hover_expose_category_details(self):
        chart = WorkflowDistributionChart()
        self.assertEqual(
            [metric_key for metric_key, _, _ in chart.SEGMENTS],
            [
                WORKFLOW_HAS_PHONE_METRIC,
                WORKFLOW_NO_PHONE_METRIC,
                WORKFLOW_INDIVIDUAL_METRIC,
                WORKFLOW_NO_OPERATION_METRIC,
                WORKFLOW_NO_TRANSPORT_METRIC,
            ],
        )
        chart.resize(900, 280)
        chart.set_values(
            {
                WORKFLOW_TOTAL_METRIC: 4,
                WORKFLOW_EMPTY_METRIC: 4,
                WORKFLOW_NO_TRANSPORT_METRIC: 0,
                WORKFLOW_NO_OPERATION_METRIC: 0,
                WORKFLOW_INDIVIDUAL_METRIC: 0,
                WORKFLOW_NO_PHONE_METRIC: 0,
                WORKFLOW_HAS_PHONE_METRIC: 4,
            },
            "测试站点",
        )
        chart.show()
        chart.grab()
        self.app.processEvents()
        self.assertEqual(len(chart._slice_hitboxes), 1)
        self.assertTrue(all(rect.height() == chart.BAR_HEIGHT for rect in chart._bar_rects))
        slice_item = chart._slice_hitboxes[0]
        outer = slice_item["outer"]
        inner = slice_item["inner"]
        donut_local = QPoint(
            int(outer.center().x() + (outer.width() + inner.width()) / 4),
            int(outer.center().y()),
        )
        donut_event = QMouseEvent(
            QEvent.MouseMove,
            donut_local,
            chart.mapToGlobal(donut_local),
            Qt.NoButton,
            Qt.NoButton,
            Qt.NoModifier,
        )
        chart.mouseMoveEvent(donut_event)
        self.assertEqual(chart._hovered_slice, 0)
        self.assertEqual(chart._hover_card.title_text, "有公司名、有电话")
        self.assertEqual(
            chart._hover_card.details,
            [("数量", "4 条"), ("占比", "100.0%")],
        )
        self.assertTrue(chart._hover_card.isVisible())
        self.assertTrue(chart._hover_card.isWindow())

        has_phone_bar = chart._bar_rects[0]
        bar_local = QPoint(
            int(has_phone_bar.center().x()),
            int(has_phone_bar.center().y()),
        )
        bar_event = QMouseEvent(
            QEvent.MouseMove,
            bar_local,
            chart.mapToGlobal(bar_local),
            Qt.NoButton,
            Qt.NoButton,
            Qt.NoModifier,
        )
        chart.mouseMoveEvent(bar_event)
        self.assertIsNone(chart._hovered_slice)
        self.assertEqual(chart._hover_card.title_text, "有公司名、有电话")
        self.assertEqual(
            chart._hover_card.details,
            [("数量", "4 条"), ("占比", "100.0%")],
        )
        self.assertTrue(chart._hover_card.isVisible())
        chart.close()

    def test_violation_chart_can_switch_to_reason_share_bars(self):
        chart = ViolationReasonChart()
        chart.resize(900, 280)
        chart.set_rows(
            [
                {"reason": "原因A", "total": 2, "has_phone": 1, "other": 1},
                {"reason": "原因B", "total": 8, "has_phone": 3, "other": 5},
            ],
            "测试站点",
        )
        chart.set_bar_mode("share")
        chart.show()
        chart._animation_timer.stop()
        chart._animation_phase = 0.0
        image = chart.grab().toImage()
        self.app.processEvents()
        self.assertEqual(chart._bar_mode, "share")
        self.assertTrue(chart.mode_button.isChecked())
        self.assertEqual(chart.mode_button.text(), "切换为电话拆分")
        self.assertEqual(chart._format_bar_label(2, 10), "2 条  ·  20.0%")
        self.assertEqual(len(chart._bar_hitboxes), 2)
        self.assertEqual(
            int(
                chart._bar_hitboxes[1][0].top()
                - chart._bar_hitboxes[0][0].top()
            ),
            chart.BAR_ROW_HEIGHT,
        )
        bar_rect, _ = chart._bar_hitboxes[0]
        self.assertLess(chart.mode_button.geometry().bottom(), int(bar_rect.top()))
        bar_y = int(bar_rect.center().y())
        value_x = int(bar_rect.left() + bar_rect.width() * 0.1)
        empty_x = int(bar_rect.left() + bar_rect.width() * 0.5)
        self.assertEqual(
            image.pixelColor(value_x, bar_y).name(), chart.COLORS[0].name()
        )
        self.assertEqual(image.pixelColor(empty_x, bar_y).name(), "#e8f1ef")
        chart.set_bar_mode("split")
        self.assertEqual(chart._format_bar_label(2, 10), "2 条")
        self.assertFalse(chart.mode_button.isChecked())
        chart.close()

    def test_hover_card_is_not_clipped_by_a_narrow_chart(self):
        chart = ViolationReasonChart()
        chart.resize(300, 280)
        chart.set_rows(
            [{"reason": "恶意U", "total": 12, "has_phone": 3, "other": 9}],
            "测试站点",
        )
        chart.show()
        chart.grab()
        self.app.processEvents()
        slice_item = chart._slice_hitboxes[0]
        outer = slice_item["outer"]
        inner = slice_item["inner"]
        donut_local = QPoint(
            int(outer.center().x() + (outer.width() + inner.width()) / 4),
            int(outer.center().y()),
        )
        donut_event = QMouseEvent(
            QEvent.MouseMove,
            donut_local,
            chart.mapToGlobal(donut_local),
            Qt.NoButton,
            Qt.NoButton,
            Qt.NoModifier,
        )
        chart.mouseMoveEvent(donut_event)
        card = chart._hover_card
        chart_right = chart.mapToGlobal(QPoint(chart.width(), 0)).x()
        self.assertTrue(card.isWindow())
        self.assertGreater(card.geometry().right(), chart_right)
        self.assertGreaterEqual(card.height(), 50 + card._detail_row_count * 22)
        screen = QApplication.screenAt(card.geometry().center())
        self.assertIsNotNone(screen)
        self.assertTrue(screen.availableGeometry().contains(card.geometry()))
        chart.close()

    def test_workflow_saves_aiqicha_fields_to_original_workbook(self):
        file_path = Path(self.temp_dir.name) / "workflow.xlsx"
        dataframe = pd.DataFrame(
            [{"车辆标识": "粤A12345", "已协助补缴": "是", "车辆所有人/企业": "示例企业"}]
        )
        dataframe.to_excel(file_path, index=False, engine="openpyxl")

        page = WorkflowPage()
        page.file_path = str(file_path)
        page.df = dataframe.copy()
        for column in TARGET_COLUMNS:
            page.df[column] = ""
        page.df.at[0, LEGAL_COL_NAME] = "张三"
        page.df.at[0, ADDR_COL_NAME] = "测试路1号"
        page.df.at[0, PHONE_COL_NAME] = "13800000000"

        self.assertTrue(page._save_aiqicha_results())
        saved = pd.read_excel(file_path, dtype=str, engine="openpyxl").fillna("")
        self.assertEqual(saved.at[0, LEGAL_COL_NAME], "张三")
        self.assertEqual(saved.at[0, ADDR_COL_NAME], "测试路1号")
        self.assertEqual(saved.at[0, PHONE_COL_NAME], "13800000000")
        page.shutdown()

    def test_workflow_counts_only_once_with_mutually_exclusive_categories(self):
        dataframe = pd.DataFrame(
            [
                {"运输证号_纯数字": "", "车辆所有人/企业": "", PHONE_COL_NAME: "13800000000", "原因": "超限|证件异常"},
                {"运输证号_纯数字": "440100001", "车辆所有人/企业": "已补缴", PHONE_COL_NAME: "", "原因": "恶意U/J形行驶|车型异常"},
                {"运输证号_纯数字": "", "车辆所有人/企业": "无运输证号", PHONE_COL_NAME: "", "原因": "证件异常/证件异常"},
                {"运输证号_纯数字": "", "车辆所有人/企业": "示例运输公司", PHONE_COL_NAME: "020-88886666", "原因": "车型异常"},
                {"运输证号_纯数字": "440100002", "车辆所有人/企业": "无营运信息", PHONE_COL_NAME: "", "原因": "营运异常"},
                {"运输证号_纯数字": "440100003", "车辆所有人/企业": "张三个体工商户", PHONE_COL_NAME: "", "原因": "经营类型"},
                {"运输证号_纯数字": "440100004", "车辆所有人/企业": "甲运输公司", PHONE_COL_NAME: "未公示", "原因": "超限/电话异常"},
                {"运输证号_纯数字": "440100005", "车辆所有人/企业": "乙运输公司", PHONE_COL_NAME: "13800000000", "原因": "超限/证件异常"},
            ]
        )
        calls = []

        def recorder(counts, source, details, task_id):
            calls.append((counts, source, details, task_id))
            return True

        page = WorkflowPage(recorder)
        page.file_path = str(Path(self.temp_dir.name) / "classified.xlsx")
        page.df = dataframe
        page._task_id = "classification-task"
        page._record_workflow_stats()
        page._record_workflow_stats()

        self.assertEqual(len(calls), 1)
        counts = calls[0][0]
        self.assertEqual(counts[WORKFLOW_TOTAL_METRIC], 8)
        self.assertEqual(counts[WORKFLOW_EMPTY_METRIC], 2)
        self.assertEqual(counts[WORKFLOW_NO_TRANSPORT_METRIC], 2)
        self.assertEqual(counts[WORKFLOW_NO_OPERATION_METRIC], 1)
        self.assertEqual(counts[WORKFLOW_INDIVIDUAL_METRIC], 1)
        self.assertEqual(counts[WORKFLOW_NO_PHONE_METRIC], 1)
        self.assertEqual(counts[WORKFLOW_HAS_PHONE_METRIC], 1)
        self.assertEqual(
            counts[WORKFLOW_TOTAL_METRIC],
            counts[WORKFLOW_EMPTY_METRIC]
            + counts[WORKFLOW_NO_TRANSPORT_METRIC]
            + counts[WORKFLOW_NO_OPERATION_METRIC]
            + counts[WORKFLOW_INDIVIDUAL_METRIC]
            + counts[WORKFLOW_NO_PHONE_METRIC]
            + counts[WORKFLOW_HAS_PHONE_METRIC],
        )
        violations = calls[0][2]["violation_counts"]
        self.assertEqual(
            violations["超限"], {"total": 1, "has_phone": 0, "other": 1}
        )
        self.assertEqual(
            violations["证件异常"], {"total": 1, "has_phone": 0, "other": 1}
        )
        self.assertEqual(
            violations["车型异常"], {"total": 2, "has_phone": 0, "other": 2}
        )
        self.assertEqual(
            violations["恶意U/J形行驶"],
            {"total": 1, "has_phone": 0, "other": 1},
        )
        self.assertIn("证件异常/证件异常", violations)
        self.assertIn("超限/电话异常", violations)
        self.assertIn("超限/证件异常", violations)
        self.assertNotIn("恶意U", violations)
        self.assertNotIn("J形行驶", violations)
        page.shutdown()

    def test_dashboard_station_legend_lines_do_not_overlap(self):
        chart = StationShareChart()
        chart.resize(597, 246)
        chart.set_rows(
            [
                {
                    "station": name,
                    "total": total,
                    "total_ms": total_ms,
                }
                for name, total, total_ms in (
                    ("萝岗中心站", 12, 720_000),
                    ("太平中心站", 4, 0),
                    ("道滘中心站", 4, 0),
                    ("宝安中心站", 4, 0),
                    ("南头中心站", 4, 0),
                )
            ]
        )
        chart.show()
        self.app.processEvents()
        chart.grab()

        legend_rects = chart._legend_text_rects
        self.assertEqual(len(legend_rects), 5)
        for station_rect, detail_rect in legend_rects:
            self.assertLessEqual(station_rect.bottom(), detail_rect.top())
        for current, following in zip(legend_rects, legend_rects[1:]):
            self.assertLessEqual(current[1].bottom(), following[0].top())

        chart.close()
        chart.deleteLater()

    def test_dashboard_summarizes_daily_phone_violation_and_station_data(self):
        self.db.ensure_default_station_users()
        accounts = {account.username: account for account in self.db.list_accounts()}
        station = accounts["luogang"]
        today = date.today()
        yesterday = today - timedelta(days=1)

        self.db.record_activity_batch(
            station.id,
            {
                WORKFLOW_TOTAL_METRIC: 10,
                WORKFLOW_HAS_PHONE_METRIC: 7,
            },
            source="unified_workflow",
            details={
                "violation_counts": {
                    "证件异常": {"total": 6, "has_phone": 4, "other": 2},
                    "超限": {"total": 4, "has_phone": 3, "other": 1},
                }
            },
            task_id="dashboard-today",
        )
        self.db.record_activity_batch(
            station.id,
            {
                WORKFLOW_TOTAL_METRIC: 5,
                WORKFLOW_HAS_PHONE_METRIC: 4,
            },
            source="unified_workflow",
            details={
                "violation_counts": {
                    "证件异常": {"total": 3, "has_phone": 2, "other": 1},
                    "车型异常": {"total": 2, "has_phone": 2, "other": 0},
                }
            },
            task_id="dashboard-yesterday",
        )
        with self.db._connect() as conn:
            conn.execute(
                "UPDATE activity_events SET created_at=? WHERE task_id=?",
                (f"{today.isoformat()}T09:00:00+08:00", "dashboard-today"),
            )
            conn.execute(
                "UPDATE activity_events SET created_at=? WHERE task_id=?",
                (f"{yesterday.isoformat()}T09:00:00+08:00", "dashboard-yesterday"),
            )

        daily_rows = self.db.get_daily_metric_totals(
            WORKFLOW_TOTAL_METRIC,
            users_only=True,
            start_date=yesterday,
            end_date=today,
        )
        self.assertEqual(
            daily_rows,
            [
                {"date": yesterday.isoformat(), "total": 5},
                {"date": today.isoformat(), "total": 10},
            ],
        )

        page = DashboardPage(self.db, self.admin)
        page.resize(1220, 780)
        page.show()
        self.app.processEvents()

        self.assertEqual(page.metric_cards["total"].value_label.text(), "15")
        self.assertEqual(page.metric_cards["today"].value_label.text(), "10")
        self.assertIn("+100.0%", page.metric_cards["today"].detail_label.text())
        self.assertEqual(page.metric_cards["phone"].value_label.text(), "11")
        self.assertEqual(page.metric_cards["no_phone"].value_label.text(), "4")
        expected_icon_kinds = {
            "total": "search",
            "today": "calendar",
            "phone": "phone",
            "no_phone": "phone-off",
            "time": "clock",
        }
        for key, icon_kind in expected_icon_kinds.items():
            badge = page.metric_cards[key].icon_badge
            self.assertIsInstance(badge, DashboardMetricIcon)
            self.assertNotIsInstance(badge, QLabel)
            self.assertEqual(badge.icon_kind, icon_kind)
            self.assertFalse(badge.grab().isNull())
        self.assertEqual(page.phone_chart.phone_count, 11)
        self.assertEqual(page.phone_chart.no_phone_count, 4)
        self.assertFalse(page.phone_chart.grab().isNull())
        center_gap = (
            page.phone_chart._center_value_rect.top()
            - page.phone_chart._center_label_rect.bottom()
        )
        self.assertGreaterEqual(center_gap, 5.5)
        self.assertEqual(page.trend_chart._rows[-2]["total"], 5)
        self.assertEqual(page.trend_chart._rows[-1]["total"], 10)
        self.assertEqual(page.violation_table.item(0, 1).text(), "证件异常")
        self.assertEqual(page.violation_table.item(0, 2).text(), "9")
        self.assertEqual(page.violation_table.item(0, 3).text(), "60.0%")

        station_row = next(
            row
            for row in range(page.station_table.rowCount())
            if page.station_table.item(row, 0).text() == station.name_label
        )
        self.assertEqual(page.station_table.item(station_row, 1).text(), "15")
        self.assertEqual(page.station_table.item(station_row, 2).text(), "11")
        self.assertEqual(page.station_table.item(station_row, 3).text(), "4")
        self.assertFalse(page.grab().isNull())
        page.close()
        page.deleteLater()

    def test_dashboard_station_filter_and_hover_details(self):
        self.db.ensure_default_station_users()
        accounts = {account.username: account for account in self.db.list_accounts()}
        luogang = accounts["luogang"]
        taiping = accounts["taiping"]
        today = date.today()
        yesterday = today - timedelta(days=1)

        batches = (
            (
                luogang,
                "dashboard-hover-luogang-today",
                10,
                7,
                today,
            ),
            (
                luogang,
                "dashboard-hover-luogang-yesterday",
                5,
                4,
                yesterday,
            ),
            (
                taiping,
                "dashboard-hover-taiping-today",
                8,
                5,
                today,
            ),
        )
        for station, task_id, total, phone, activity_date in batches:
            self.db.record_activity_batch(
                station.id,
                {
                    WORKFLOW_TOTAL_METRIC: total,
                    WORKFLOW_HAS_PHONE_METRIC: phone,
                },
                source="unified_workflow",
                task_id=task_id,
            )
            with self.db._connect() as conn:
                conn.execute(
                    "UPDATE activity_events SET created_at=? WHERE task_id=?",
                    (
                        f"{activity_date.isoformat()}T09:00:00+08:00",
                        task_id,
                    ),
                )

        def record_timing(station, run_id, seconds):
            clock = [1000.0]
            service = WorkflowTimingService(
                self.db,
                station.id,
                clock=lambda: clock[0],
            )
            service.start_run(
                Path(self.temp_dir.name) / f"{run_id}.xlsx",
                pd.DataFrame({"站点": [station.username]}),
                run_id=run_id,
            )
            service.start_step(1)
            clock[0] += seconds
            service.finish_run("succeeded")

        record_timing(luogang, "dashboard-hover-luogang-today", 3600)
        record_timing(taiping, "dashboard-hover-taiping-today", 1800)

        page = DashboardPage(self.db, self.admin)
        page.resize(1220, 780)
        page.show()
        self.app.processEvents()

        self.assertIsNone(page.station_combo.currentData())
        self.assertEqual(page.metric_cards["total"].value_label.text(), "23")
        self.assertEqual(page.metric_cards["today"].value_label.text(), "18")
        self.assertEqual(page.metric_cards["phone"].value_label.text(), "16")
        self.assertEqual(page.metric_cards["no_phone"].value_label.text(), "7")
        self.assertEqual(page.trend_chart._rows[-1]["has_phone"], 12)
        self.assertEqual(page.trend_chart._rows[-1]["no_phone"], 6)

        page.station_combo.setCurrentIndex(
            page.station_combo.findData(luogang.id)
        )
        self.app.processEvents()
        self.assertEqual(page.metric_cards["total"].value_label.text(), "15")
        self.assertEqual(page.metric_cards["today"].value_label.text(), "10")
        self.assertEqual(page.metric_cards["phone"].value_label.text(), "11")
        self.assertEqual(page.metric_cards["no_phone"].value_label.text(), "4")
        self.assertEqual(page.trend_chart._rows[-1]["has_phone"], 7)
        self.assertEqual(page.trend_chart._rows[-1]["no_phone"], 3)

        page.phone_chart.grab()
        phone_rect = page.phone_chart._donut_rect
        phone_local = QPoint(
            int(phone_rect.center().x() + phone_rect.width() * 0.38),
            int(phone_rect.center().y()),
        )
        phone_event = QMouseEvent(
            QEvent.MouseMove,
            phone_local,
            page.phone_chart.mapToGlobal(phone_local),
            Qt.NoButton,
            Qt.NoButton,
            Qt.NoModifier,
        )
        page.phone_chart.mouseMoveEvent(phone_event)
        self.assertIsInstance(page.phone_chart, AnimatedDonutChart)
        self.assertTrue(page.phone_chart._animation_timer.isActive())
        self.assertEqual(page.phone_chart._hover_card.title_text, "有电话")
        self.assertEqual(
            page.phone_chart._hover_card.details,
            [
                ("数量", "11 条"),
                ("占比", "73.3%"),
                ("总查询量", "15 条"),
            ],
        )

        page.trend_chart.grab()
        trend_point = page.trend_chart._points[-1].toPoint()
        trend_event = QMouseEvent(
            QEvent.MouseMove,
            trend_point,
            page.trend_chart.mapToGlobal(trend_point),
            Qt.NoButton,
            Qt.NoButton,
            Qt.NoModifier,
        )
        page.trend_chart.mouseMoveEvent(trend_event)
        self.assertEqual(
            page.trend_chart._hover_card.details,
            [
                ("查询总量", "10 条"),
                ("有电话", "7 条"),
                ("无电话", "3 条"),
            ],
        )

        page.station_view_toggle.click()
        self.app.processEvents()
        self.assertIs(
            page.station_view_stack.currentWidget(),
            page.station_share_chart,
        )
        self.assertEqual(page.station_view_toggle.text(), "查看数据表")
        page.station_share_chart.grab()
        self.assertIsInstance(page.station_share_chart, AnimatedDonutChart)
        self.assertTrue(page.station_share_chart._animation_timer.isActive())
        self.assertEqual(len(page.station_share_chart._donuts), 2)

        total_donut = page.station_share_chart._donuts[0]["rect"]
        total_local = QPoint(
            int(total_donut.center().x() + total_donut.width() * 0.38),
            int(total_donut.center().y()),
        )
        total_event = QMouseEvent(
            QEvent.MouseMove,
            total_local,
            page.station_share_chart.mapToGlobal(total_local),
            Qt.NoButton,
            Qt.NoButton,
            Qt.NoModifier,
        )
        page.station_share_chart.mouseMoveEvent(total_event)
        self.assertEqual(
            page.station_share_chart._hover_card.title_text,
            f"{luogang.name_label} · 总数量占比",
        )
        self.assertIn(
            ("占比", "65.2%"),
            page.station_share_chart._hover_card.details,
        )

        time_donut = page.station_share_chart._donuts[1]["rect"]
        time_local = QPoint(
            int(time_donut.center().x() + time_donut.width() * 0.38),
            int(time_donut.center().y()),
        )
        time_event = QMouseEvent(
            QEvent.MouseMove,
            time_local,
            page.station_share_chart.mapToGlobal(time_local),
            Qt.NoButton,
            Qt.NoButton,
            Qt.NoModifier,
        )
        page.station_share_chart.mouseMoveEvent(time_event)
        self.assertEqual(
            page.station_share_chart._hover_card.title_text,
            f"{luogang.name_label} · 总用时占比",
        )
        self.assertIn(
            ("小时数", "1.0 小时"),
            page.station_share_chart._hover_card.details,
        )
        page.close()
        page.deleteLater()

    def test_failed_workflow_does_not_record_counts(self):
        file_path = Path(self.temp_dir.name) / "failed.xlsx"
        dataframe = pd.DataFrame(
            [{"车辆标识": "粤A12345", "已协助补缴": "是", "运输证号_纯数字": ""}]
        )
        dataframe.to_excel(file_path, index=False, engine="openpyxl")
        calls = []

        class FinishedWorker:
            def isRunning(self):
                return False

            def deleteLater(self):
                pass

        page = WorkflowPage(lambda *args: calls.append(args) or True)
        page.file_path = str(file_path)
        page.df = dataframe
        page.pipeline_running = True
        page.current_step = 3
        worker = FinishedWorker()
        page.current_worker = worker
        page._save_aiqicha_results = lambda: True
        page._aiqicha_finished(worker, False)

        self.assertEqual(calls, [])
        page.shutdown()

    def test_workflow_automatically_advances_between_three_steps(self):
        file_path = Path(self.temp_dir.name) / "sequence.xlsx"
        pd.DataFrame(
            [{"车辆标识": "粤A12345", "已协助补缴": "是", "车辆所有人/企业": "示例企业"}]
        ).to_excel(file_path, index=False, engine="openpyxl")

        class FinishedWorker:
            def isRunning(self):
                return False

            def deleteLater(self):
                pass

        page = WorkflowPage()
        page.file_path = str(file_path)
        page.pipeline_running = True
        page.current_step = 1
        page._start_backfill_worker = lambda: setattr(page, "_step_two_started", True)
        first_worker = FinishedWorker()
        page.current_worker = first_worker
        page._transport_finished(first_worker, "完成")
        self.app.processEvents()
        self.assertEqual(page.current_step, 2)
        self.assertTrue(page._step_two_started)

        page._start_aiqicha_worker = lambda: setattr(page, "_step_three_started", True)
        second_worker = FinishedWorker()
        page.current_worker = second_worker
        page._backfill_finished(second_worker, "结束")
        self.app.processEvents()
        self.assertEqual(page.current_step, 3)
        self.assertTrue(page._step_three_started)
        page.shutdown()


if __name__ == "__main__":
    unittest.main()
