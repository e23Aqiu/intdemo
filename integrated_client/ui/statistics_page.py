import math
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PyQt5.QtCore import QRect, QRectF, Qt, QTimer
from PyQt5.QtGui import (
    QColor,
    QFont,
    QLinearGradient,
    QPainter,
    QPainterPath,
    QPen,
)
from PyQt5.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QButtonGroup,
    QComboBox,
    QFrame,
    QGridLayout,
    QHBoxLayout,
    QLabel,
    QPushButton,
    QScrollBar,
    QTableWidget,
    QTableWidgetItem,
    QTabWidget,
    QVBoxLayout,
    QWidget,
)

from ..database import (
    DEFAULT_STATION_USERS,
    Database,
    WORKFLOW_EMPTY_METRIC,
    WORKFLOW_HAS_PHONE_METRIC,
    WORKFLOW_INDIVIDUAL_METRIC,
    WORKFLOW_NO_OPERATION_METRIC,
    WORKFLOW_NO_PHONE_METRIC,
    WORKFLOW_NO_TRANSPORT_METRIC,
    WORKFLOW_TOTAL_METRIC,
)
from ..models import Account
from .date_range import DateRangeSelector
from .file_dialogs import SystemFileDialog as QFileDialog
from .frameless import FramelessMessageBox as QMessageBox
from .loading_dialog import run_ui_with_loading
from .table_utils import make_table_columns_resizable


MILLISECONDS_PER_HOUR = 60 * 60 * 1000
DISTRIBUTION_CHART_HEIGHT = 300
DISTRIBUTION_PANEL_RADIUS = 8
DISTRIBUTION_TITLE_LEFT = 20
DISTRIBUTION_TITLE_BASELINE = 29
DISTRIBUTION_CHART_LEFT = 32
DISTRIBUTION_CHART_TOP = 56


def format_hours(milliseconds):
    return f"{max(0, float(milliseconds)) / MILLISECONDS_PER_HOUR:.1f} 小时"


def format_precise_duration(milliseconds):
    total_ms = max(0, int(round(float(milliseconds))))
    hours, remainder = divmod(total_ms, MILLISECONDS_PER_HOUR)
    minutes, remainder = divmod(remainder, 60 * 1000)
    seconds, milliseconds = divmod(remainder, 1000)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}.{milliseconds:03d}"


class ChartHoverCard(QFrame):
    """不受图表边界裁切、可由图表内容单独定制的悬浮信息卡。"""

    def __init__(self, parent=None):
        super().__init__(
            parent,
            Qt.ToolTip | Qt.FramelessWindowHint | Qt.NoDropShadowWindowHint,
        )
        self.title_text = ""
        self.details = []
        self._detail_row_count = 0
        self.setObjectName("ChartHoverCard")
        self.setAttribute(Qt.WA_TranslucentBackground, True)
        self.setAttribute(Qt.WA_StyledBackground, True)
        self.setAttribute(Qt.WA_TransparentForMouseEvents, True)
        self.setAttribute(Qt.WA_ShowWithoutActivating, True)
        self.setFixedWidth(240)
        self.setStyleSheet(
            """
            QFrame#ChartHoverCard {
                background: transparent;
                border: none;
            }
            QLabel#HoverCardTitle {
                color: #ffffff;
                font-size: 13px;
                font-weight: 700;
            }
            QLabel#HoverCardKey {
                color: #aebbd1;
                font-size: 12px;
            }
            QLabel#HoverCardValue {
                color: #ffffff;
                font-size: 12px;
                font-weight: 700;
            }
            """
        )
        outer = QHBoxLayout(self)
        outer.setContentsMargins(10, 10, 12, 10)
        outer.setSpacing(10)
        self._accent = QFrame()
        self._accent.setFixedWidth(4)
        outer.addWidget(self._accent)

        content = QVBoxLayout()
        content.setContentsMargins(0, 0, 0, 0)
        content.setSpacing(7)
        self._title = QLabel()
        self._title.setObjectName("HoverCardTitle")
        self._title.setWordWrap(True)
        self._title.setFixedWidth(202)
        content.addWidget(self._title)
        self._details_layout = QGridLayout()
        self._details_layout.setContentsMargins(0, 0, 0, 0)
        self._details_layout.setHorizontalSpacing(18)
        self._details_layout.setVerticalSpacing(4)
        content.addLayout(self._details_layout)
        outer.addLayout(content, 1)
        self.hide()

    def paintEvent(self, event):
        super().paintEvent(event)
        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing, True)
        card_rect = QRectF(self.rect()).adjusted(0.5, 0.5, -0.5, -0.5)
        card_path = QPainterPath()
        card_path.addRoundedRect(card_rect, 10, 10)
        painter.fillPath(card_path, QColor(23, 35, 60, 245))
        painter.setPen(QPen(QColor("#40506b"), 1))
        painter.drawPath(card_path)

    def _clear_details(self):
        while self._details_layout.count():
            item = self._details_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()

    def show_details(
        self,
        title,
        details,
        accent_color,
        anchor,
        card_width=240,
        compact=None,
        punctuated=False,
    ):
        """在全局坐标 anchor 附近显示，并保持在当前屏幕可用区域内。"""
        self.title_text = str(title)
        self.details = [(str(key), str(value)) for key, value in details]
        card_width = max(240, int(card_width))
        self.setFixedWidth(card_width)
        self._title.setFixedWidth(card_width - 38)
        self._title.setText(self.title_text)
        self._title.ensurePolished()
        title_rect = self._title.fontMetrics().boundingRect(
            QRect(0, 0, self._title.width(), 1000),
            Qt.TextWordWrap,
            self.title_text,
        )
        self._title.setFixedHeight(
            max(self._title.fontMetrics().lineSpacing(), title_rect.height())
        )
        self._accent.setStyleSheet(
            f"background:{QColor(accent_color).name()};border-radius:2px;"
        )
        self._clear_details()
        compact_grid = len(self.details) > 2 if compact is None else bool(compact)
        self._details_layout.setHorizontalSpacing(12 if compact_grid else 18)
        self._detail_row_count = (
            math.ceil(len(self.details) / 2) if compact_grid else len(self.details)
        )
        for index, (key, value) in enumerate(self.details):
            row = index // 2 if compact_grid else index
            group = index % 2 if compact_grid else 0
            key_column = group * 2
            value_column = key_column + 1
            key_label = QLabel(f"{key}：" if punctuated else key)
            key_label.setObjectName("HoverCardKey")
            key_label.setContentsMargins(0, 0, 4 if punctuated else 0, 0)
            value_label = QLabel(value)
            value_label.setObjectName("HoverCardValue")
            value_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
            self._details_layout.addWidget(key_label, row, key_column)
            self._details_layout.addWidget(value_label, row, value_column)
        self._details_layout.activate()
        self.layout().activate()
        hint = self.sizeHint()
        card_width = max(
            self.minimumWidth(), min(self.maximumWidth(), hint.width())
        )
        card_height = max(
            self.minimumSizeHint().height(),
            hint.height(),
            50 + self._detail_row_count * 22,
        )
        self.resize(card_width, card_height)
        x = anchor.x() + 16
        y = anchor.y() + 14
        screen = QApplication.screenAt(anchor)
        if screen is None:
            screen = QApplication.primaryScreen()
        available = screen.availableGeometry()
        if x + card_width > available.right() - 8:
            x = anchor.x() - card_width - 16
        if y + card_height > available.bottom() - 8:
            y = anchor.y() - card_height - 14
        x = max(available.left() + 8, min(int(x), available.right() - card_width - 8))
        y = max(available.top() + 8, min(int(y), available.bottom() - card_height - 8))
        self.move(x, y)
        self.show()
        self.raise_()


class TimingMetricCard(QFrame):
    """使用主题悬浮卡展示计时指标的精简明细。"""

    def __init__(self, hover_title, hover_details, accent_color, parent=None):
        super().__init__(parent)
        self._hover_title = str(hover_title)
        self._hover_details = list(hover_details)
        self._hover_accent = QColor(accent_color)
        self._hover_card = ChartHoverCard(self)
        self.setMouseTracking(True)

    def _show_hover_card(self, anchor):
        self._hover_card.show_details(
            self._hover_title,
            self._hover_details,
            self._hover_accent,
            anchor,
            card_width=280,
            compact=False,
            punctuated=True,
        )

    def enterEvent(self, event):
        self._show_hover_card(self.mapToGlobal(self.rect().center()))
        super().enterEvent(event)

    def mouseMoveEvent(self, event):
        self._show_hover_card(event.globalPos())
        super().mouseMoveEvent(event)

    def leaveEvent(self, event):
        self._hover_card.hide()
        super().leaveEvent(event)

    def hideEvent(self, event):
        self._hover_card.hide()
        super().hideEvent(event)


class AnimatedDonutChart(QWidget):
    """为仪表盘环形图提供悬停命中、浮出效果和计量条流光。"""

    BAR_HEIGHT = 10.0
    DONUT_HOVER_OFFSET = 9.0
    METER_SCROLLBAR_WIDTH = 12

    def __init__(self, parent=None):
        super().__init__(parent)
        self._slice_hitboxes = []
        self._hovered_slice = None
        self._animation_phase = 0.0
        self._animation_timer = QTimer(self)
        self._animation_timer.setInterval(40)
        self._animation_timer.timeout.connect(self._advance_animation)
        self._hover_card = ChartHoverCard(self)
        self._meter_viewport = QRectF()
        self._meter_row_height = 32
        self._meter_scrollbar = QScrollBar(Qt.Vertical, self)
        self._meter_scrollbar.setObjectName("ChartMeterScrollBar")
        self._meter_scrollbar.setFocusPolicy(Qt.NoFocus)
        self._meter_scrollbar.valueChanged.connect(self._on_meter_scroll)
        self._meter_scrollbar.hide()
        self.setMouseTracking(True)

    def _on_meter_scroll(self, _value):
        self._hover_card.hide()
        self.update()

    def _reset_meter_scroll(self):
        self._meter_scrollbar.setValue(0)

    def _configure_meter_scroll(self, viewport, content_height, row_height):
        """Configure the child scrollbar and return the current pixel offset."""
        self._meter_viewport = QRectF(viewport)
        self._meter_row_height = max(1, int(row_height))
        viewport_height = max(0, int(self._meter_viewport.height()))
        maximum = max(0, int(math.ceil(content_height - viewport_height)))
        self._meter_scrollbar.setGeometry(
            max(0, self.width() - self.METER_SCROLLBAR_WIDTH - 5),
            max(0, int(self._meter_viewport.top())),
            self.METER_SCROLLBAR_WIDTH,
            viewport_height,
        )
        self._meter_scrollbar.setSingleStep(self._meter_row_height)
        self._meter_scrollbar.setPageStep(max(self._meter_row_height, viewport_height))
        self._meter_scrollbar.setRange(0, maximum)
        self._meter_scrollbar.setVisible(maximum > 0)
        if maximum > 0:
            self._meter_scrollbar.raise_()
        return self._meter_scrollbar.value()

    def wheelEvent(self, event):
        if (
            self._meter_scrollbar.maximum() > 0
            and self._meter_viewport.contains(event.pos())
        ):
            pixel_delta = event.pixelDelta().y()
            angle_delta = event.angleDelta().y()
            if pixel_delta:
                amount = -pixel_delta
            else:
                steps = angle_delta / 120 if angle_delta else 0
                amount = int(-steps * self._meter_row_height * 3)
            if amount:
                self._meter_scrollbar.setValue(
                    self._meter_scrollbar.value() + amount
                )
                event.accept()
                return
        super().wheelEvent(event)

    def _advance_animation(self):
        self._animation_phase = (self._animation_phase + 0.025) % 1.0
        self.update()

    def showEvent(self, event):
        self._animation_timer.start()
        super().showEvent(event)

    def hideEvent(self, event):
        self._animation_timer.stop()
        self._hovered_slice = None
        self._hover_card.hide()
        super().hideEvent(event)

    def leaveEvent(self, event):
        self._hovered_slice = None
        self._hover_card.hide()
        self.update()
        super().leaveEvent(event)

    def _slice_at(self, position):
        for index, item in enumerate(self._slice_hitboxes):
            outer = item["outer"]
            inner = item["inner"]
            center = outer.center()
            dx = position.x() - center.x()
            dy = center.y() - position.y()
            distance = math.hypot(dx, dy)
            if distance < inner.width() / 2 or distance > outer.width() / 2:
                continue
            angle = math.degrees(math.atan2(dy, dx)) % 360
            clockwise_delta = (item["start"] - angle) % 360
            if clockwise_delta <= item["sweep"] + 0.25:
                return index, item["payload"]
        return None, None

    def _draw_donut_slice(
        self, painter, outer, inner, start_degrees, span_degrees, color, index
    ):
        wedge = QPainterPath()
        wedge.moveTo(outer.center())
        wedge.arcTo(outer, start_degrees, span_degrees)
        wedge.closeSubpath()
        hole = QPainterPath()
        hole.addEllipse(inner)
        ring_slice = wedge.subtracted(hole)
        if index == self._hovered_slice:
            midpoint = math.radians(start_degrees + span_degrees / 2)
            floating = self.DONUT_HOVER_OFFSET + math.sin(
                self._animation_phase * math.tau
            ) * 1.5
            offset_x = math.cos(midpoint) * floating
            offset_y = -math.sin(midpoint) * floating
            ring_slice.translate(offset_x, offset_y)
        painter.setPen(Qt.NoPen)
        painter.setBrush(color)
        painter.drawPath(ring_slice)

    def _draw_flow_highlight(self, painter, clip_path, rect):
        band_width = max(42.0, rect.width() * 0.12)
        travel = rect.width() + band_width * 2
        center_x = rect.left() - band_width + travel * self._animation_phase
        gradient = QLinearGradient(
            center_x - band_width,
            rect.top(),
            center_x + band_width,
            rect.top(),
        )
        gradient.setColorAt(0.0, QColor(255, 255, 255, 0))
        gradient.setColorAt(0.5, QColor(255, 255, 255, 105))
        gradient.setColorAt(1.0, QColor(255, 255, 255, 0))
        painter.save()
        painter.setClipPath(clip_path)
        painter.fillRect(rect, gradient)
        painter.restore()


class WorkflowDistributionChart(AnimatedDonutChart):
    """不依赖额外图表库的完整流程结果环形图。"""

    SEGMENTS = (
        (WORKFLOW_HAS_PHONE_METRIC, "有公司名、有电话", QColor("#1d9a84")),
        (WORKFLOW_NO_TRANSPORT_METRIC, "无运输证号", QColor("#d86464")),
        (WORKFLOW_INDIVIDUAL_METRIC, "个体经营", QColor("#6f8f89")),
        (WORKFLOW_NO_OPERATION_METRIC, "无营运信息", QColor("#df8b55")),
        (WORKFLOW_NO_PHONE_METRIC, "有公司名、无电话", QColor("#e2a64a")),
    )
    BAR_ROW_HEIGHT = 38

    def __init__(self, parent=None):
        super().__init__(parent)
        self._values = {}
        self._scope = ""
        self._bar_rects = []
        self._bar_payloads = []
        self.setMinimumHeight(280)

    def set_values(self, values, scope=""):
        self._values = {key: int(value or 0) for key, value in values.items()}
        self._scope = scope
        self._reset_meter_scroll()
        self.update()

    def paintEvent(self, event):
        self._slice_hitboxes = []
        self._bar_rects = []
        self._bar_payloads = []
        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing)
        bounds = QRectF(self.rect()).adjusted(1, 1, -1, -1)
        painter.setPen(QPen(QColor("#d8e7e3"), 1))
        painter.setBrush(QColor("#ffffff"))
        painter.drawRoundedRect(
            bounds,
            DISTRIBUTION_PANEL_RADIUS,
            DISTRIBUTION_PANEL_RADIUS,
        )

        painter.setPen(QColor("#173a3d"))
        painter.setFont(QFont("Microsoft YaHei UI", 11, QFont.Bold))
        painter.drawText(
            DISTRIBUTION_TITLE_LEFT,
            DISTRIBUTION_TITLE_BASELINE,
            "完整流程结果分布",
        )
        painter.setPen(QColor("#647c7b"))
        painter.setFont(QFont("Microsoft YaHei UI", 9))
        painter.drawText(160, 29, self._scope)

        chart_size = min(190, max(150, self.height() - 82))
        pie_rect = QRectF(
            DISTRIBUTION_CHART_LEFT,
            DISTRIBUTION_CHART_TOP,
            chart_size,
            chart_size,
        )
        inner = pie_rect.adjusted(
            chart_size * 0.28,
            chart_size * 0.28,
            -chart_size * 0.28,
            -chart_size * 0.28,
        )
        segment_values = [self._values.get(key, 0) for key, _, _ in self.SEGMENTS]
        classified_total = sum(segment_values)
        start_degrees = 90.0
        if classified_total:
            for index, (value, (_, label, color)) in enumerate(
                zip(segment_values, self.SEGMENTS)
            ):
                span_degrees = -(value / classified_total * 360)
                if value:
                    self._slice_hitboxes.append(
                        {
                            "outer": QRectF(pie_rect),
                            "inner": QRectF(inner),
                            "start": start_degrees,
                            "sweep": abs(span_degrees),
                            "payload": {
                                "label": label,
                                "value": value,
                                "total": classified_total,
                                "color": color,
                            },
                        }
                    )
                    slice_index = len(self._slice_hitboxes) - 1
                    self._draw_donut_slice(
                        painter,
                        pie_rect,
                        inner,
                        start_degrees,
                        span_degrees,
                        color,
                        slice_index,
                    )
                start_degrees += span_degrees
        else:
            painter.setPen(Qt.NoPen)
            painter.setBrush(QColor("#e9eef5"))
            painter.drawEllipse(pie_rect)
            painter.setBrush(QColor("#ffffff"))
            painter.drawEllipse(inner)
        total = self._values.get(WORKFLOW_TOTAL_METRIC, classified_total)
        painter.setPen(QColor("#173a3d"))
        painter.setFont(QFont("Microsoft YaHei UI", 18, QFont.Bold))
        painter.drawText(inner, Qt.AlignCenter, str(total))
        painter.setPen(QColor("#647c7b"))
        painter.setFont(QFont("Microsoft YaHei UI", 8))
        painter.drawText(
            QRectF(inner.left(), inner.center().y() + 17, inner.width(), 20),
            Qt.AlignHCenter | Qt.AlignTop,
            "总计",
        )

        legend_left = pie_rect.right() + 42
        legend_width = max(120, self.width() - legend_left - 28)
        row_height = self.BAR_ROW_HEIGHT
        legend_rows = sorted(
            zip(self.SEGMENTS, segment_values),
            key=lambda item: (-int(item[1]), str(item[0][1])),
        )
        meter_viewport = QRectF(
            legend_left,
            DISTRIBUTION_CHART_TOP,
            max(0, self.width() - legend_left - 22),
            max(0, self.height() - DISTRIBUTION_CHART_TOP - 14),
        )
        scroll_offset = self._configure_meter_scroll(
            meter_viewport,
            len(legend_rows) * row_height,
            row_height,
        )
        painter.save()
        painter.setClipRect(meter_viewport)
        for index, ((metric_key, label, color), value) in enumerate(legend_rows):
            top = DISTRIBUTION_CHART_TOP + index * row_height - scroll_offset
            painter.setPen(Qt.NoPen)
            painter.setBrush(color)
            painter.drawRoundedRect(QRectF(legend_left, top + 2, 10, 10), 3, 3)
            painter.setPen(QColor("#526e6d"))
            painter.setFont(QFont("Microsoft YaHei UI", 9))
            painter.drawText(int(legend_left + 18), int(top + 12), label)
            percent = value / classified_total * 100 if classified_total else 0
            painter.setPen(QColor("#173a3d"))
            painter.setFont(QFont("Microsoft YaHei UI", 10, QFont.Bold))
            painter.drawText(
                QRectF(legend_left, top - 3, legend_width, 20),
                Qt.AlignRight | Qt.AlignVCenter,
                f"{value} 条  ·  {percent:.1f}%",
            )
            bar_rect = QRectF(
                legend_left, top + 20, legend_width, self.BAR_HEIGHT
            )
            if bar_rect.intersects(meter_viewport):
                self._bar_rects.append(QRectF(bar_rect))
                self._bar_payloads.append((metric_key, label, color))
            painter.setPen(Qt.NoPen)
            painter.setBrush(QColor("#e8f1ef"))
            painter.drawRoundedRect(bar_rect, 5, 5)
            if classified_total and value:
                value_rect = QRectF(bar_rect)
                value_rect.setWidth(max(10, bar_rect.width() * value / classified_total))
                painter.setBrush(color)
                value_path = QPainterPath()
                value_path.addRoundedRect(value_rect, 5, 5)
                painter.drawPath(value_path)
                self._draw_flow_highlight(painter, value_path, value_rect)
        painter.restore()

    def _show_category_details(self, label, value, total, color, anchor):
        percent = value / max(total, 1) * 100
        self._hover_card.show_details(
            label,
            [("数量", f"{value} 条"), ("占比", f"{percent:.1f}%")],
            color,
            anchor,
        )

    def mouseMoveEvent(self, event):
        index, payload = self._slice_at(event.pos())
        if payload is not None:
            self._hovered_slice = index
            self._show_category_details(
                payload["label"],
                payload["value"],
                payload["total"],
                payload["color"],
                event.globalPos(),
            )
            self.update()
            return
        classified_total = sum(
            self._values.get(metric_key, 0)
            for metric_key, _, _ in self.SEGMENTS
        )
        for bar_index, bar_rect in enumerate(self._bar_rects):
            if bar_rect.contains(event.pos()):
                if self._hovered_slice is not None:
                    self._hovered_slice = None
                    self.update()
                metric_key, label, color = self._bar_payloads[bar_index]
                self._show_category_details(
                    label,
                    self._values.get(metric_key, 0),
                    classified_total,
                    color,
                    event.globalPos(),
                )
                return
        if self._hovered_slice is not None:
            self._hovered_slice = None
            self.update()
        self._hover_card.hide()
        super().mouseMoveEvent(event)


class ViolationReasonChart(AnimatedDonutChart):
    """违规原因环形图及可切换的电话拆分/原因占比计量条。"""

    BAR_ROW_HEIGHT = 33
    COLORS = (
        QColor("#1d8178"),
        QColor("#1d9a84"),
        QColor("#e2a64a"),
        QColor("#d86464"),
        QColor("#6f8f89"),
        QColor("#df8b55"),
        QColor("#4a9c91"),
        QColor("#7f93ad"),
        QColor("#c765a7"),
        QColor("#65a84f"),
    )
    PHONE_COLOR = QColor("#1d9a84")
    OTHER_COLOR = QColor("#df8b55")

    def __init__(self, parent=None):
        super().__init__(parent)
        self._rows = []
        self._scope = ""
        self._bar_hitboxes = []
        self._bar_mode = "share"
        self.mode_button = QPushButton("切换为电话拆分", self)
        self.mode_button.setObjectName("ViolationModeButton")
        self.mode_button.setCheckable(True)
        self.mode_button.setChecked(True)
        self.mode_button.setFixedSize(150, 32)
        self.mode_button.toggled.connect(self._on_mode_button_toggled)
        self.setMinimumHeight(280)

    def set_rows(self, rows, scope=""):
        self._rows = sorted(
            (dict(row) for row in rows),
            key=lambda row: (
                -int(row.get("total") or 0),
                str(row.get("reason") or ""),
            ),
        )
        self._scope = scope
        self._reset_meter_scroll()
        self.update()

    def set_bar_mode(self, mode):
        if mode not in {"split", "share"}:
            raise ValueError(f"Unsupported violation bar mode: {mode}")
        self._bar_mode = mode
        checked = mode == "share"
        self.mode_button.blockSignals(True)
        self.mode_button.setChecked(checked)
        self.mode_button.blockSignals(False)
        self.mode_button.setText(
            "切换为电话拆分" if checked else "切换为原因占比"
        )
        self._hover_card.hide()
        self.update()

    def _on_mode_button_toggled(self, checked):
        self.set_bar_mode("share" if checked else "split")

    def resizeEvent(self, event):
        chart_size = min(190, max(150, self.height() - 82))
        preferred_x = int(32 + chart_size + 46)
        right_aligned_x = max(8, self.width() - self.mode_button.width() - 12)
        self.mode_button.move(min(preferred_x, right_aligned_x), 7)
        self.mode_button.raise_()
        super().resizeEvent(event)

    @staticmethod
    def _format_total_label(total):
        return f"{int(total)} 条"

    def _format_bar_label(self, total, all_total):
        if self._bar_mode == "share":
            percent = int(total) / max(int(all_total), 1) * 100
            return f"{int(total)} 条  ·  {percent:.1f}%"
        return self._format_total_label(total)

    def paintEvent(self, event):
        self._slice_hitboxes = []
        self._bar_hitboxes = []
        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing)
        bounds = QRectF(self.rect()).adjusted(1, 1, -1, -1)
        painter.setPen(QPen(QColor("#d8e7e3"), 1))
        painter.setBrush(QColor("#ffffff"))
        painter.drawRoundedRect(
            bounds,
            DISTRIBUTION_PANEL_RADIUS,
            DISTRIBUTION_PANEL_RADIUS,
        )

        painter.setPen(QColor("#173a3d"))
        painter.setFont(QFont("Microsoft YaHei UI", 11, QFont.Bold))
        painter.drawText(
            DISTRIBUTION_TITLE_LEFT,
            DISTRIBUTION_TITLE_BASELINE,
            "违规原因分布",
        )
        painter.setPen(QColor("#647c7b"))
        painter.setFont(QFont("Microsoft YaHei UI", 9))
        painter.drawText(130, 29, self._scope)

        legend_x = max(320, self.width() - 214)
        if self._bar_mode == "split":
            painter.setPen(Qt.NoPen)
            painter.setBrush(self.PHONE_COLOR)
            painter.drawRoundedRect(QRectF(legend_x, 19, 9, 9), 3, 3)
            painter.setPen(QColor("#647c7b"))
            painter.drawText(legend_x + 14, 29, "有电话")
            painter.setPen(Qt.NoPen)
            painter.setBrush(self.OTHER_COLOR)
            painter.drawRoundedRect(QRectF(legend_x + 82, 19, 9, 9), 3, 3)
            painter.setPen(QColor("#647c7b"))
            painter.drawText(legend_x + 96, 29, "其他数据")
        else:
            painter.setPen(QColor("#647c7b"))
            painter.drawText(
                QRectF(legend_x, 14, 186, 20),
                Qt.AlignRight | Qt.AlignVCenter,
                "条长表示原因占比",
            )

        if not self._rows:
            self._configure_meter_scroll(
                QRectF(320, 46, max(0, self.width() - 342), max(0, self.height() - 60)),
                0,
                self.BAR_ROW_HEIGHT,
            )
            painter.setPen(QColor("#8a98aa"))
            painter.drawText(self.rect(), Qt.AlignCenter, "当前筛选范围暂无“原因”数据")
            return

        total = sum(int(row["total"]) for row in self._rows)
        chart_size = min(190, max(150, self.height() - 82))
        pie_rect = QRectF(
            DISTRIBUTION_CHART_LEFT,
            DISTRIBUTION_CHART_TOP,
            chart_size,
            chart_size,
        )
        inner = pie_rect.adjusted(
            chart_size * 0.28,
            chart_size * 0.28,
            -chart_size * 0.28,
            -chart_size * 0.28,
        )
        start_degrees = 90.0
        for index, row in enumerate(self._rows):
            span_degrees = -(int(row["total"]) / max(total, 1) * 360)
            if int(row["total"]):
                color = self.COLORS[index % len(self.COLORS)]
                self._slice_hitboxes.append(
                    {
                        "outer": QRectF(pie_rect),
                        "inner": QRectF(inner),
                        "start": start_degrees,
                        "sweep": abs(span_degrees),
                        "payload": {**dict(row), "color": color},
                    }
                )
                slice_index = len(self._slice_hitboxes) - 1
                self._draw_donut_slice(
                    painter,
                    pie_rect,
                    inner,
                    start_degrees,
                    span_degrees,
                    color,
                    slice_index,
                )
            start_degrees += span_degrees
        painter.setPen(QColor("#173a3d"))
        painter.setFont(QFont("Microsoft YaHei UI", 18, QFont.Bold))
        painter.drawText(inner, Qt.AlignCenter, str(total))
        painter.setPen(QColor("#647c7b"))
        painter.setFont(QFont("Microsoft YaHei UI", 8))
        painter.drawText(
            QRectF(inner.left(), inner.center().y() + 17, inner.width(), 20),
            Qt.AlignHCenter | Qt.AlignTop,
            "原因条目",
        )

        rows = self._rows
        legend_left = pie_rect.right() + 46
        legend_width = max(120, self.width() - legend_left - 28)
        label_width = max(90, int(legend_width * 0.56))
        row_height = self.BAR_ROW_HEIGHT
        top = 46
        meter_viewport = QRectF(
            legend_left,
            top,
            max(0, self.width() - legend_left - 22),
            max(0, self.height() - top - 14),
        )
        scroll_offset = self._configure_meter_scroll(
            meter_viewport,
            len(rows) * row_height,
            row_height,
        )
        painter.setFont(QFont("Microsoft YaHei UI", 9))

        painter.save()
        painter.setClipRect(meter_viewport)
        for index, row in enumerate(rows):
            row_top = top + index * row_height - scroll_offset
            color = self.COLORS[index % len(self.COLORS)]
            painter.setPen(Qt.NoPen)
            painter.setBrush(color)
            painter.drawRoundedRect(QRectF(legend_left, row_top + 2, 10, 10), 3, 3)
            painter.setFont(QFont("Microsoft YaHei UI", 9))
            reason = painter.fontMetrics().elidedText(
                str(row["reason"]), Qt.ElideRight, label_width - 24
            )
            label_rect = QRectF(legend_left + 18, row_top, label_width - 18, 16)
            painter.setPen(QColor("#526e6d"))
            painter.drawText(label_rect, Qt.AlignVCenter, reason)

            painter.setPen(QColor("#173a3d"))
            painter.setFont(QFont("Microsoft YaHei UI", 9, QFont.Bold))
            painter.drawText(
                QRectF(legend_left, row_top, legend_width, 16),
                Qt.AlignRight | Qt.AlignVCenter,
                self._format_bar_label(row["total"], total),
            )

            bar_rect = QRectF(
                legend_left, row_top + 18, legend_width, self.BAR_HEIGHT
            )
            if bar_rect.intersects(meter_viewport):
                self._bar_hitboxes.append((bar_rect, dict(row)))
            painter.setPen(Qt.NoPen)
            bar_path = QPainterPath()
            bar_path.addRoundedRect(bar_rect, 5, 5)
            if self._bar_mode == "split":
                phone_count = max(0, int(row["has_phone"]))
                other_count = max(0, int(row["other"]))
                split_total = phone_count + other_count
                phone_ratio = phone_count / split_total if split_total else 0
                painter.save()
                painter.setClipPath(bar_path)
                painter.fillRect(bar_rect, self.OTHER_COLOR)
                if phone_ratio:
                    phone_rect = QRectF(bar_rect)
                    phone_rect.setWidth(bar_rect.width() * phone_ratio)
                    painter.fillRect(phone_rect, self.PHONE_COLOR)
                painter.restore()
                self._draw_flow_highlight(painter, bar_path, bar_rect)
            else:
                painter.setBrush(QColor("#e8f1ef"))
                painter.drawPath(bar_path)
                if int(row["total"]):
                    value_rect = QRectF(bar_rect)
                    value_rect.setWidth(
                        max(10, bar_rect.width() * int(row["total"]) / max(total, 1))
                    )
                    value_path = QPainterPath()
                    value_path.addRoundedRect(value_rect, 5, 5)
                    painter.setBrush(color)
                    painter.drawPath(value_path)
                    self._draw_flow_highlight(painter, value_path, value_rect)
        painter.restore()

    def mouseMoveEvent(self, event):
        slice_index, slice_row = self._slice_at(event.pos())
        if slice_row is not None:
            self._hovered_slice = slice_index
            total = int(slice_row["total"])
            all_total = sum(int(row["total"]) for row in self._rows)
            percent = total / max(all_total, 1) * 100
            self._hover_card.show_details(
                slice_row["reason"],
                [
                    ("总计", f"{total} 条"),
                    ("占比", f"{percent:.1f}%"),
                    ("有电话", f"{slice_row['has_phone']} 条"),
                    ("其他数据", f"{slice_row['other']} 条"),
                ],
                slice_row["color"],
                event.globalPos(),
            )
            self.update()
            return
        for rect, row in self._bar_hitboxes:
            if rect.contains(event.pos()):
                if self._hovered_slice is not None:
                    self._hovered_slice = None
                    self.update()
                self._hover_card.show_details(
                    "电话数据拆分",
                    [
                        ("有电话", f"{row['has_phone']} 条"),
                        ("其他数据", f"{row['other']} 条"),
                    ],
                    self.PHONE_COLOR,
                    event.globalPos(),
                )
                return
        if self._hovered_slice is not None:
            self._hovered_slice = None
            self.update()
        self._hover_card.hide()
        super().mouseMoveEvent(event)


class StationDistributionChart(AnimatedDonutChart):
    """Compare station counts with bars and station timing with one donut."""

    TOTAL_COLOR = QColor("#1d8178")
    PHONE_COLOR = QColor("#f0a45d")
    COLORS = (
        QColor("#1d8178"),
        QColor("#f0a45d"),
        QColor("#d2a441"),
        QColor("#d86464"),
        QColor("#6f8f89"),
        QColor("#df8b55"),
        QColor("#4a9c91"),
    )
    COUNT_ROW_HEIGHT = 42
    TIMING_ROW_HEIGHT = 34

    def __init__(self, parent=None):
        super().__init__(parent)
        self._rows = []
        self._legend_hitboxes = []
        self._bar_rects = []
        self._series_mode = "counts"
        self._timing_basis = "total"
        self._count_metric = "total"
        self.count_metric_group = QButtonGroup(self)
        self.count_metric_group.setExclusive(True)
        self.total_count_button = QPushButton("总计数", self)
        self.phone_count_button = QPushButton("有电话计数", self)
        for button, metric in (
            (self.total_count_button, "total"),
            (self.phone_count_button, "has_phone"),
        ):
            button.setObjectName("StationMetricSegment")
            button.setCheckable(True)
            button.setProperty("metric", metric)
            button.setFixedHeight(30)
            self.count_metric_group.addButton(button)
            button.clicked.connect(
                lambda _checked=False, value=metric: self.set_count_metric(value)
            )
        self.total_count_button.setChecked(True)
        self.setMinimumHeight(280)

    def set_rows(self, rows):
        self._rows = [dict(row) for row in rows]
        self._sort_rows()
        self._hovered_slice = None
        self._hover_card.hide()
        self._reset_meter_scroll()
        self.update()

    def _sort_rows(self):
        if self._series_mode == "timing":
            value_key = (
                "active_ms" if self._timing_basis == "active" else "total_time_ms"
            )
            tie_key = "total"
        else:
            value_key = self._count_metric
            tie_key = "has_phone" if value_key == "total" else "total"
        self._rows.sort(
            key=lambda row: (
                -int(row.get(value_key) or 0),
                -int(row.get(tie_key) or 0),
                str(row.get("station") or "").casefold(),
            )
        )

    def set_series_mode(self, mode):
        if mode not in {"counts", "timing"}:
            raise ValueError(f"Unsupported station distribution mode: {mode}")
        self._series_mode = mode
        show_count_switch = mode == "counts"
        self.total_count_button.setVisible(show_count_switch)
        self.phone_count_button.setVisible(show_count_switch)
        self._sort_rows()
        self._hovered_slice = None
        self._hover_card.hide()
        self._reset_meter_scroll()
        self.update()

    def set_count_metric(self, metric):
        if metric not in {"total", "has_phone"}:
            raise ValueError(f"Unsupported station count metric: {metric}")
        self._count_metric = metric
        target = (
            self.total_count_button
            if metric == "total"
            else self.phone_count_button
        )
        target.setChecked(True)
        self._sort_rows()
        self._hover_card.hide()
        self._reset_meter_scroll()
        self.update()

    def resizeEvent(self, event):
        right = self.width() - 18
        phone_width = 108
        total_width = 82
        self.phone_count_button.setGeometry(
            right - phone_width,
            8,
            phone_width,
            30,
        )
        self.total_count_button.setGeometry(
            right - phone_width - total_width + 1,
            8,
            total_width,
            30,
        )
        self.total_count_button.raise_()
        self.phone_count_button.raise_()
        super().resizeEvent(event)

    def set_timing_basis(self, basis):
        if basis not in {"total", "active"}:
            raise ValueError(f"Unsupported timing basis: {basis}")
        self._timing_basis = basis
        self._sort_rows()
        self._hovered_slice = None
        self._hover_card.hide()
        self._reset_meter_scroll()
        self.update()

    @staticmethod
    def _format_share(value):
        return f"{float(value):.1f}%"

    def _draw_series_donut(
        self,
        painter,
        outer,
        value_key,
        share_key,
        title,
        center_label,
        is_duration=False,
        show_title=True,
    ):
        inner = outer.adjusted(
            outer.width() * 0.28,
            outer.height() * 0.28,
            -outer.width() * 0.28,
            -outer.height() * 0.28,
        )
        total = sum(max(0, int(row[value_key])) for row in self._rows)
        if show_title:
            painter.setPen(QColor("#526e6d"))
            painter.setFont(QFont("Microsoft YaHei UI", 9, QFont.Bold))
            painter.drawText(
                QRectF(outer.left(), outer.top() - 27, outer.width(), 20),
                Qt.AlignCenter,
                title,
            )

        if total:
            start_degrees = 90.0
            for index, row in enumerate(self._rows):
                value = int(row[value_key])
                span_degrees = -(value / total * 360)
                if value:
                    color = self.COLORS[index % len(self.COLORS)]
                    self._slice_hitboxes.append(
                        {
                            "outer": QRectF(outer),
                            "inner": QRectF(inner),
                            "start": start_degrees,
                            "sweep": abs(span_degrees),
                            "payload": {
                                "station": row["station"],
                                "series": title,
                                "value": value,
                                "share": row[share_key],
                                "color": color,
                                "is_duration": is_duration,
                            },
                        }
                    )
                    slice_index = len(self._slice_hitboxes) - 1
                    self._draw_donut_slice(
                        painter,
                        outer,
                        inner,
                        start_degrees,
                        span_degrees,
                        color,
                        slice_index,
                    )
                start_degrees += span_degrees
        else:
            painter.setPen(Qt.NoPen)
            painter.setBrush(QColor("#e9eef5"))
            painter.drawEllipse(outer)
            painter.setBrush(QColor("#ffffff"))
            painter.drawEllipse(inner)

        painter.setPen(QColor("#173a3d"))
        painter.setFont(QFont("Microsoft YaHei UI", 17, QFont.Bold))
        painter.drawText(
            QRectF(inner.left(), inner.top() - 5, inner.width(), inner.height()),
            Qt.AlignCenter,
            (
                f"{total / MILLISECONDS_PER_HOUR:.1f}"
                if is_duration
                else str(total)
            ),
        )
        painter.setPen(QColor("#647c7b"))
        painter.setFont(QFont("Microsoft YaHei UI", 8))
        painter.drawText(
            QRectF(inner.left(), inner.center().y() + 14, inner.width(), 18),
            Qt.AlignHCenter | Qt.AlignTop,
            center_label,
        )

    def _paint_count_bars(self, painter):
        value_key = self._count_metric
        share_key = "total_share" if value_key == "total" else "phone_share"
        metric_label = "总计数" if value_key == "total" else "有电话计数"
        maximum = max(
            (max(0, int(row.get(value_key) or 0)) for row in self._rows),
            default=0,
        )
        left = DISTRIBUTION_CHART_LEFT
        right = self.width() - 22
        width = max(120, right - left)
        top = DISTRIBUTION_CHART_TOP
        viewport = QRectF(
            left,
            top,
            width,
            max(0, self.height() - top - 14),
        )
        scroll_offset = self._configure_meter_scroll(
            viewport,
            len(self._rows) * self.COUNT_ROW_HEIGHT,
            self.COUNT_ROW_HEIGHT,
        )
        painter.save()
        painter.setClipRect(viewport)
        for index, row in enumerate(self._rows):
            row_top = top + index * self.COUNT_ROW_HEIGHT - scroll_offset
            row_rect = QRectF(
                left,
                row_top,
                width,
                self.COUNT_ROW_HEIGHT - 3,
            )
            if row_rect.intersects(viewport):
                self._legend_hitboxes.append((QRectF(row_rect), dict(row)))
            if index % 2:
                painter.setPen(Qt.NoPen)
                painter.setBrush(QColor("#f5f8f7"))
                painter.drawRoundedRect(row_rect, 5, 5)

            color = self.COLORS[index % len(self.COLORS)]
            painter.setPen(Qt.NoPen)
            painter.setBrush(color)
            painter.drawRoundedRect(QRectF(left + 2, row_top + 4, 10, 10), 3, 3)
            painter.setPen(QColor("#526e6d"))
            painter.setFont(QFont("Microsoft YaHei UI", 9))
            station_width = max(120, int(width * 0.58))
            station = painter.fontMetrics().elidedText(
                str(row.get("station") or "-"),
                Qt.ElideRight,
                station_width,
            )
            painter.drawText(
                QRectF(left + 20, row_top, station_width, 19),
                Qt.AlignLeft | Qt.AlignVCenter,
                station,
            )
            value = max(0, int(row.get(value_key) or 0))
            share = float(row.get(share_key) or 0)
            painter.setPen(QColor("#173a3d"))
            painter.setFont(QFont("Microsoft YaHei UI", 9, QFont.Bold))
            painter.drawText(
                QRectF(left, row_top, width - 8, 19),
                Qt.AlignRight | Qt.AlignVCenter,
                f"{value} 条  ·  {share:.1f}%",
            )
            bar_rect = QRectF(left, row_top + 23, width, self.BAR_HEIGHT)
            if bar_rect.intersects(viewport):
                self._bar_rects.append(QRectF(bar_rect))
            painter.setPen(Qt.NoPen)
            painter.setBrush(QColor("#e8f1ef"))
            painter.drawRoundedRect(bar_rect, 5, 5)
            if maximum and value:
                value_rect = QRectF(bar_rect)
                value_rect.setWidth(max(10, bar_rect.width() * value / maximum))
                value_path = QPainterPath()
                value_path.addRoundedRect(value_rect, 5, 5)
                painter.setBrush(color)
                painter.drawPath(value_path)
                self._draw_flow_highlight(painter, value_path, value_rect)
        painter.restore()

        painter.setPen(QColor("#647c7b"))
        painter.setFont(QFont("Microsoft YaHei UI", 9))
        painter.drawText(155, DISTRIBUTION_TITLE_BASELINE, f"{metric_label}降序")

    def paintEvent(self, event):
        self._slice_hitboxes = []
        self._legend_hitboxes = []
        self._bar_rects = []
        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing)
        bounds = QRectF(self.rect()).adjusted(1, 1, -1, -1)
        painter.setPen(QPen(QColor("#d8e7e3"), 1))
        painter.setBrush(QColor("#ffffff"))
        painter.drawRoundedRect(
            bounds,
            DISTRIBUTION_PANEL_RADIUS,
            DISTRIBUTION_PANEL_RADIUS,
        )

        is_timing_mode = self._series_mode == "timing"
        panel_title = "各站用时效率分布" if is_timing_mode else "各站业务分布"
        timing_label = "有效用时" if self._timing_basis == "active" else "总用时"
        panel_subtitle = f"{timing_label}占比" if is_timing_mode else ""
        painter.setPen(QColor("#173a3d"))
        painter.setFont(QFont("Microsoft YaHei UI", 11, QFont.Bold))
        painter.drawText(
            DISTRIBUTION_TITLE_LEFT,
            DISTRIBUTION_TITLE_BASELINE,
            panel_title,
        )
        painter.setPen(QColor("#647c7b"))
        painter.setFont(QFont("Microsoft YaHei UI", 9))
        if panel_subtitle:
            painter.drawText(155, DISTRIBUTION_TITLE_BASELINE, panel_subtitle)

        if not self._rows:
            self._configure_meter_scroll(
                QRectF(320, 59, max(0, self.width() - 342), max(0, self.height() - 73)),
                0,
                self.COUNT_ROW_HEIGHT,
            )
            painter.setPen(QColor("#8a98aa"))
            painter.drawText(self.rect(), Qt.AlignCenter, "暂无站点数据")
            return

        if not is_timing_mode:
            self._paint_count_bars(painter)
            return

        if self._timing_basis == "active":
            series = (
                (
                    "active_ms",
                    "active_time_share",
                    "各站有效耗时占比",
                    "小时",
                    True,
                ),
            )
        else:
            series = (
                (
                    "total_time_ms",
                    "total_time_share",
                    "各站总耗时占比",
                    "小时",
                    True,
                ),
            )
        # Match WorkflowDistributionChart: the donut starts at the same
        # left/top position and the meter list follows it directly.
        chart_size = min(190, max(150, self.height() - 82))
        chart_top = DISTRIBUTION_CHART_TOP
        first_left = DISTRIBUTION_CHART_LEFT
        for index, (value_key, share_key, title, center_label, is_duration) in enumerate(series):
            left = first_left + index * (chart_size + 24)
            self._draw_series_donut(
                painter,
                QRectF(left, chart_top, chart_size, chart_size),
                value_key,
                share_key,
                title,
                center_label,
                is_duration=is_duration,
                show_title=False,
            )

        legend_left = first_left + chart_size + 42
        legend_width = max(
            180,
            self.width() - legend_left - 28,
        )
        first_column_width = min(
            180,
            max(118, int(legend_width * 0.48)),
        )
        first_column = legend_left + legend_width - first_column_width
        painter.setFont(QFont("Microsoft YaHei UI", 8, QFont.Bold))
        row_height = self.TIMING_ROW_HEIGHT
        top = 50
        meter_viewport = QRectF(
            legend_left,
            top,
            max(0, self.width() - legend_left - 22),
            max(0, self.height() - top - 14),
        )
        scroll_offset = self._configure_meter_scroll(
            meter_viewport,
            len(self._rows) * row_height,
            row_height,
        )
        label_width = max(90, first_column - legend_left - 24)
        painter.setFont(QFont("Microsoft YaHei UI", 9))

        painter.save()
        painter.setClipRect(meter_viewport)
        for index, row in enumerate(self._rows):
            row_top = top + index * row_height - scroll_offset
            hitbox = QRectF(
                legend_left, row_top, legend_width, max(27, row_height - 2)
            )
            if hitbox.intersects(meter_viewport):
                self._legend_hitboxes.append((hitbox, dict(row)))
            color = self.COLORS[index % len(self.COLORS)]
            label_height = 20
            painter.setPen(Qt.NoPen)
            painter.setBrush(color)
            painter.drawRoundedRect(
                QRectF(legend_left + 2, row_top + label_height / 2 - 5, 10, 10),
                3,
                3,
            )
            painter.setPen(QColor("#526e6d"))
            station = painter.fontMetrics().elidedText(
                str(row["station"]), Qt.ElideRight, label_width
            )
            painter.drawText(
                QRectF(legend_left + 18, row_top, label_width, label_height),
                Qt.AlignLeft | Qt.AlignVCenter,
                station,
            )
            painter.setPen(QColor("#173a3d"))
            painter.drawText(
                QRectF(first_column, row_top, first_column_width, label_height),
                Qt.AlignRight | Qt.AlignVCenter,
                f'{format_hours(row[series[0][0]])} · '
                f'{self._format_share(row[series[0][1]])}',
            )
            bar_rect = QRectF(
                legend_left,
                row_top + 20,
                legend_width,
                self.BAR_HEIGHT,
            )
            if bar_rect.intersects(meter_viewport):
                self._bar_rects.append(QRectF(bar_rect))
            painter.setPen(Qt.NoPen)
            painter.setBrush(QColor("#e8f1ef"))
            painter.drawRoundedRect(bar_rect, 5, 5)
            share = max(0.0, min(100.0, float(row[series[0][1]])))
            if share:
                value_rect = QRectF(bar_rect)
                value_rect.setWidth(max(10, bar_rect.width() * share / 100))
                painter.setBrush(color)
                value_path = QPainterPath()
                value_path.addRoundedRect(value_rect, 5, 5)
                painter.drawPath(value_path)
                self._draw_flow_highlight(painter, value_path, value_rect)
        painter.restore()

    def mouseMoveEvent(self, event):
        slice_index, payload = self._slice_at(event.pos())
        if payload is not None:
            self._hovered_slice = slice_index
            details = (
                [
                    ("小时数", format_hours(payload["value"])),
                    ("精确耗时", format_precise_duration(payload["value"])),
                    ("占比", self._format_share(payload["share"])),
                ]
                if payload["is_duration"]
                else [
                    ("数量", f'{payload["value"]} 条'),
                    ("占比", self._format_share(payload["share"])),
                ]
            )
            self._hover_card.show_details(
                f'{payload["station"]} · {payload["series"]}',
                details,
                payload["color"],
                event.globalPos(),
                card_width=300 if payload["is_duration"] else 240,
            )
            self.update()
            return
        for rect, row in self._legend_hitboxes:
            if rect.contains(event.pos()):
                if self._hovered_slice is not None:
                    self._hovered_slice = None
                    self.update()
                if self._series_mode == "timing":
                    timing_label = (
                        "有效用时"
                        if self._timing_basis == "active"
                        else "总用时"
                    )
                    timing_key = (
                        "active_ms"
                        if self._timing_basis == "active"
                        else "total_time_ms"
                    )
                    share_key = (
                        "active_time_share"
                        if self._timing_basis == "active"
                        else "total_time_share"
                    )
                    details = [
                        (timing_label, format_hours(row[timing_key])),
                        (
                            f"精确{timing_label}",
                            format_precise_duration(row[timing_key]),
                        ),
                        (
                            f"{timing_label}占比",
                            self._format_share(row[share_key]),
                        ),
                    ]
                else:
                    details = [
                        ("总计数", f'{row["total"]} 条'),
                        ("总数占比", self._format_share(row["total_share"])),
                        ("有电话数", f'{row["has_phone"]} 条'),
                        ("有电话占比", self._format_share(row["phone_share"])),
                    ]
                self._hover_card.show_details(
                    row["station"],
                    details,
                    self.COLORS[
                        self._rows.index(row) % len(self.COLORS)
                    ],
                    event.globalPos(),
                    card_width=340,
                    compact=self._series_mode != "timing",
                )
                return
        if self._hovered_slice is not None:
            self._hovered_slice = None
            self.update()
        self._hover_card.hide()
        super().mouseMoveEvent(event)


class StatisticsPage(QWidget):
    HUMAN_BASELINE_ITEMS = 260
    HUMAN_BASELINE_MS = 6 * 60 * 60 * 1000
    NAVIGATION_VIEWS = {
        "station_distribution": (
            "全站分布",
            "查看各站查询量与电话覆盖分布。",
        ),
        "timing": (
            "用时效率",
            "切换总用时或有效用时口径，查看处理效率和效率提升。",
        ),
        "completion": (
            "完成类型",
            "按站点和日期查看完整流程结果分类。",
        ),
        "violation": (
            "违规原因",
            "查看违规原因构成以及有电话数据占比。",
        ),
        "anomaly": (
            "异常数据",
            "追溯空单元格异常数量、来源站点和占比。",
        ),
    }

    def __init__(self, database: Database, account: Account, dependency_warning="", parent=None):
        super().__init__(parent)
        self.database = database
        self.account = account
        self._sidebar_navigation = False
        self._timing_basis = "total"

        layout = QVBoxLayout(self)
        layout.setContentsMargins(22, 20, 22, 22)
        layout.setSpacing(10)

        header = QHBoxLayout()
        title_box = QVBoxLayout()
        self.title_label = QLabel("全站分布")
        self.title_label.setObjectName("PageTitle")
        title_box.addWidget(self.title_label)
        self.subtitle_label = QLabel(
            "按站点查看完成类型或违规原因统计。"
            if account.can_view_all_stats
            else "默认显示当前站点，可切换查看全部或其他站点的数据。"
        )
        self.subtitle_label.setObjectName("Muted")
        title_box.addWidget(self.subtitle_label)
        header.addLayout(title_box)
        header.addStretch()
        self._showing_anomalies = False
        if account.is_admin:
            self.anomaly_button = QPushButton("异常数据（0）")
            self.anomaly_button.setObjectName("AnomalyButton")
            self.anomaly_button.setCheckable(True)
            self.anomaly_button.setToolTip("查看空单元格异常的数量和来源站点")
            self.anomaly_button.clicked.connect(self._toggle_anomaly_view)
            self.anomaly_button.hide()
            header.addWidget(self.anomaly_button)
        self.export_btn = QPushButton("导出 Excel")
        self.export_btn.setObjectName("PrimaryButton")
        self.export_btn.setToolTip("按当前站点、数据分类和日期范围导出表格数据")
        self.export_btn.clicked.connect(self._export_dashboard_excel)
        header.addWidget(self.export_btn)
        refresh = QPushButton("刷新统计")
        refresh.clicked.connect(self._refresh_with_loading)
        header.addWidget(refresh)
        layout.addLayout(header)

        if dependency_warning:
            warning = QLabel(dependency_warning)
            warning.setWordWrap(True)
            warning.setStyleSheet(
                "color:#9a6415;background:#fff7e5;border:1px solid #f4d79b;"
                "padding:10px;border-radius:8px;"
            )
            layout.addWidget(warning)

        self.filter_card = QFrame()
        self.filter_card.setObjectName("Card")
        self.filter_layout = QHBoxLayout(self.filter_card)
        self.filter_layout.setContentsMargins(12, 8, 12, 8)
        self.filter_layout.setSpacing(10)
        self.station_combo = QComboBox()
        self.station_combo.setMinimumWidth(210)
        self.category_combo = QComboBox()
        self.category_combo.setMinimumWidth(180)
        self.category_combo.addItem("各站分布", "station_distribution")
        self.category_combo.addItem("用时效率", "timing")
        self.category_combo.addItem("按完成类型", "completion")
        self.category_combo.addItem("按违规类型", "violation")
        self.date_range_selector = DateRangeSelector()
        (
            self.station_filter_group,
            self.station_filter_label,
        ) = self._create_filter_group("站点", self.station_combo)
        (
            self.category_filter_group,
            self.category_filter_label,
        ) = self._create_filter_group("数据分类", self.category_combo)
        (
            self.date_filter_group,
            self.date_filter_label,
        ) = self._create_filter_group(
            "日期范围",
            self.date_range_selector,
        )
        self.filter_layout.addWidget(self.station_filter_group)
        self.filter_layout.addWidget(self.category_filter_group)
        self.filter_layout.addWidget(self.date_filter_group)
        self.filter_layout.addStretch(1)
        self.filter_layout.activate()
        self.filter_card.setMinimumWidth(
            self.filter_layout.minimumSize().width()
        )
        layout.addWidget(self.filter_card)

        self.timing_section = QWidget()
        timing_section_layout = QVBoxLayout(self.timing_section)
        timing_section_layout.setContentsMargins(0, 0, 0, 0)
        timing_section_layout.setSpacing(8)
        timing_header = QHBoxLayout()
        timing_header.setContentsMargins(2, 1, 2, 0)
        timing_title = QLabel("效率统计")
        timing_title.setStyleSheet("color:#173a3d;font-size:14px;font-weight:700;")
        timing_header.addWidget(timing_title)
        self.timing_scope_label = QLabel("仅统计三步全部完成的数据")
        self.timing_scope_label.setObjectName("Muted")
        timing_header.addWidget(self.timing_scope_label)
        timing_header.addStretch(1)
        self.timing_basis_button = QPushButton("总用时口径")
        self.timing_basis_button.setObjectName("TimingBasisToggleButton")
        self.timing_basis_button.setCheckable(True)
        self.timing_basis_button.setAccessibleName("切换总用时和有效用时")
        self.timing_basis_button.setToolTip("点击切换为有效用时口径")
        self.timing_basis_button.toggled.connect(self._toggle_timing_basis)
        timing_header.addWidget(self.timing_basis_button)
        timing_section_layout.addLayout(timing_header)

        self.timing_kpi_layout = QGridLayout()
        self.timing_kpi_layout.setSpacing(10)
        self.timing_kpi_cards = {}
        timing_section_layout.addLayout(self.timing_kpi_layout)
        layout.addWidget(self.timing_section)

        self._active_category = self.category_combo.currentData()
        self._station_before_distribution = None
        self._has_saved_station_before_distribution = False
        self._populate_station_options()
        if self._active_category == "station_distribution":
            self._station_before_distribution = self.station_combo.currentData()
            self._has_saved_station_before_distribution = True
        self.station_combo.currentIndexChanged.connect(self.refresh)
        self.category_combo.currentIndexChanged.connect(self._on_category_changed)
        self.date_range_selector.range_changed.connect(self.refresh)

        self.kpi_layout = QGridLayout()
        self.kpi_layout.setSpacing(10)
        layout.addLayout(self.kpi_layout)

        self.detail_tabs = QTabWidget()
        self.detail_tabs.setObjectName("DashboardTabs")
        self.detail_tabs.setDocumentMode(False)
        self.detail_tabs.setMinimumHeight(350)

        self.chart_tab = QWidget()
        chart_layout = QVBoxLayout(self.chart_tab)
        chart_layout.setContentsMargins(12, 12, 12, 12)
        chart_layout.setSpacing(0)
        self.distribution_chart = WorkflowDistributionChart()
        self.violation_chart = ViolationReasonChart()
        self.station_distribution_chart = StationDistributionChart()
        for chart in (
            self.distribution_chart,
            self.violation_chart,
            self.station_distribution_chart,
        ):
            chart.setFixedHeight(DISTRIBUTION_CHART_HEIGHT)
        self.violation_mode_button = self.violation_chart.mode_button
        chart_layout.addWidget(self.distribution_chart)
        chart_layout.addWidget(self.violation_chart)
        chart_layout.addWidget(self.station_distribution_chart)
        chart_layout.addStretch(1)
        self.detail_tabs.addTab(self.chart_tab, "图表分析")

        self.data_tab = QWidget()
        data_layout = QVBoxLayout(self.data_tab)
        data_layout.setContentsMargins(14, 14, 14, 14)
        data_layout.setSpacing(12)

        data_header = QFrame()
        data_header.setObjectName("DataSummaryHeader")
        data_header.setStyleSheet(
            """
            QFrame#DataSummaryHeader {
                background: #f5f8f7;
                border: 1px solid #d8e7e3;
                border-radius: 10px;
            }
            QLabel#DataTitle {
                color: #173a3d;
                font-size: 16px;
                font-weight: 700;
                border: none;
                background: transparent;
            }
            QLabel#DataDescription {
                color: #647c7b;
                font-size: 12px;
                border: none;
                background: transparent;
            }
            QLabel#DataCountBadge {
                color: #176f68;
                background: #e7f6f2;
                border: 1px solid #bfe1d9;
                border-radius: 12px;
                padding: 4px 11px;
                font-weight: 700;
            }
            """
        )
        data_header_layout = QHBoxLayout(data_header)
        data_header_layout.setContentsMargins(16, 11, 16, 11)
        data_header_layout.setSpacing(12)
        data_title_layout = QVBoxLayout()
        data_title_layout.setContentsMargins(0, 0, 0, 0)
        data_title_layout.setSpacing(3)
        self.data_title = QLabel("完整数据")
        self.data_title.setObjectName("DataTitle")
        self.data_description = QLabel("按当前筛选条件展示详细统计")
        self.data_description.setObjectName("DataDescription")
        data_title_layout.addWidget(self.data_title)
        data_title_layout.addWidget(self.data_description)
        data_header_layout.addLayout(data_title_layout)
        data_header_layout.addStretch()
        self.data_count_label = QLabel("0 行")
        self.data_count_label.setObjectName("DataCountBadge")
        self.data_count_label.setAlignment(Qt.AlignCenter)
        data_header_layout.addWidget(self.data_count_label)
        data_layout.addWidget(data_header)

        self.summary_table = QTableWidget()
        self.summary_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.summary_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.summary_table.setSelectionMode(QAbstractItemView.SingleSelection)
        self.summary_table.setAlternatingRowColors(True)
        self.summary_table.setShowGrid(False)
        self.summary_table.setWordWrap(False)
        self.summary_table.setHorizontalScrollMode(QAbstractItemView.ScrollPerPixel)
        self.summary_table.setVerticalScrollMode(QAbstractItemView.ScrollPerPixel)
        self.summary_table.verticalHeader().setVisible(False)
        self.summary_table.verticalHeader().setDefaultSectionSize(42)
        self.summary_table.horizontalHeader().setMinimumHeight(42)
        self.summary_table.horizontalHeader().setMinimumSectionSize(80)
        self.summary_table.setStyleSheet(
            """
            QTableWidget {
                background: #ffffff;
                alternate-background-color: #f7faf9;
                border: 1px solid #d8e7e3;
                border-radius: 10px;
                color: #264b4c;
                selection-background-color: #e1f2ee;
                selection-color: #145f58;
            }
            QTableWidget::item {
                border-bottom: 1px solid #e8f1ef;
                padding: 7px 10px;
            }
            QHeaderView::section {
                background: #eef5f3;
                color: #536e6c;
                border: none;
                border-bottom: 1px solid #d8e7e3;
                padding: 8px 10px;
                font-weight: 700;
            }
            QTableCornerButton::section {
                background: #eef5f3;
                border: none;
                border-bottom: 1px solid #d8e7e3;
            }
            """
        )
        data_layout.addWidget(self.summary_table)
        self.detail_tabs.addTab(self.data_tab, "完整数据")

        layout.addWidget(self.detail_tabs, 1)
        self.refresh()

    @property
    def navigation_view(self):
        if self._showing_anomalies:
            return "anomaly"
        return self.category_combo.currentData() or "station_distribution"

    def _refresh_with_loading(self, _checked=False):
        return run_ui_with_loading(self, "正在加载统计数据…", self.refresh)

    def set_sidebar_navigation(self, enabled=True):
        """由主窗口的数据中心侧边栏接管详细统计分类。"""
        self._sidebar_navigation = bool(enabled)
        self.category_filter_group.setVisible(not self._sidebar_navigation)
        anomaly_button = getattr(self, "anomaly_button", None)
        if anomaly_button is not None:
            anomaly_button.setVisible(
                self.category_combo.currentData() == "completion"
            )
        self._sync_view_header()
        self.refresh()

    def set_navigation_view(self, view):
        if view not in self.NAVIGATION_VIEWS:
            raise ValueError(f"未知的数据视图：{view}")
        if view == "anomaly" and not self.account.is_admin:
            raise PermissionError("普通用户无权查看异常数据")

        category = "completion" if view == "anomaly" else view
        category_index = self.category_combo.findData(category)
        if category_index < 0:
            raise ValueError(f"数据分类不可用：{category}")
        if self.category_combo.currentIndex() != category_index:
            self.category_combo.setCurrentIndex(category_index)

        self._showing_anomalies = view == "anomaly"
        anomaly_button = getattr(self, "anomaly_button", None)
        if anomaly_button is not None:
            anomaly_button.setChecked(self._showing_anomalies)
        self.detail_tabs.setCurrentWidget(
            self.data_tab if self._showing_anomalies else self.chart_tab
        )
        self._sync_view_header()
        self.refresh()

    def _sync_view_header(self):
        title, subtitle = self.NAVIGATION_VIEWS[self.navigation_view]
        self.title_label.setText(title)
        self.subtitle_label.setText(subtitle)

    @staticmethod
    def _create_filter_group(label, control):
        group = QFrame()
        group.setObjectName("DashboardFilterGroup")
        group_layout = QHBoxLayout(group)
        group_layout.setContentsMargins(10, 6, 10, 6)
        group_layout.setSpacing(8)
        label_widget = QLabel(label)
        label_widget.setObjectName("DashboardFilterLabel")
        group_layout.addWidget(label_widget)
        group_layout.addWidget(control)
        return group, label_widget

    def _populate_station_options(self):
        had_options = self.station_combo.count() > 0
        current_id = self.station_combo.currentData()
        restricted_online_scope = (
            self.account.server_account_id is not None
            and not self.account.can_view_all_stats
        )
        hidden_test_scope = self.account.is_test and not self.account.can_view_all_stats
        self.station_combo.blockSignals(True)
        self.station_combo.clear()
        self.station_combo.addItem(
            "不参与统计" if hidden_test_scope else "全部站点",
            self.account.id if hidden_test_scope else None,
        )
        order = {username: index for index, (_, username) in enumerate(DEFAULT_STATION_USERS)}
        accounts = [
            account
            for account in self.database.list_accounts()
            if not account.is_admin and not account.is_test
        ]
        if restricted_online_scope:
            accounts = [
                account for account in accounts if account.id == self.account.id
            ]
        accounts.sort(key=lambda item: (order.get(item.username, 999), item.name_label))
        for station in accounts:
            self.station_combo.addItem(station.name_label, station.id)
            self.station_combo.setItemData(
                self.station_combo.count() - 1,
                f"账号：{station.username}",
                Qt.ToolTipRole,
            )
        target_id = current_id if had_options else (
            None if self.account.can_view_all_stats else self.account.id
        )
        index = self.station_combo.findData(target_id)
        self.station_combo.setCurrentIndex(index if index >= 0 else 0)
        self.station_combo.setEnabled(not restricted_online_scope)
        self.station_combo.blockSignals(False)

    def _selected_user_id(self):
        return self.station_combo.currentData()

    def _selected_scope(self):
        scope = self.station_combo.currentText()
        if not self.date_range_selector.all_dates_check.isChecked():
            scope = f"{scope} · {self.date_range_selector.range_label()}"
        return scope

    def _date_range(self):
        return self.date_range_selector.date_range()

    def _export_station_accounts(self):
        default_order = {
            username: index
            for index, (_, username) in enumerate(DEFAULT_STATION_USERS)
        }
        accounts = [
            account
            for account in self.database.list_accounts()
            if not account.is_admin and not account.is_test
        ]
        accounts.sort(
            key=lambda account: (
                default_order.get(account.username, 999),
                account.name_label,
            )
        )
        return accounts

    @staticmethod
    def _excel_value(text):
        value = str(text).strip()
        if value.endswith("%"):
            try:
                return float(value[:-1]) / 100, "0.0%"
            except ValueError:
                return value, None
        if re.fullmatch(r"[-+]?\d+", value):
            return int(value), "0"
        if re.fullmatch(r"[-+]?(?:\d+\.\d*|\d*\.\d+)", value):
            return float(value), "0.00"
        return value, None

    def _timing_detail_rows(self, user_id):
        start_date, end_date = self._date_range()
        totals = self.database.get_workflow_timing_totals(
            user_id,
            users_only=(user_id is None),
            start_date=start_date,
            end_date=end_date,
        )
        completed_items = int(totals["completed_items"])
        basis_key = "active_ms" if self._timing_basis == "active" else "total_ms"
        basis_label = "有效用时" if self._timing_basis == "active" else "总用时"
        basis_ms = totals[basis_key]
        average_basis_ms = (
            basis_ms / completed_items
            if completed_items and basis_ms > 0
            else None
        )
        throughput_per_hour = (
            completed_items * MILLISECONDS_PER_HOUR / basis_ms
            if completed_items and basis_ms > 0
            else None
        )
        manual_estimated_ms = (
            completed_items
            * self.HUMAN_BASELINE_MS
            / self.HUMAN_BASELINE_ITEMS
        )
        efficiency_gain = (
            (manual_estimated_ms - basis_ms)
            / manual_estimated_ms
            * 100
            if manual_estimated_ms
            else None
        )
        rows = [
            (
                basis_label,
                format_precise_duration(basis_ms),
                (
                    "实际运行和重试耗时，已剔除暂停等待"
                    if self._timing_basis == "active"
                    else "有效用时＋暂停等待"
                ),
            ),
            (
                "暂停等待",
                format_precise_duration(totals["paused_ms"]),
                "登录、验证、浏览器加载及人工等待",
            ),
            ("完成数据", f"{completed_items} 条", "仅统计三步全部成功的批次"),
            (
                "平均处理效率",
                (
                    f"{throughput_per_hour:.0f} 条/小时"
                    if throughput_per_hour is not None
                    else "—"
                ),
                f"完成数据÷{basis_label}",
            ),
            (
                "较纯人工效率提升",
                f"{efficiency_gain:.1f}%" if efficiency_gain is not None else "—",
                "正值表示效率提升，负值表示低于参考效率",
            ),
        ]
        return totals, rows

    def _station_export_dataset(self, account):
        start_date, end_date = self._date_range()
        category = self.category_combo.currentData()

        if self._showing_anomalies:
            rows = self._get_anomaly_rows(account.id)
            return (
                ["用户（站）", "登录账号", "异常条数", "本站总计数", "异常占比"],
                [
                    [
                        row["station"],
                        row["username"],
                        row["empty"],
                        row["total"],
                        f'{row["empty_rate"]:.1f}%',
                    ]
                    for row in rows
                ],
            )

        if category == "timing":
            _, rows = self._timing_detail_rows(account.id)
            return ["指标", "数值", "统计说明"], rows

        if category == "violation":
            rows = self.database.get_violation_totals(
                account.id,
                start_date=start_date,
                end_date=end_date,
            )
            values = []
            for row in rows:
                percent = (
                    row["has_phone"] / row["total"] * 100
                    if row["total"]
                    else 0
                )
                values.append(
                    [
                        row["reason"],
                        row["total"],
                        row["has_phone"],
                        row["other"],
                        f"{percent:.1f}%",
                    ]
                )
            return (
                ["违规原因", "总计", "有电话", "其他数据", "有电话占比"],
                values,
            )

        totals_by_metric = {
            row["metric_key"]: int(row["total"])
            for row in self.database.get_user_totals(
                account.id,
                start_date,
                end_date,
            )
        }
        if category == "station_distribution":
            total = totals_by_metric.get(WORKFLOW_TOTAL_METRIC, 0)
            has_phone = totals_by_metric.get(WORKFLOW_HAS_PHONE_METRIC, 0)
            no_phone = totals_by_metric.get(
                WORKFLOW_NO_PHONE_METRIC,
                max(0, total - has_phone),
            )
            timing, _ = self._timing_detail_rows(account.id)
            phone_rate = has_phone / total * 100 if total else 0
            return (
                ["指标", "数值", "单位 / 统计说明"],
                [
                    ("总计数", total, "条"),
                    ("有电话数", has_phone, "条"),
                    ("无电话数", no_phone, "条"),
                    ("有电话占比", f"{phone_rate:.1f}%", "有电话数÷总计数"),
                    (
                        "总用时",
                        format_precise_duration(timing["total_ms"]),
                        "有效用时＋暂停等待",
                    ),
                    (
                        "有效用时",
                        format_precise_duration(timing["active_ms"]),
                        "实际运行和重试耗时",
                    ),
                    (
                        "暂停等待",
                        format_precise_duration(timing["paused_ms"]),
                        "登录、验证及人工等待",
                    ),
                    ("完成数据", timing["completed_items"], "条"),
                ],
            )

        metrics = self._ordered_completion_metrics(
            self.database.get_metric_definitions()
        )
        return (
            ["完成类型", "累计数量", "单位"],
            [
                [
                    metric["label"],
                    totals_by_metric.get(metric["metric_key"], 0),
                    metric["unit"],
                ]
                for metric in metrics
            ],
        )

    @staticmethod
    def _unique_excel_sheet_title(workbook, preferred):
        base = re.sub(r'[\\/*?:\[\]]+', "_", str(preferred)).strip().strip("'")
        base = (base or "站点")[:31]
        title = base
        suffix_index = 2
        while title in workbook.sheetnames:
            suffix = f"_{suffix_index}"
            title = f"{base[:31 - len(suffix)]}{suffix}"
            suffix_index += 1
        return title

    def _append_station_export_sheet(
        self,
        workbook,
        account,
        category_name,
        range_label,
    ):
        headers, rows = self._station_export_dataset(account)
        sheet = workbook.create_sheet(
            self._unique_excel_sheet_title(workbook, account.name_label)
        )
        sheet.sheet_view.showGridLines = False
        final_column = max(6, len(headers))
        sheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=final_column,
        )
        title_cell = sheet.cell(
            1,
            1,
            f"{account.name_label} · {category_name}",
        )
        title_cell.font = Font(color="FFFFFF", bold=True, size=14)
        title_cell.fill = PatternFill("solid", fgColor="1C5ED6")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 28

        metadata = (
            ("站名", account.name_label),
            ("登录账号", account.username),
            ("时间范围", range_label),
        )
        for index, (label, value) in enumerate(metadata):
            label_cell = sheet.cell(2, index * 2 + 1, label)
            value_cell = sheet.cell(2, index * 2 + 2, value)
            label_cell.font = Font(bold=True, color="526177")
            label_cell.fill = PatternFill("solid", fgColor="EDF4FF")
            value_cell.fill = PatternFill("solid", fgColor="F8FAFD")

        header_row = 4
        header_fill = PatternFill("solid", fgColor="3478F6")
        header_font = Font(color="FFFFFF", bold=True)
        thin_border = Border(
            left=Side(style="thin", color="DCE4EF"),
            right=Side(style="thin", color="DCE4EF"),
            top=Side(style="thin", color="DCE4EF"),
            bottom=Side(style="thin", color="DCE4EF"),
        )
        for column, header_text in enumerate(headers, 1):
            cell = sheet.cell(header_row, column, header_text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for source_row, row_values in enumerate(rows):
            target_row = header_row + source_row + 1
            for column, raw_value in enumerate(row_values, 1):
                value, number_format = self._excel_value(raw_value)
                cell = sheet.cell(target_row, column, value)
                if number_format:
                    cell.number_format = number_format
                cell.border = thin_border
                cell.alignment = Alignment(
                    horizontal="center" if column > 1 else "left",
                    vertical="center",
                )
                if source_row % 2:
                    cell.fill = PatternFill("solid", fgColor="F7F9FD")

        last_row = max(header_row, sheet.max_row)
        last_column_letter = get_column_letter(len(headers))
        sheet.auto_filter.ref = (
            f"A{header_row}:{last_column_letter}{last_row}"
        )
        sheet.freeze_panes = f"A{header_row + 1}"
        for column in range(1, final_column + 1):
            letter = get_column_letter(column)
            content_width = max(
                len(str(sheet.cell(row, column).value or ""))
                for row in range(1, sheet.max_row + 1)
            )
            sheet.column_dimensions[letter].width = min(
                40,
                max(12, content_width + 3),
            )
        return sheet

    def _save_dashboard_excel(self, file_path):
        """Export the current summary and all-station detail sheets."""
        station_name = self.station_combo.currentText() or "全部站点"
        category_name = (
            "异常数据"
            if self._showing_anomalies
            else self.category_combo.currentText()
        )
        range_label = self.date_range_selector.range_label()
        table_headers = [
            self.summary_table.horizontalHeaderItem(column).text()
            for column in range(self.summary_table.columnCount())
        ]
        headers = ["站名", "时间范围", *table_headers]

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "仪表盘数据"
        sheet.sheet_view.showGridLines = False

        final_column = max(6, len(headers))
        sheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=final_column,
        )
        title_cell = sheet.cell(1, 1, "数据仪表盘导出")
        title_cell.font = Font(color="FFFFFF", bold=True, size=15)
        title_cell.fill = PatternFill("solid", fgColor="1C5ED6")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 30

        metadata = (
            ("站名", station_name),
            ("数据分类", category_name),
            ("时间范围", range_label),
        )
        for index, (label, value) in enumerate(metadata):
            label_cell = sheet.cell(2, index * 2 + 1, label)
            value_cell = sheet.cell(2, index * 2 + 2, value)
            label_cell.font = Font(bold=True, color="526177")
            label_cell.fill = PatternFill("solid", fgColor="EDF4FF")
            value_cell.fill = PatternFill("solid", fgColor="F8FAFD")
            for cell in (label_cell, value_cell):
                cell.alignment = Alignment(vertical="center")

        sheet.cell(3, 1, "导出时间").font = Font(bold=True, color="526177")
        sheet.cell(3, 2, datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

        header_row = 5
        header_fill = PatternFill("solid", fgColor="3478F6")
        header_font = Font(color="FFFFFF", bold=True)
        thin_border = Border(
            left=Side(style="thin", color="DCE4EF"),
            right=Side(style="thin", color="DCE4EF"),
            top=Side(style="thin", color="DCE4EF"),
            bottom=Side(style="thin", color="DCE4EF"),
        )
        for column, header_text in enumerate(headers, 1):
            cell = sheet.cell(header_row, column, header_text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        sheet.row_dimensions[header_row].height = 25

        for source_row in range(self.summary_table.rowCount()):
            target_row = header_row + source_row + 1
            row_values = [station_name, range_label]
            row_values.extend(
                self.summary_table.item(source_row, column).text()
                if self.summary_table.item(source_row, column) is not None
                else ""
                for column in range(self.summary_table.columnCount())
            )
            for column, raw_value in enumerate(row_values, 1):
                value, number_format = self._excel_value(raw_value)
                cell = sheet.cell(target_row, column, value)
                if number_format:
                    cell.number_format = number_format
                cell.border = thin_border
                cell.alignment = Alignment(
                    horizontal="center" if column > 2 else "left",
                    vertical="center",
                )
                if source_row % 2:
                    cell.fill = PatternFill("solid", fgColor="F7F9FD")

        last_row = max(header_row, sheet.max_row)
        last_column_letter = get_column_letter(len(headers))
        sheet.auto_filter.ref = f"A{header_row}:{last_column_letter}{last_row}"
        sheet.freeze_panes = f"A{header_row + 1}"

        for column in range(1, final_column + 1):
            letter = get_column_letter(column)
            content_width = max(
                len(str(sheet.cell(row, column).value or ""))
                for row in range(1, sheet.max_row + 1)
            )
            sheet.column_dimensions[letter].width = min(
                36,
                max(12, content_width + 3),
            )
        sheet.column_dimensions["B"].width = max(
            sheet.column_dimensions["B"].width,
            25,
        )
        sheet.column_dimensions["F"].width = max(
            sheet.column_dimensions["F"].width,
            25,
        )

        station_sheet_count = 0
        if self.station_combo.currentData() is None:
            for station in self._export_station_accounts():
                self._append_station_export_sheet(
                    workbook,
                    station,
                    category_name,
                    range_label,
                )
                station_sheet_count += 1

        target = Path(file_path)
        workbook.save(target)
        return {
            "file_path": str(target),
            "row_count": self.summary_table.rowCount(),
            "station_name": station_name,
            "category_name": category_name,
            "range_label": range_label,
            "station_sheet_count": station_sheet_count,
        }

    def _export_dashboard_excel(self):
        start_date, end_date = self._date_range()
        range_suffix = (
            "all"
            if start_date is None
            else f"{start_date:%Y%m%d}-{end_date:%Y%m%d}"
        )
        station_part = re.sub(
            r'[\\/:*?"<>|]+',
            "_",
            self.station_combo.currentText() or "全部站点",
        )
        default_name = (
            f"仪表盘_{station_part}_{range_suffix}_"
            f"{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        )
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "导出数据仪表盘",
            default_name,
            "Excel 工作簿 (*.xlsx)",
        )
        if not file_path:
            return
        target = Path(file_path)
        if target.suffix.lower() != ".xlsx":
            target = target.with_suffix(".xlsx")
        try:
            result = self._save_dashboard_excel(target)
        except (OSError, ValueError) as exc:
            QMessageBox.warning(self, "导出失败", str(exc))
            return
        station_sheet_message = (
            f"另附 {result['station_sheet_count']} 个站点子表。\n"
            if result["station_sheet_count"]
            else ""
        )
        QMessageBox.information(
            self,
            "导出成功",
            f"站点：{result['station_name']}\n"
            f"时间范围：{result['range_label']}\n"
            f"已导出 {result['row_count']} 行数据。\n"
            f"{station_sheet_message}"
            f"{result['file_path']}",
        )

    def _on_category_changed(self, _index):
        category = self.category_combo.currentData()
        if category != "completion":
            self._showing_anomalies = False
            anomaly_button = getattr(self, "anomaly_button", None)
            if anomaly_button is not None:
                anomaly_button.setChecked(False)
        if category == "station_distribution" and self._active_category != category:
            self._station_before_distribution = self.station_combo.currentData()
            self._has_saved_station_before_distribution = True
        elif (
            self._active_category == "station_distribution"
            and category != "station_distribution"
            and self._has_saved_station_before_distribution
        ):
            station_index = self.station_combo.findData(
                self._station_before_distribution
            )
            self.station_combo.blockSignals(True)
            self.station_combo.setCurrentIndex(station_index if station_index >= 0 else 0)
            self.station_combo.blockSignals(False)
            self._has_saved_station_before_distribution = False
        self._active_category = category
        self._sync_view_header()
        self.refresh()

    def _toggle_anomaly_view(self, checked):
        if not self.account.is_admin:
            return
        self._showing_anomalies = bool(checked)
        self._sync_view_header()
        if self._showing_anomalies:
            self.detail_tabs.setCurrentWidget(self.data_tab)
        self.refresh()
        if not self._showing_anomalies:
            self.detail_tabs.setCurrentWidget(self.chart_tab)

    def _clear_kpis(self):
        while self.kpi_layout.count():
            item = self.kpi_layout.takeAt(0)
            if item.widget():
                item.widget().hide()
                item.widget().deleteLater()

    def _toggle_timing_basis(self, checked):
        self._timing_basis = "active" if checked else "total"
        basis_label = "有效用时" if checked else "总用时"
        next_label = "总用时" if checked else "有效用时"
        self.timing_basis_button.setText(f"{basis_label}口径")
        self.timing_basis_button.setToolTip(f"点击切换为{next_label}口径")
        self.refresh()

    @staticmethod
    def _metric_card(
        label,
        value,
        unit,
        hover_details=None,
        hover_accent="#1d8178",
    ):
        card = (
            TimingMetricCard(
                f"{label}明细",
                hover_details,
                hover_accent,
            )
            if hover_details is not None
            else QFrame()
        )
        card.setObjectName("Card")
        box = QVBoxLayout(card)
        box.setContentsMargins(14, 9, 14, 9)
        caption = QLabel(label)
        caption.setObjectName("Muted")
        number = QLabel(f"{value} {unit}".rstrip())
        number.setObjectName("MetricValue")
        if hover_details is not None:
            caption.setAttribute(Qt.WA_TransparentForMouseEvents, True)
            number.setAttribute(Qt.WA_TransparentForMouseEvents, True)
        box.addWidget(caption)
        box.addWidget(number)
        return card

    def _refresh_timing_kpis(self, user_id):
        while self.timing_kpi_layout.count():
            item = self.timing_kpi_layout.takeAt(0)
            if item.widget():
                item.widget().hide()
                item.widget().deleteLater()
        self.timing_kpi_cards = {}

        start_date, end_date = self._date_range()
        totals = self.database.get_workflow_timing_totals(
            user_id,
            users_only=(user_id is None),
            start_date=start_date,
            end_date=end_date,
        )
        completed_items = int(totals["completed_items"])
        basis_key = "active_ms" if self._timing_basis == "active" else "total_ms"
        basis_label = "有效用时" if self._timing_basis == "active" else "总用时"
        basis_ms = totals[basis_key]
        average_basis_ms = (
            basis_ms / completed_items
            if completed_items and basis_ms > 0
            else None
        )
        throughput_per_hour = (
            completed_items * MILLISECONDS_PER_HOUR / basis_ms
            if completed_items and basis_ms > 0
            else None
        )
        manual_estimated_ms = (
            completed_items
            * self.HUMAN_BASELINE_MS
            / self.HUMAN_BASELINE_ITEMS
        )
        efficiency_gain = (
            (manual_estimated_ms - basis_ms)
            / manual_estimated_ms
            * 100
            if manual_estimated_ms
            else None
        )
        completed_text = f"{completed_items} 条"
        precise_basis = format_precise_duration(basis_ms)
        precise_paused = format_precise_duration(totals["paused_ms"])
        seconds_per_item = (
            average_basis_ms / 1000
            if average_basis_ms is not None
            else None
        )
        precise_average = (
            format_precise_duration(average_basis_ms)
            if average_basis_ms is not None
            else "暂无数据"
        )
        values = [
            (
                basis_label,
                f"{basis_ms / MILLISECONDS_PER_HOUR:.1f}",
                "小时",
                [
                    (f"精确{basis_label}", precise_basis),
                    ("暂停等待", precise_paused),
                    ("完成数据", completed_text),
                ],
                "#1d8178",
            ),
            (
                "平均处理效率",
                (
                    f"{throughput_per_hour:.0f}"
                    if throughput_per_hour is not None
                    else "—"
                ),
                "条/小时" if throughput_per_hour is not None else "",
                [
                    ("计量口径", basis_label),
                    (
                        "平均耗时",
                        f"{seconds_per_item:.1f} 秒/条"
                        if seconds_per_item is not None
                        else "暂无数据",
                    ),
                    ("精确平均", precise_average),
                    ("完成数据", completed_text),
                ],
                "#d39a2c",
            ),
            (
                "较纯人工效率提升",
                f"{efficiency_gain:.1f}" if efficiency_gain is not None else "—",
                "%" if efficiency_gain is not None else "",
                [
                    ("计量口径", basis_label),
                    (basis_label, precise_basis),
                    (
                        "人工预计用时",
                        format_precise_duration(manual_estimated_ms),
                    ),
                ],
                "#4a8bc4",
            ),
        ]
        for column in range(4):
            self.timing_kpi_layout.setColumnStretch(
                column,
                1 if column < len(values) else 0,
            )
        for index, (
            label,
            value,
            unit,
            hover_details,
            hover_accent,
        ) in enumerate(values):
            card = self._metric_card(
                label,
                value,
                unit,
                hover_details=hover_details,
                hover_accent=hover_accent,
            )
            self.timing_kpi_cards[label] = card
            self.timing_kpi_layout.addWidget(card, 0, index)

        self.timing_scope_label.setText(
            f"完成数据：{completed_items} 条 · 当前按{basis_label}计量"
        )

    def _add_kpis(self, values, columns=4):
        for column in range(max(columns, self.kpi_layout.columnCount())):
            self.kpi_layout.setColumnStretch(
                column,
                1 if column < columns else 0,
            )
        for index, (label, value, unit) in enumerate(values):
            self.kpi_layout.addWidget(
                self._metric_card(label, value, unit),
                index // columns,
                index % columns,
            )

    def _finish_summary_table(self, scope, detail, stretch_columns=(0,)):
        self.data_description.setText(f"{scope} · {detail}")
        self.data_count_label.setText(f"{self.summary_table.rowCount()} 行")
        make_table_columns_resizable(
            self.summary_table,
            {
                column: 240 if column in stretch_columns else 130
                for column in range(self.summary_table.columnCount())
            },
            minimum_width=80,
        )
        self.summary_table.scrollToTop()

    @staticmethod
    def _ordered_completion_metrics(metrics):
        """Keep the aggregate first, followed by the five result types."""
        metrics_by_key = {metric["metric_key"]: metric for metric in metrics}
        ordered_keys = [WORKFLOW_TOTAL_METRIC]
        ordered_keys.extend(
            metric_key
            for metric_key, _label, _color in WorkflowDistributionChart.SEGMENTS
        )
        return [
            metrics_by_key[metric_key]
            for metric_key in ordered_keys
            if metric_key in metrics_by_key
        ]

    def _get_station_distribution_rows(self):
        stations = {}
        start_date, end_date = self._date_range()
        restricted_user_id = (
            self.account.id
            if self.account.server_account_id is not None
            and not self.account.can_view_all_stats
            else None
        )
        for row in self.database.get_all_account_totals(start_date, end_date):
            if row["role"] != "user":
                continue
            if restricted_user_id is not None and row["user_id"] != restricted_user_id:
                continue
            station = stations.setdefault(
                row["user_id"],
                {
                    "user_id": row["user_id"],
                    "station": row["display_name"] or row["username"],
                    "username": row["username"],
                    "total": 0,
                    "has_phone": 0,
                },
            )
            if row["metric_key"] == WORKFLOW_TOTAL_METRIC:
                station["total"] = int(row["total"])
            elif row["metric_key"] == WORKFLOW_HAS_PHONE_METRIC:
                station["has_phone"] = int(row["total"])

        rows = list(stations.values())
        for row in rows:
            timing = self.database.get_workflow_timing_totals(
                row["user_id"],
                start_date=start_date,
                end_date=end_date,
            )
            row["total_time_ms"] = timing["total_ms"]
            row["active_ms"] = timing["active_ms"]
        all_total = sum(row["total"] for row in rows)
        all_has_phone = sum(row["has_phone"] for row in rows)
        all_total_time_ms = sum(row["total_time_ms"] for row in rows)
        all_active_ms = sum(row["active_ms"] for row in rows)
        for row in rows:
            row["total_share"] = row["total"] / all_total * 100 if all_total else 0
            row["phone_share"] = (
                row["has_phone"] / all_has_phone * 100 if all_has_phone else 0
            )
            row["total_time_share"] = (
                row["total_time_ms"] / all_total_time_ms * 100
                if all_total_time_ms
                else 0
            )
            row["active_time_share"] = (
                row["active_ms"] / all_active_ms * 100
                if all_active_ms
                else 0
            )
        rows.sort(
            key=lambda row: (
                -int(row["total"]),
                -int(row["has_phone"]),
                str(row["station"]).casefold(),
            )
        )
        return rows

    def _get_anomaly_rows(self, user_id=None):
        """Return empty-cell anomaly totals grouped by ordinary user station."""
        default_order = {
            username: index for index, (_, username) in enumerate(DEFAULT_STATION_USERS)
        }
        stations = {}
        start_date, end_date = self._date_range()
        for row in self.database.get_all_account_totals(start_date, end_date):
            if row["role"] != "user":
                continue
            if user_id is not None and row["user_id"] != user_id:
                continue
            station = stations.setdefault(
                row["user_id"],
                {
                    "user_id": row["user_id"],
                    "station": row["display_name"] or row["username"],
                    "username": row["username"],
                    "total": 0,
                    "empty": 0,
                },
            )
            if row["metric_key"] == WORKFLOW_TOTAL_METRIC:
                station["total"] = int(row["total"])
            elif row["metric_key"] == WORKFLOW_EMPTY_METRIC:
                station["empty"] = int(row["total"])

        rows = [row for row in stations.values() if row["empty"] > 0]
        rows.sort(
            key=lambda row: (
                default_order.get(row["username"], 999),
                row["station"],
            )
        )
        for row in rows:
            row["empty_rate"] = (
                row["empty"] / row["total"] * 100 if row["total"] else 0
            )
        return rows

    def _render_anomalies(self, user_id, scope, rows=None):
        rows = self._get_anomaly_rows(user_id) if rows is None else rows
        self.distribution_chart.hide()
        self.violation_chart.hide()
        self.station_distribution_chart.hide()
        self.detail_tabs.setTabEnabled(0, False)
        self.detail_tabs.setCurrentWidget(self.data_tab)
        self.data_title.setText("异常数据")

        anomaly_total = sum(row["empty"] for row in rows)
        workflow_total = sum(row["total"] for row in rows)
        anomaly_rate = anomaly_total / workflow_total * 100 if workflow_total else 0
        self._add_kpis(
            [
                ("空单元格异常", anomaly_total, "条"),
                ("涉及站点", len(rows), "个"),
                ("涉及站点总数据", workflow_total, "条"),
                ("异常占比", f"{anomaly_rate:.1f}", "%"),
            ]
        )

        headers = ["用户（站）", "登录账号", "异常条数", "本站总计数", "异常占比"]
        self.summary_table.setColumnCount(len(headers))
        self.summary_table.setHorizontalHeaderLabels(headers)
        self.summary_table.setRowCount(len(rows))
        for row_index, row in enumerate(rows):
            values = [
                row["station"],
                row["username"],
                row["empty"],
                row["total"],
                f'{row["empty_rate"]:.1f}%',
            ]
            for column, value in enumerate(values):
                item = QTableWidgetItem(str(value))
                if column >= 2:
                    item.setTextAlignment(Qt.AlignCenter)
                if column == 2:
                    font = item.font()
                    font.setBold(True)
                    item.setFont(font)
                    item.setForeground(QColor("#d17a22"))
                elif column == 4:
                    item.setForeground(QColor("#d18400"))
                self.summary_table.setItem(row_index, column, item)
        self._finish_summary_table(scope, "空单元格异常来源", (0,))

    def _render_timing(self, user_id, scope):
        _, rows = self._timing_detail_rows(user_id)

        self.distribution_chart.hide()
        self.violation_chart.hide()
        chart_was_enabled = self.detail_tabs.isTabEnabled(0)
        if user_id is None:
            self.station_distribution_chart.set_series_mode("timing")
            self.station_distribution_chart.set_timing_basis(self._timing_basis)
            self.station_distribution_chart.set_rows(
                self._get_station_distribution_rows()
            )
            self.station_distribution_chart.show()
            self.detail_tabs.setTabEnabled(0, True)
            if not chart_was_enabled:
                self.detail_tabs.setCurrentWidget(self.chart_tab)
        else:
            self.station_distribution_chart.hide()
            self.detail_tabs.setTabEnabled(0, False)
            self.detail_tabs.setCurrentWidget(self.data_tab)
        self.data_title.setText("用时效率明细")

        headers = ["指标", "数值", "统计说明"]
        self.summary_table.setColumnCount(len(headers))
        self.summary_table.setHorizontalHeaderLabels(headers)
        self.summary_table.setRowCount(len(rows))
        for row_index, values in enumerate(rows):
            for column, value in enumerate(values):
                item = QTableWidgetItem(str(value))
                if column == 1:
                    item.setTextAlignment(Qt.AlignCenter)
                    item.setForeground(QColor("#176f68"))
                    font = item.font()
                    font.setBold(True)
                    item.setFont(font)
                elif column == 2:
                    item.setForeground(QColor("#718096"))
                self.summary_table.setItem(row_index, column, item)
        self._finish_summary_table(scope, "批次计时与效率口径", (0, 2))

    def _render_completion(self, user_id, scope):
        metrics = self._ordered_completion_metrics(
            self.database.get_metric_definitions()
        )
        self.distribution_chart.show()
        self.violation_chart.hide()
        self.station_distribution_chart.hide()
        self.detail_tabs.setTabEnabled(0, True)
        self.data_title.setText("完整数据")
        start_date, end_date = self._date_range()

        if user_id is None:
            totals_by_metric = {metric["metric_key"]: 0 for metric in metrics}
            for row in self.database.get_all_account_totals(start_date, end_date):
                if row["role"] == "user":
                    totals_by_metric[row["metric_key"]] = (
                        totals_by_metric.get(row["metric_key"], 0)
                        + int(row["total"])
                    )
        else:
            totals_by_metric = {
                row["metric_key"]: int(row["total"])
                for row in self.database.get_user_totals(
                    user_id,
                    start_date,
                    end_date,
                )
            }

        self._add_kpis(
            [
                (
                    metric["label"],
                    totals_by_metric.get(metric["metric_key"], 0),
                    metric["unit"],
                )
                for metric in metrics
            ],
            columns=3,
        )
        self.distribution_chart.set_values(totals_by_metric, scope)
        self.summary_table.setColumnCount(3)
        self.summary_table.setHorizontalHeaderLabels(["完成类型", "累计数量", "单位"])
        self.summary_table.setRowCount(len(metrics))
        for row_index, metric in enumerate(metrics):
            for column, value in enumerate(
                [
                    metric["label"],
                    totals_by_metric.get(metric["metric_key"], 0),
                    metric["unit"],
                ]
            ):
                item = QTableWidgetItem(str(value))
                if column == 1:
                    font = item.font()
                    font.setBold(True)
                    item.setFont(font)
                    item.setForeground(QColor("#176f68"))
                    item.setTextAlignment(Qt.AlignCenter)
                elif column == 2:
                    item.setForeground(QColor("#718096"))
                    item.setTextAlignment(Qt.AlignCenter)
                self.summary_table.setItem(row_index, column, item)
        self._finish_summary_table(scope, "完成类型累计明细")

    def _render_violation(self, user_id, scope):
        start_date, end_date = self._date_range()
        rows = self.database.get_violation_totals(
            user_id,
            users_only=(user_id is None),
            start_date=start_date,
            end_date=end_date,
        )
        rows = sorted(
            rows,
            key=lambda row: (
                -int(row.get("total") or 0),
                str(row.get("reason") or ""),
            ),
        )
        self.distribution_chart.hide()
        self.violation_chart.show()
        self.station_distribution_chart.hide()
        self.detail_tabs.setTabEnabled(0, True)
        self.data_title.setText("完整数据")

        total = sum(row["total"] for row in rows)
        has_phone = sum(row["has_phone"] for row in rows)
        other = sum(row["other"] for row in rows)
        self._add_kpis(
            [
                ("违规原因条目", total, "条"),
                ("原因种类", len(rows), "种"),
                ("有电话", has_phone, "条"),
                ("其他数据", other, "条"),
            ]
        )
        self.violation_chart.set_rows(rows, scope)
        headers = ["违规原因", "总计", "有电话", "其他数据", "有电话占比"]
        self.summary_table.setColumnCount(len(headers))
        self.summary_table.setHorizontalHeaderLabels(headers)
        self.summary_table.setRowCount(len(rows))
        for row_index, row in enumerate(rows):
            percent = row["has_phone"] / row["total"] * 100 if row["total"] else 0
            values = [
                row["reason"], row["total"], row["has_phone"], row["other"], f"{percent:.1f}%"
            ]
            for column, value in enumerate(values):
                item = QTableWidgetItem(str(value))
                if column > 0:
                    item.setTextAlignment(Qt.AlignCenter)
                if column == 1:
                    font = item.font()
                    font.setBold(True)
                    item.setFont(font)
                    item.setForeground(QColor("#26344d"))
                elif column == 2:
                    item.setForeground(QColor("#15946c"))
                elif column == 3:
                    item.setForeground(QColor("#d17a22"))
                elif column == 4:
                    item.setForeground(QColor("#176f68"))
                self.summary_table.setItem(row_index, column, item)
        self._finish_summary_table(scope, "违规原因明细")

    def _render_station_distribution(self):
        rows = self._get_station_distribution_rows()
        self.distribution_chart.hide()
        self.violation_chart.hide()
        self.station_distribution_chart.set_series_mode("counts")
        self.station_distribution_chart.show()
        self.detail_tabs.setTabEnabled(0, True)
        self.data_title.setText("完整数据")
        self.station_distribution_chart.set_rows(rows)

        all_total = sum(row["total"] for row in rows)
        all_has_phone = sum(row["has_phone"] for row in rows)
        overall_phone_rate = all_has_phone / all_total * 100 if all_total else 0
        self._add_kpis(
            [
                ("站点数量", len(rows), "个"),
                ("总计数", all_total, "条"),
                ("有电话数", all_has_phone, "条"),
                ("有电话占总计数", f"{overall_phone_rate:.1f}", "%"),
            ]
        )

        headers = [
            "站点",
            "总计数",
            "总数占比",
            "有电话数",
            "有电话数占比",
        ]
        self.summary_table.setColumnCount(len(headers))
        self.summary_table.setHorizontalHeaderLabels(headers)
        self.summary_table.setRowCount(len(rows))
        for row_index, row in enumerate(rows):
            values = [
                row["station"],
                row["total"],
                f'{row["total_share"]:.1f}%',
                row["has_phone"],
                f'{row["phone_share"]:.1f}%',
            ]
            for column, value in enumerate(values):
                item = QTableWidgetItem(str(value))
                if column in (1, 2):
                    item.setForeground(QColor("#1d8178"))
                    item.setTextAlignment(Qt.AlignCenter)
                elif column in (3, 4):
                    item.setForeground(QColor("#15946c"))
                    item.setTextAlignment(Qt.AlignCenter)
                if column in (1, 3):
                    font = item.font()
                    font.setBold(True)
                    item.setFont(font)
                self.summary_table.setItem(row_index, column, item)
        self._finish_summary_table(
            self._selected_scope(),
            "各站总计数占比与有电话数占比",
            (0,),
        )

    def refresh(self):
        self._sync_view_header()
        self._populate_station_options()
        self._clear_kpis()
        category = self.category_combo.currentData()
        self.timing_section.setVisible(
            not self._sidebar_navigation or category == "timing"
        )
        anomaly_button = getattr(self, "anomaly_button", None)
        if anomaly_button is not None:
            anomaly_button.setVisible(category == "completion")
        is_station_distribution = category == "station_distribution"
        if is_station_distribution:
            self.station_combo.blockSignals(True)
            self.station_combo.setCurrentIndex(0)
            self.station_combo.blockSignals(False)
            self.station_combo.setEnabled(False)
            self.station_combo.setToolTip("各站分布固定统计全部站点")
            if not self._sidebar_navigation:
                self._refresh_timing_kpis(None)
            self._render_station_distribution()
            return

        self.station_combo.setEnabled(True)
        self.station_combo.setToolTip("")
        user_id = self._selected_user_id()
        if category == "timing":
            self._refresh_timing_kpis(user_id)
            self._render_timing(user_id, self._selected_scope())
            return
        if not self._sidebar_navigation:
            self._refresh_timing_kpis(user_id)
        scope = self._selected_scope()
        if anomaly_button is not None and category == "completion":
            anomaly_rows = self._get_anomaly_rows(user_id)
            anomaly_total = sum(row["empty"] for row in anomaly_rows)
            anomaly_button.setText(
                ("返回完成类型" if self._showing_anomalies else "异常数据")
                + f"（{anomaly_total}）"
            )
            anomaly_button.setChecked(self._showing_anomalies)
            if self._showing_anomalies:
                self._render_anomalies(user_id, scope, anomaly_rows)
                return
        if category == "violation":
            self._render_violation(user_id, scope)
        else:
            self._render_completion(user_id, scope)
