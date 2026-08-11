import os
import re
import uuid
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from PyQt5.QtCore import (
    QAbstractTableModel,
    QModelIndex,
    QUrl,
    Qt,
    QThread,
    QTimer,
    pyqtSignal,
)
from PyQt5.QtGui import QColor, QDesktopServices
from PyQt5.QtWidgets import (
    QAbstractItemView,
    QCheckBox,
    QComboBox,
    QFrame,
    QGridLayout,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QInputDialog,
    QLabel,
    QLineEdit,
    QProgressBar,
    QPushButton,
    QSpinBox,
    QSplitter,
    QTableView,
    QTabWidget,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

from ..browser import check_builtin_chromium, get_builtin_chromium_path
from .file_dialogs import SystemFileDialog as QFileDialog
from ..database import (
    WORKFLOW_EMPTY_METRIC,
    WORKFLOW_HAS_PHONE_METRIC,
    WORKFLOW_INDIVIDUAL_METRIC,
    WORKFLOW_NO_OPERATION_METRIC,
    WORKFLOW_NO_PHONE_METRIC,
    WORKFLOW_NO_TRANSPORT_METRIC,
    WORKFLOW_TOTAL_METRIC,
    split_violation_reasons,
)
from ..tencent_docs import (
    TencentDocsImportResult,
    TencentDocsImportWorker,
    validate_tencent_docs_url,
)
from ..tools.aiqicha_tool import (
    ADDR_COL_NAME,
    COMPANY_COL_NAME,
    LEGAL_COL_NAME,
    PHONE_COL_NAME,
    TARGET_COLUMNS,
    QueryWorker,
    has_meaningful_value,
)
from ..tools.transport_tool import (
    BusinessBackfillWorker,
    Worker,
    is_excel_file_open,
)
from .frameless import FramelessMessageBox as QMessageBox
from .tencent_docs_dialog import (
    TencentDocsLinkDialog,
    TencentDocsProgressDialog,
)


class DataFrameTableModel(QAbstractTableModel):
    """轻量 DataFrame 表格模型，适合批量任务期间频繁刷新。"""

    def __init__(self, dataframe=None, parent=None):
        super().__init__(parent)
        self._df = dataframe if dataframe is not None else pd.DataFrame()

    @property
    def dataframe(self):
        return self._df

    def rowCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self._df.index)

    def columnCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self._df.columns)

    def data(self, index, role=Qt.DisplayRole):
        if not index.isValid():
            return None
        if role == Qt.DisplayRole:
            value = self._df.iat[index.row(), index.column()]
            return "" if pd.isna(value) else str(value)
        if role == Qt.TextAlignmentRole:
            return Qt.AlignLeft | Qt.AlignVCenter
        if role == Qt.BackgroundRole:
            column = str(self._df.columns[index.column()])
            if column in TARGET_COLUMNS or column in {"运输证号_纯数字", "回填状态", "查询状态"}:
                return QColor("#f2fbf6")
        return None

    def headerData(self, section, orientation, role=Qt.DisplayRole):
        if role != Qt.DisplayRole:
            return None
        if orientation == Qt.Horizontal and section < len(self._df.columns):
            return str(self._df.columns[section])
        if orientation == Qt.Vertical:
            return str(section + 1)
        return None

    def set_dataframe(self, dataframe):
        self.beginResetModel()
        self._df = dataframe if dataframe is not None else pd.DataFrame()
        self.endResetModel()

    def update_row(self, row_index, values):
        if row_index < 0 or row_index >= len(self._df.index):
            return
        changed_columns = []
        for name, value in values.items():
            if name not in self._df.columns:
                self.beginInsertColumns(QModelIndex(), len(self._df.columns), len(self._df.columns))
                self._df[name] = ""
                self.endInsertColumns()
            if self._df[name].dtype != object:
                self._df[name] = self._df[name].astype(object)
            self._df.at[row_index, name] = value
            changed_columns.append(self._df.columns.get_loc(name))
        if changed_columns:
            left = self.index(row_index, min(changed_columns))
            right = self.index(row_index, max(changed_columns))
            self.dataChanged.emit(left, right, [Qt.DisplayRole, Qt.BackgroundRole])


class BrowserCheckWorker(QThread):
    """在后台实际启动兼容 Chromium，避免健康检查阻塞界面。"""

    result_ready = pyqtSignal(bool, str, str)

    def run(self):
        try:
            executable, version = check_builtin_chromium()
            self.result_ready.emit(True, f"Chromium {version}", executable)
        except Exception as exc:
            self.result_ready.emit(False, "启动失败", str(exc))


class CollapsiblePanel(QFrame):
    """供 QSplitter 使用的可收起内容面板。"""

    expanded_changed = pyqtSignal(bool)

    def __init__(self, title, content, parent=None):
        super().__init__(parent)
        self.setObjectName("Card")
        self._expanded = True

        layout = QVBoxLayout(self)
        layout.setContentsMargins(12, 10, 12, 12)
        layout.setSpacing(8)

        header = QHBoxLayout()
        self.header_layout = header
        heading = QLabel(title)
        heading.setStyleSheet("font-size:14px;font-weight:700;color:#264b4c;")
        self.toggle_button = QPushButton("▾ 收起")
        self.toggle_button.setFixedWidth(78)
        self.toggle_button.clicked.connect(self.toggle)
        header.addWidget(heading)
        header.addStretch()
        header.addWidget(self.toggle_button)
        layout.addLayout(header)

        self.content = content
        layout.addWidget(content, 1)

    def add_header_widget(self, widget):
        """在收起按钮左侧添加仅属于当前面板的操作按钮。"""
        self.header_layout.insertWidget(
            self.header_layout.indexOf(self.toggle_button),
            widget,
        )

    def is_expanded(self):
        return self._expanded

    def toggle(self):
        self.set_expanded(not self._expanded)

    def set_expanded(self, expanded):
        expanded = bool(expanded)
        if expanded == self._expanded:
            return
        self._expanded = expanded
        self.content.setVisible(expanded)
        self.toggle_button.setText("▾ 收起" if expanded else "▸ 展开")
        self.setMaximumHeight(16_777_215 if expanded else self.sizeHint().height())
        self.updateGeometry()
        self.expanded_changed.emit(expanded)


class WorkflowPage(QWidget):
    """统一编排运输证、营运回填、爱企查查询的三步流水线。"""

    browser_check_completed = pyqtSignal(bool, str, str)
    REQUIRED_COLUMNS = ("车辆标识", "已协助补缴")
    # 约 60 帧/秒刷新计时文本；16 不是 10 的整数倍，可避免毫秒
    # 末位长期只在少数固定数字间变化。
    TIMING_DISPLAY_INTERVAL_MS = 16
    TIMING_HEARTBEAT_INTERVAL_MS = 5_000
    BROWSER_CHECK_ANIMATION_FRAMES = (
        "正在检测内置浏览器",
        "正在检测内置浏览器 ·",
        "正在检测内置浏览器 ··",
        "正在检测内置浏览器 ···",
    )

    def __init__(
        self,
        stats_recorder=None,
        timing_service=None,
        parent=None,
        *,
        client_preferences=None,
        account_key="",
        untracked_mode=False,
        captcha_reporter=None,
        captcha_model_manager=None,
        captcha_collection_enabled=None,
        captcha_sample_collection_enabled=None,
    ):
        super().__init__(parent)
        self._stats_recorder = stats_recorder
        self._timing_service = timing_service
        self._untracked_mode = bool(untracked_mode)
        self.client_preferences = client_preferences
        self.account_key = str(account_key or "").strip().lower()
        self.captcha_reporter = captcha_reporter
        self.captcha_model_manager = captcha_model_manager
        self.captcha_collection_enabled = captcha_collection_enabled
        self.captcha_sample_collection_enabled = (
            captcha_sample_collection_enabled
        )
        self.file_path = ""
        self.df = pd.DataFrame()
        self.model = DataFrameTableModel(self.df, self)
        self.current_worker = None
        self.current_step = 0
        self.pipeline_running = False
        self.stopping = False
        self.awaiting_login = False
        self._backfill_browser_loading = False
        self._task_id = None
        self._stats_recorded = False
        self._last_file_mtime = None
        self._retired_workers = []
        self.browser_check_worker = None
        self.browser_check_state = "unchecked"
        self._settings_browser_check_requested = False
        self._browser_start_dialog = None
        self._timing_finished_steps = set()
        self.last_force_shutdown_error = ""
        self.tencent_import_worker = None
        self.tencent_import_dialog = None

        self._build_ui()
        self._restore_run_settings()
        self._connect_run_settings_persistence()
        self._sync_mode_controls()

        self.preview_timer = QTimer(self)
        self.preview_timer.timeout.connect(self._poll_file_preview)
        self.preview_timer.start(800)

        self.timing_timer = QTimer(self)
        self.timing_timer.setTimerType(Qt.PreciseTimer)
        self.timing_timer.timeout.connect(self._timing_tick)
        self.timing_timer.start(self.TIMING_DISPLAY_INTERVAL_MS)

        self.timing_heartbeat_timer = QTimer(self)
        self.timing_heartbeat_timer.timeout.connect(self._timing_heartbeat)
        self.timing_heartbeat_timer.start(self.TIMING_HEARTBEAT_INTERVAL_MS)

    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(18, 16, 18, 18)
        root.setSpacing(10)

        header = QHBoxLayout()
        title_box = QVBoxLayout()
        title = QLabel("一键业务处理")
        title.setObjectName("PageTitle")
        subtitle = QLabel("一次选择表格，自动顺序完成运输证查询、营运企业回填和爱企查信息补齐。")
        subtitle.setObjectName("Muted")
        title_box.addWidget(title)
        title_box.addWidget(subtitle)
        header.addLayout(title_box)
        header.addStretch()
        root.addLayout(header)

        self.page_tabs = QTabWidget()
        self.page_tabs.setObjectName("WorkflowTabs")
        self.page_tabs.setDocumentMode(False)
        self.workflow_tab = QWidget()
        workflow_root = QVBoxLayout(self.workflow_tab)
        workflow_root.setContentsMargins(8, 10, 8, 8)
        workflow_root.setSpacing(10)
        self.settings_tab = QWidget()
        settings_root = QVBoxLayout(self.settings_tab)
        settings_root.setContentsMargins(14, 14, 14, 14)
        settings_root.setSpacing(12)
        self.page_tabs.addTab(self.workflow_tab, "业务处理")
        self.page_tabs.addTab(self.settings_tab, "运行设置")
        self.page_tabs.currentChanged.connect(self._on_page_tab_changed)
        root.addWidget(self.page_tabs, 1)

        file_group = QGroupBox("业务表格")
        self.file_group = file_group
        file_layout = QHBoxLayout(file_group)
        self.file_edit = QLineEdit()
        self.file_edit.setReadOnly(True)
        self.file_edit.setPlaceholderText("请选择 .xlsx 业务表格")
        choose_btn = QPushButton("选择表格")
        choose_btn.clicked.connect(self._choose_file)
        tencent_docs_btn = QPushButton("导入腾讯文档")
        tencent_docs_btn.clicked.connect(self._prompt_tencent_docs_import)
        reload_btn = QPushButton("刷新预览")
        reload_btn.clicked.connect(lambda: self._reload_preview(force=True))
        self.choose_btn = choose_btn
        self.tencent_docs_btn = tencent_docs_btn
        self.reload_btn = reload_btn
        file_layout.addWidget(self.file_edit, 1)
        file_layout.addWidget(choose_btn)
        file_layout.addWidget(tencent_docs_btn)
        file_layout.addWidget(reload_btn)
        workflow_root.addWidget(file_group)

        browser_group = QGroupBox("内置浏览器状态")
        browser_layout = QGridLayout(browser_group)
        self.browser_info = QLabel("Chromium（自动适配内置/系统）")
        self.browser_info.setStyleSheet("font-size:14px;font-weight:700;color:#264b4c;")
        self.browser_status_label = QLabel("● 尚未检测")
        self.browser_status_label.setStyleSheet("color:#647c7b;font-weight:600;")
        self.browser_detail_label = QLabel("打开本页面时会自动执行启动检查。")
        self.browser_detail_label.setObjectName("Muted")
        self.browser_detail_label.setWordWrap(True)
        self.browser_detail_label.hide()
        self.browser_check_btn = QPushButton("立即检测")
        self.browser_check_btn.clicked.connect(self.check_browser)
        self.browser_detail_toggle_btn = QPushButton("查看详情")
        self.browser_detail_toggle_btn.clicked.connect(
            self._toggle_browser_details
        )
        browser_layout.addWidget(self.browser_info, 0, 0)
        browser_layout.addWidget(self.browser_status_label, 0, 1)
        browser_layout.addWidget(self.browser_check_btn, 0, 2)
        browser_layout.addWidget(self.browser_detail_toggle_btn, 0, 3)
        browser_layout.addWidget(self.browser_detail_label, 1, 0, 1, 4)
        browser_layout.setColumnStretch(1, 1)
        settings_root.addWidget(browser_group)

        settings_group = QGroupBox("处理参数")
        settings_layout = QVBoxLayout(settings_group)
        settings_layout.setContentsMargins(16, 18, 16, 16)
        settings_layout.setSpacing(12)
        settings_intro = QLabel("根据业务场景配置自动化程度、验证码处理方式和查询范围。")
        settings_intro.setObjectName("Muted")
        settings_layout.addWidget(settings_intro)

        self.mode_combo = QComboBox()
        self.mode_combo.addItem("人工接管模式", False)
        self.mode_combo.addItem("全自动模式", True)
        self.mode_combo.setMinimumHeight(36)
        self.mode_combo.currentIndexChanged.connect(self._sync_mode_controls)
        self.manual_captcha = QCheckBox("验证码人工处理")
        self.manual_captcha.setChecked(True)
        self.manual_captcha.stateChanged.connect(self._sync_mode_controls)
        self.auto_continue = QCheckBox("人工处理后自动继续")
        self.auto_continue.setChecked(True)
        self.only_yellow = QCheckBox("运输证仅查询黄牌")
        self.only_yellow.setChecked(True)
        self.infinite_captcha = QCheckBox("验证码无限重试")
        self.infinite_captcha.stateChanged.connect(self._sync_mode_controls)
        self.page_retry = QSpinBox()
        self.page_retry.setRange(1, 99)
        self.page_retry.setValue(5)
        self.page_retry.setSuffix(" 次")
        self.page_retry.setMinimumWidth(100)
        self.captcha_retry = QSpinBox()
        self.captcha_retry.setRange(1, 9999)
        self.captcha_retry.setValue(10)
        self.captcha_retry.setSuffix(" 次")
        self.captcha_retry.setMinimumWidth(100)

        cards_layout = QHBoxLayout()
        cards_layout.setSpacing(12)

        self.mode_card, mode_layout = self._create_setting_card(
            "运行模式",
            "选择自动化程度，运行期间仍可随时暂停或停止流水线。",
        )
        mode_label = QLabel("自动化模式")
        mode_label.setObjectName("SettingFieldLabel")
        mode_layout.addWidget(mode_label)
        mode_layout.addWidget(self.mode_combo)
        self.mode_hint = QLabel()
        self.mode_hint.setObjectName("ModeHint")
        self.mode_hint.setWordWrap(True)
        mode_layout.addWidget(self.mode_hint)
        mode_layout.addStretch()
        cards_layout.addWidget(self.mode_card, 1)

        self.captcha_card, captcha_layout = self._create_setting_card(
            "验证码策略",
            "设置验证码由人工接管还是自动识别，以及识别失败后的重试方式。",
        )
        captcha_layout.addWidget(self.manual_captcha)
        captcha_layout.addWidget(self.auto_continue)
        captcha_retry_row = QHBoxLayout()
        captcha_retry_label = QLabel("识别重试")
        captcha_retry_label.setObjectName("SettingFieldLabel")
        captcha_retry_row.addWidget(captcha_retry_label)
        captcha_retry_row.addStretch()
        captcha_retry_row.addWidget(self.captcha_retry)
        captcha_layout.addLayout(captcha_retry_row)
        captcha_layout.addWidget(self.infinite_captcha)
        captcha_layout.addStretch()
        cards_layout.addWidget(self.captcha_card, 1)

        self.query_card, query_layout = self._create_setting_card(
            "查询策略",
            "控制运输证查询的车牌范围，以及结果列表加载失败时的重试次数。",
        )
        query_layout.addWidget(self.only_yellow)
        page_retry_row = QHBoxLayout()
        page_retry_label = QLabel("列表重试")
        page_retry_label.setObjectName("SettingFieldLabel")
        page_retry_row.addWidget(page_retry_label)
        page_retry_row.addStretch()
        page_retry_row.addWidget(self.page_retry)
        query_layout.addLayout(page_retry_row)
        query_hint = QLabel("关闭黄牌限制后，将按表格中的实际车牌颜色查询。")
        query_hint.setObjectName("SettingCardDescription")
        query_hint.setWordWrap(True)
        query_layout.addWidget(query_hint)
        query_layout.addStretch()
        cards_layout.addWidget(self.query_card, 1)

        settings_layout.addLayout(cards_layout)
        self.settings_group = settings_group
        settings_root.addWidget(settings_group)
        settings_root.addStretch()

        steps_layout = QHBoxLayout()
        self.step_labels = {}
        self.step_cards = []
        step_definitions = (
            (1, "步骤 1", "查询运输证号"),
            (2, "步骤 2", "回填营运企业"),
            (3, "步骤 3", "爱企查补齐信息"),
        )
        for step, caption, description in step_definitions:
            card = QFrame()
            card.setObjectName("Card")
            card_layout = QVBoxLayout(card)
            heading = QLabel(caption)
            heading.setStyleSheet("font-size:13px;font-weight:700;color:#526e6d;")
            name = QLabel(description)
            name.setStyleSheet("font-size:15px;font-weight:700;color:#173a3d;")
            status = QLabel("等待执行")
            status.setStyleSheet("color:#647c7b;")
            card_layout.addWidget(heading)
            card_layout.addWidget(name)
            card_layout.addWidget(status)
            steps_layout.addWidget(card, 1)
            self.step_cards.append(card)
            self.step_labels[step] = status
        workflow_root.addLayout(steps_layout)

        action_layout = QHBoxLayout()
        self.start_btn = QPushButton("▶ 一键执行三个步骤")
        self.start_btn.setObjectName("PrimaryButton")
        self.start_btn.setMinimumHeight(40)
        self.start_btn.clicked.connect(self.start_pipeline)
        self.pause_btn = QPushButton("暂停")
        self.pause_btn.setEnabled(False)
        self.pause_btn.clicked.connect(self.pause_pipeline)
        self.continue_btn = QPushButton("继续执行")
        self.continue_btn.setEnabled(False)
        self.continue_btn.clicked.connect(self.continue_pipeline)
        self.stop_btn = QPushButton("停止")
        self.stop_btn.setObjectName("DangerButton")
        self.stop_btn.setEnabled(False)
        self.stop_btn.clicked.connect(self.stop_pipeline)
        action_layout.addWidget(self.start_btn, 2)
        action_layout.addWidget(self.pause_btn)
        action_layout.addWidget(self.continue_btn)
        action_layout.addWidget(self.stop_btn)
        workflow_root.addLayout(action_layout)

        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 100)
        self.progress_bar.setFormat("整体进度 %p%")
        workflow_root.addWidget(self.progress_bar)

        self.timing_label = QLabel(
            "本次有效用时 00:00:00.000 · 本批次总用时 00:00:00.000"
        )
        self.timing_label.setObjectName("WorkflowTimingLabel")
        self.timing_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        self.timing_label.setStyleSheet(
            "color:#526e6d;font-size:12px;font-weight:600;padding:0 2px;"
        )
        self.timing_label.setToolTip(
            "总用时包含有效用时和暂停等待；停止后再次处理同一份输入数据时会累计前次用时。"
        )
        workflow_root.addWidget(self.timing_label)

        splitter = QSplitter(Qt.Vertical)
        splitter.setChildrenCollapsible(False)
        self.content_splitter = splitter
        self.table = QTableView()
        self.table.setModel(self.model)
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.table.horizontalHeader().setStretchLastSection(True)
        self.preview_panel = CollapsiblePanel(
            "数据预览（随处理结果实时刷新）",
            self.table,
        )
        self.open_workbook_btn = QPushButton("打开表格")
        self.open_workbook_btn.setEnabled(False)
        self.open_workbook_btn.clicked.connect(self._open_current_workbook)
        self.open_workbook_folder_btn = QPushButton("打开表格所在文件夹")
        self.open_workbook_folder_btn.setEnabled(False)
        self.open_workbook_folder_btn.clicked.connect(
            self._open_current_workbook_folder
        )
        self.preview_panel.add_header_widget(self.open_workbook_btn)
        self.preview_panel.add_header_widget(self.open_workbook_folder_btn)
        self.preview_toggle_btn = self.preview_panel.toggle_button
        self.preview_panel.expanded_changed.connect(self._rebalance_content_panels)
        splitter.addWidget(self.preview_panel)

        self.log_text = QTextEdit()
        self.log_text.setReadOnly(True)
        self.log_text.document().setMaximumBlockCount(5000)
        self.log_panel = CollapsiblePanel("流水线日志", self.log_text)
        self.log_toggle_btn = self.log_panel.toggle_button
        self.log_panel.expanded_changed.connect(self._rebalance_content_panels)
        splitter.addWidget(self.log_panel)
        splitter.setStretchFactor(0, 4)
        splitter.setStretchFactor(1, 2)
        splitter.setSizes([480, 220])
        self.collapsed_content_spacer = QWidget()
        self.collapsed_content_spacer.hide()
        workflow_root.addWidget(splitter, 1)
        workflow_root.addWidget(self.collapsed_content_spacer, 1)

    def _on_page_tab_changed(self, index):
        if (
            self.page_tabs.widget(index) is self.settings_tab
            and not self._settings_browser_check_requested
        ):
            self._settings_browser_check_requested = True
            self.check_browser()

    @staticmethod
    def _create_setting_card(title, description):
        card = QFrame()
        card.setObjectName("SettingCard")
        layout = QVBoxLayout(card)
        layout.setContentsMargins(18, 16, 18, 16)
        layout.setSpacing(11)
        heading = QLabel(title)
        heading.setObjectName("SettingCardTitle")
        detail = QLabel(description)
        detail.setObjectName("SettingCardDescription")
        detail.setWordWrap(True)
        layout.addWidget(heading)
        layout.addWidget(detail)
        layout.addSpacing(3)
        return card, layout

    def _toggle_browser_details(self):
        show_details = self.browser_detail_label.isHidden()
        self.browser_detail_label.setVisible(show_details)
        self.browser_detail_toggle_btn.setText(
            "收起详情" if show_details else "查看详情"
        )

    def check_browser(self):
        worker = self.browser_check_worker
        if worker is not None and worker.isRunning():
            return
        if self.pipeline_running:
            return
        self.browser_check_state = "checking"
        self.browser_status_label.setText("● 正在启动检测…")
        self.browser_status_label.setStyleSheet("color:#176f68;font-weight:600;")
        self.browser_detail_label.setText("正在实际启动 Chromium 并访问空白页。")
        self.browser_check_btn.setEnabled(False)

        worker = BrowserCheckWorker(self)
        self.browser_check_worker = worker
        worker.result_ready.connect(self._browser_check_finished)
        worker.finished.connect(
            lambda checked_worker=worker: self._browser_check_worker_finished(
                checked_worker
            )
        )
        worker.start()

    def _browser_check_finished(self, success, summary, details):
        self.browser_check_state = "ready" if success else "failed"
        if success:
            self.browser_status_label.setText(f"● 运行正常 · {summary}")
            self.browser_status_label.setStyleSheet("color:#1d8d70;font-weight:600;")
            self.browser_detail_label.setText(f"已成功启动并关闭测试实例。\n{details}")
        else:
            self.browser_status_label.setText("● 检测失败")
            self.browser_status_label.setStyleSheet("color:#d33f49;font-weight:600;")
            self.browser_detail_label.setText(details)
        self.browser_check_btn.setEnabled(not self.pipeline_running)
        self.browser_check_completed.emit(success, summary, details)

    def _browser_check_worker_finished(self, worker):
        if self.browser_check_worker is worker:
            self.browser_check_worker = None
        worker.deleteLater()

    def _rebalance_content_panels(self, _expanded=None):
        preview_expanded = self.preview_panel.is_expanded()
        log_expanded = self.log_panel.is_expanded()
        self.collapsed_content_spacer.setVisible(
            not preview_expanded and not log_expanded
        )
        if preview_expanded and log_expanded:
            self.content_splitter.setSizes([480, 220])
        elif preview_expanded:
            self.content_splitter.setSizes([1_000, 1])
        elif log_expanded:
            self.content_splitter.setSizes([1, 1_000])
        else:
            self.content_splitter.setSizes([1, 1])

    def _refresh_timing_label(self, snapshot=None):
        service = self._timing_service
        if service is None:
            self.timing_label.setText(
                "本次有效用时 00:00:00.000 · 本批次总用时 00:00:00.000"
            )
            return
        snapshot = snapshot or service.snapshot()
        run_text = service.format_duration(snapshot.get("run_active_ms", 0))
        batch_total_ms = snapshot.get("batch_total_ms")
        if batch_total_ms is None:
            batch_total_ms = snapshot.get("batch_active_ms", 0) + snapshot.get(
                "batch_paused_ms",
                0,
            )
        batch_text = service.format_duration(batch_total_ms)
        paused_text = service.format_duration(snapshot.get("run_paused_ms", 0))
        prefix = "暂停中 · " if snapshot.get("state") == "paused" else ""
        self.timing_label.setText(
            f"{prefix}本次有效用时 {run_text} · 本批次总用时 {batch_text}"
        )
        self.timing_label.setToolTip(
            f"本次暂停等待 {paused_text}。总用时包含有效用时和暂停等待；"
            "停止后再次处理相同输入数据时会累计前次总用时。"
        )

    def _timing_tick(self):
        service = self._timing_service
        if service is None or not service.is_active:
            return
        self._refresh_timing_label()

    def _timing_heartbeat(self):
        service = self._timing_service
        if service is None or not service.is_active:
            return
        try:
            service.heartbeat()
        except Exception as exc:
            self._log(f"⚠️ 计时心跳保存失败：{exc}")

    def _timing_start_step(self, step):
        service = self._timing_service
        if service is None:
            return
        try:
            service.start_step(step)
        except Exception as exc:
            self._log(f"⚠️ 步骤 {step} 计时启动失败：{exc}")

    def _timing_finish_step(self, step, status, error_summary=""):
        service = self._timing_service
        if service is None or step in self._timing_finished_steps:
            return
        try:
            service.finish_step(status, error_summary)
            self._timing_finished_steps.add(step)
            self._refresh_timing_label()
        except Exception as exc:
            self._log(f"⚠️ 步骤 {step} 计时结束失败：{exc}")

    def _timing_retry(self, retry_type, reason=""):
        service = self._timing_service
        if service is None:
            return
        try:
            service.record_retry(str(retry_type), str(reason))
        except Exception as exc:
            self._log(f"⚠️ 重试计时记录失败：{exc}")

    def _timing_pause(self, reason):
        if self._timing_service is None:
            return
        try:
            self._timing_service.pause(reason)
            self._refresh_timing_label()
        except Exception as exc:
            self._log(f"⚠️ 暂停计时记录失败：{exc}")

    def _timing_resume(self, reason):
        if self._timing_service is None:
            return
        try:
            self._timing_service.resume(reason)
            self._refresh_timing_label()
        except Exception as exc:
            self._log(f"⚠️ 恢复计时记录失败：{exc}")

    def _timing_request_stop(self, reason):
        if self._timing_service is None:
            return
        try:
            self._timing_service.request_stop(reason)
            self._refresh_timing_label()
        except Exception as exc:
            self._log(f"⚠️ 停止计时记录失败：{exc}")

    def _timing_finish_run(self, status, reason=""):
        service = self._timing_service
        if service is None or not service.is_active:
            return
        try:
            snapshot = service.finish_run(status, reason, reason if status == "failed" else "")
            self._refresh_timing_label(snapshot)
        except Exception as exc:
            self._log(f"⚠️ 流水线计时结束失败：{exc}")

    def _log(self, message):
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.log_text.append(f"[{timestamp}] {message}")
        scrollbar = self.log_text.verticalScrollBar()
        scrollbar.setValue(scrollbar.maximum())

    def _set_step_status(self, step, text, state="waiting"):
        colors = {
            "waiting": "#647c7b",
            "running": "#176f68",
            "paused": "#d18400",
            "success": "#1d8d70",
            "failed": "#d33f49",
            "stopped": "#9a6415",
        }
        label = self.step_labels[step]
        label.setText(text)
        label.setStyleSheet(f"color:{colors.get(state, colors['waiting'])};font-weight:600;")

    def _sync_mode_controls(self):
        auto_mode = bool(self.mode_combo.currentData())
        self.mode_hint.setText(
            "全自动运行：优先自动识别验证码，重试耗尽后按当前步骤规则跳过或停止。"
            if auto_mode
            else "人工接管：可选择直接人工处理，或先自动识别、失败后再由人工接管。"
        )
        self.manual_captcha.setEnabled(not auto_mode)
        self.auto_continue.setEnabled(not auto_mode)
        retry_enabled = auto_mode or not self.manual_captcha.isChecked()
        self.infinite_captcha.setEnabled(retry_enabled)
        self.captcha_retry.setEnabled(retry_enabled and not self.infinite_captcha.isChecked())

    def _run_settings_snapshot(self):
        return {
            "auto_mode": bool(self.mode_combo.currentData()),
            "manual_captcha": self.manual_captcha.isChecked(),
            "auto_continue": self.auto_continue.isChecked(),
            "only_yellow": self.only_yellow.isChecked(),
            "infinite_captcha": self.infinite_captcha.isChecked(),
            "page_retry": self.page_retry.value(),
            "captcha_retry": self.captcha_retry.value(),
        }

    def _restore_run_settings(self):
        if self.client_preferences is None or not self.account_key:
            return
        settings = self.client_preferences.workflow_run_settings(
            self.account_key
        )
        if not settings:
            return
        mode_index = self.mode_combo.findData(
            bool(settings.get("auto_mode", False))
        )
        if mode_index >= 0:
            self.mode_combo.setCurrentIndex(mode_index)
        self.manual_captcha.setChecked(
            bool(settings.get("manual_captcha", True))
        )
        self.auto_continue.setChecked(
            bool(settings.get("auto_continue", True))
        )
        self.only_yellow.setChecked(
            bool(settings.get("only_yellow", True))
        )
        self.infinite_captcha.setChecked(
            bool(settings.get("infinite_captcha", False))
        )
        self.page_retry.setValue(int(settings.get("page_retry", 5)))
        self.captcha_retry.setValue(int(settings.get("captcha_retry", 10)))

    def _connect_run_settings_persistence(self):
        self.mode_combo.currentIndexChanged.connect(self._save_run_settings)
        for checkbox in (
            self.manual_captcha,
            self.auto_continue,
            self.only_yellow,
            self.infinite_captcha,
        ):
            checkbox.toggled.connect(self._save_run_settings)
        self.page_retry.valueChanged.connect(self._save_run_settings)
        self.captcha_retry.valueChanged.connect(self._save_run_settings)

    def _save_run_settings(self, *_args):
        if self.client_preferences is None or not self.account_key:
            return
        try:
            self.client_preferences.set_workflow_run_settings(
                self.account_key,
                self._run_settings_snapshot(),
            )
        except (OSError, TypeError, ValueError) as exc:
            self._log(f"⚠️ 运行设置保存失败：{exc}")

    def _update_workbook_action_buttons(self):
        path_exists = bool(
            self.file_path and Path(self.file_path).is_file()
        )
        import_running = self.tencent_import_worker is not None
        enabled = path_exists and not self.pipeline_running and not import_running
        self.open_workbook_btn.setEnabled(enabled)
        self.open_workbook_folder_btn.setEnabled(enabled)

    def _open_current_workbook(self):
        path = Path(self.file_path) if self.file_path else None
        if path is None or not path.is_file():
            self._update_workbook_action_buttons()
            QMessageBox.warning(self, "无法打开表格", "当前业务表格不存在。")
            return
        if not QDesktopServices.openUrl(QUrl.fromLocalFile(str(path.resolve()))):
            QMessageBox.warning(
                self,
                "无法打开表格",
                "系统没有可用于打开该 Excel 表格的程序。",
            )

    def _open_current_workbook_folder(self):
        path = Path(self.file_path) if self.file_path else None
        if path is None or not path.is_file():
            self._update_workbook_action_buttons()
            QMessageBox.warning(
                self,
                "无法打开文件夹",
                "当前业务表格不存在。",
            )
            return
        folder = path.resolve().parent
        if not QDesktopServices.openUrl(QUrl.fromLocalFile(str(folder))):
            QMessageBox.warning(
                self,
                "无法打开文件夹",
                f"系统无法打开文件夹：\n{folder}",
            )

    def _choose_file(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "选择业务表格", "", "Excel 工作簿 (*.xlsx)"
        )
        if not path:
            return
        self.file_path = path
        self.file_edit.setText(path)
        self._last_file_mtime = None
        self._reload_preview(force=True)
        self._update_workbook_action_buttons()

    def _saved_tencent_document_url(self):
        if self.client_preferences is None:
            return ""
        return self.client_preferences.tencent_document_url(self.account_key)

    def _aiqicha_profile_directory(self):
        if self.client_preferences is None or not self.account_key:
            return None
        account_directory = re.sub(
            r"[^A-Za-z0-9._-]+",
            "-",
            self.account_key,
        ).strip("-.")
        account_directory = account_directory[:80] or "local"
        return (
            self.client_preferences.path.parent
            / "browser-profiles"
            / "aiqicha"
            / account_directory
        )

    def _prompt_tencent_docs_import(self):
        if self.pipeline_running or (
            self.tencent_import_worker is not None
            and self.tencent_import_worker.isRunning()
        ):
            return
        dialog = TencentDocsLinkDialog(
            self._saved_tencent_document_url(),
            self,
        )
        if dialog.exec_() != dialog.Accepted:
            return
        url = dialog.value()
        try:
            url = validate_tencent_docs_url(url)
            get_builtin_chromium_path()
            if self.client_preferences is not None:
                self.client_preferences.set_tencent_document_url(
                    self.account_key,
                    url,
                )
        except (OSError, RuntimeError, ValueError) as exc:
            QMessageBox.warning(self, "无法导入腾讯文档", str(exc))
            return
        self._start_tencent_docs_import(url)

    def _start_tencent_docs_import(self, url):
        data_directory = (
            self.client_preferences.path.parent
            if self.client_preferences is not None
            else None
        )
        worker = TencentDocsImportWorker(
            url,
            self.account_key,
            data_directory=data_directory,
            parent=self,
        )
        dialog = TencentDocsProgressDialog(self)
        dialog.cancel_requested.connect(worker.stop)
        worker.status_changed.connect(dialog.set_status)
        worker.progress_changed.connect(dialog.set_progress)
        worker.import_succeeded.connect(
            lambda result, current=worker, progress=dialog:
            self._tencent_docs_import_succeeded(
                current,
                progress,
                result,
            )
        )
        worker.import_failed.connect(
            lambda message, current=worker, progress=dialog:
            self._tencent_docs_import_failed(
                current,
                progress,
                message,
            )
        )
        worker.import_cancelled.connect(
            lambda current=worker, progress=dialog:
            self._tencent_docs_import_cancelled(current, progress)
        )
        worker.finished.connect(
            lambda current=worker, progress=dialog:
            self._tencent_docs_import_finished(current, progress)
        )
        self.tencent_import_worker = worker
        self.tencent_import_dialog = dialog
        self.choose_btn.setEnabled(False)
        self.tencent_docs_btn.setEnabled(False)
        self.reload_btn.setEnabled(False)
        self._update_workbook_action_buttons()
        self._log("开始导入腾讯文档；此过程不计入业务处理时间。")
        dialog.open()
        worker.start()

    def _tencent_docs_import_succeeded(
        self,
        worker,
        dialog,
        result,
    ):
        if (
            worker is not self.tencent_import_worker
            or not isinstance(result, TencentDocsImportResult)
        ):
            return
        dialog.finish()
        self.file_path = str(result.path)
        self.file_edit.setText(self.file_path)
        self._last_file_mtime = None
        loaded = self._reload_preview(force=True)
        self._update_workbook_action_buttons()
        self._log(
            "腾讯文档导入完成：从原表第 "
            f"{result.source_start_row} 行开始复制 {result.copied_rows} 行，"
            f"新表位于 {result.path}"
        )
        if loaded:
            QMessageBox.information(
                self,
                "腾讯文档导入完成",
                f"已按“{result.company_header}”列筛选：\n"
                f"从原表第 {result.source_start_row} 行开始，"
                f"复制 {result.copied_rows} 行。\n\n"
                "新 Excel 已自动导入程序，且本次导入不计入业务时间。\n"
                f"文件位置：\n{result.path}",
            )
        else:
            QMessageBox.warning(
                self,
                "新表读取失败",
                "腾讯文档内容已生成 Excel，但程序无法读取该文件：\n"
                f"{result.path}",
            )

    def _tencent_docs_import_failed(self, worker, dialog, message):
        if worker is not self.tencent_import_worker:
            return
        dialog.finish()
        self._log(f"腾讯文档导入失败：{message}")
        QMessageBox.warning(self, "腾讯文档导入失败", str(message))

    def _tencent_docs_import_cancelled(self, worker, dialog):
        if worker is not self.tencent_import_worker:
            return
        dialog.finish()
        self._log("已取消腾讯文档导入；未开始业务计时。")

    def _tencent_docs_import_finished(self, worker, dialog):
        if worker is self.tencent_import_worker:
            self.tencent_import_worker = None
        if dialog is self.tencent_import_dialog:
            self.tencent_import_dialog = None
        if dialog.isVisible():
            dialog.finish()
        dialog.deleteLater()
        worker.deleteLater()
        self.choose_btn.setEnabled(not self.pipeline_running)
        self.tencent_docs_btn.setEnabled(not self.pipeline_running)
        self.reload_btn.setEnabled(not self.pipeline_running)
        self._update_workbook_action_buttons()

    def _read_dataframe(self):
        return pd.read_excel(
            self.file_path,
            engine="openpyxl",
            dtype=str,
        ).fillna("")

    def _reload_preview(self, force=False):
        if not self.file_path or not os.path.exists(self.file_path):
            self._update_workbook_action_buttons()
            return False
        try:
            mtime = os.path.getmtime(self.file_path)
            if not force and self._last_file_mtime == mtime:
                return False
            dataframe = self._read_dataframe()
        except (PermissionError, OSError):
            return False
        except Exception as exc:
            if force:
                QMessageBox.warning(self, "读取失败", f"无法读取表格：\n{exc}")
            return False
        self.df = dataframe
        self.model.set_dataframe(self.df)
        self._last_file_mtime = mtime
        self._update_workbook_action_buttons()
        if force:
            self.table.resizeColumnsToContents()
            self._log(f"已加载表格：{len(self.df)} 行 × {len(self.df.columns)} 列")
        return True

    def _poll_file_preview(self):
        if self.current_step in {1, 2} or not self.pipeline_running:
            self._reload_preview(force=False)

    def _captcha_retry_count(self):
        return 9999 if self.infinite_captcha.isChecked() else self.captcha_retry.value()

    def _report_captcha_attempt(self, event):
        if self._untracked_mode or not callable(self.captcha_reporter):
            return
        try:
            self.captcha_reporter(event)
        except Exception as exc:
            self._log(f"⚠️ 验证码学习数据上报失败：{exc}")

    def _set_controls_running(self, running):
        self.pipeline_running = running
        self.choose_btn.setEnabled(not running)
        self.tencent_docs_btn.setEnabled(
            not running
            and not (
                self.tencent_import_worker is not None
                and self.tencent_import_worker.isRunning()
            )
        )
        self.reload_btn.setEnabled(not running)
        self._update_workbook_action_buttons()
        self.settings_group.setEnabled(not running)
        self.browser_check_btn.setEnabled(not running and self.browser_check_state != "checking")
        self.start_btn.setEnabled(not running)
        self.pause_btn.setEnabled(running)
        self.stop_btn.setEnabled(running)
        if not running:
            self.continue_btn.setEnabled(False)

    @staticmethod
    def _browser_start_result_allows_pipeline(result):
        return result in (QMessageBox.Ok, QMessageBox.Ignore)

    def _run_browser_start_check_dialog(self):
        dialog = QMessageBox(self)
        dialog.setWindowTitle("启动前浏览器检测")
        dialog.setText("正在检测内置浏览器")
        dialog.setInformativeText(
            "正在实际启动 Chromium。检测成功后将自动开始业务处理。"
        )
        dialog.setIcon(QMessageBox.Information)
        dialog.setStandardButtons(QMessageBox.Ignore | QMessageBox.Cancel)
        progress = QProgressBar(dialog)
        progress.setObjectName("BrowserCheckProgress")
        progress.setRange(0, 0)
        progress.setTextVisible(False)
        progress.setFixedHeight(7)
        dialog.layout().insertWidget(dialog.layout().count() - 1, progress)

        animation_frames = self.BROWSER_CHECK_ANIMATION_FRAMES
        animation_state = {"index": 0}
        animation_timer = QTimer(dialog)
        animation_timer.setObjectName("BrowserCheckAnimationTimer")

        def advance_animation():
            animation_state["index"] = (
                animation_state["index"] + 1
            ) % len(animation_frames)
            dialog.setText(animation_frames[animation_state["index"]])

        animation_timer.timeout.connect(advance_animation)
        animation_timer.start(320)

        direct_start_btn = dialog.button(QMessageBox.Ignore)
        direct_start_btn.setText("直接开始")
        direct_start_btn.setObjectName("PrimaryButton")
        cancel_btn = dialog.button(QMessageBox.Cancel)
        dialog.button_layout.removeWidget(direct_start_btn)
        dialog.button_layout.addWidget(direct_start_btn)
        dialog.setDefaultButton(cancel_btn)
        self._browser_start_dialog = dialog

        def browser_check_completed(success, summary, details):
            if self._browser_start_dialog is not dialog:
                return
            animation_timer.stop()
            progress.hide()
            if success:
                dialog.setText("浏览器检测完成")
                dialog.setInformativeText(
                    f"{summary} 运行正常，正在自动开始业务处理。"
                )
                dialog.done(QMessageBox.Ok)
                return
            dialog.setIcon(QMessageBox.Warning)
            dialog.setText("浏览器检测失败")
            dialog.setInformativeText(
                f"{details}\n\n你可以取消本次启动，或跳过检测直接开始。"
            )

        self.browser_check_completed.connect(browser_check_completed)
        self.check_browser()
        if self.browser_check_state == "ready":
            result = QMessageBox.Ok
        else:
            result = dialog.exec_()
        animation_timer.stop()
        try:
            self.browser_check_completed.disconnect(browser_check_completed)
        except TypeError:
            pass
        if self._browser_start_dialog is dialog:
            self._browser_start_dialog = None
        dialog.deleteLater()
        return self._browser_start_result_allows_pipeline(result)

    def start_pipeline(self):
        if self.pipeline_running:
            return
        if not self.file_path or not os.path.exists(self.file_path):
            QMessageBox.warning(self, "缺少表格", "请先选择业务表格。")
            return
        if is_excel_file_open(self.file_path):
            QMessageBox.warning(self, "表格被占用", "请先关闭 Excel/WPS 中打开的业务表格。")
            return
        if not self._reload_preview(force=True):
            return
        missing = [name for name in self.REQUIRED_COLUMNS if name not in self.df.columns]
        if missing:
            QMessageBox.warning(self, "表格列不完整", f"业务表格缺少：{', '.join(missing)}")
            return
        if (
            self.browser_check_state != "ready"
            and not self._run_browser_start_check_dialog()
        ):
            return
        try:
            get_builtin_chromium_path()
        except RuntimeError as exc:
            QMessageBox.warning(self, "内置浏览器不可用", str(exc))
            return

        self.stopping = False
        self.awaiting_login = False
        self._task_id = uuid.uuid4().hex
        self._stats_recorded = False
        self._timing_finished_steps = set()
        timing_result = None
        if self._timing_service is not None:
            try:
                timing_result = self._timing_service.start_run(
                    self.file_path,
                    self.df,
                    run_id=self._task_id,
                )
            except Exception as exc:
                QMessageBox.warning(
                    self,
                    "计时任务创建失败",
                    f"无法创建本次业务计时记录：\n{exc}",
                )
                return
        if self.log_text.toPlainText().strip():
            self.log_text.append("")
        run_number = timing_result["run_number"] if timing_result else 1
        self._log(f"━━━━━━━━━━ 第 {run_number} 次流水线执行 ━━━━━━━━━━")
        self.current_step = 1
        self.progress_bar.setValue(0)
        self._set_step_status(1, "正在执行", "running")
        self._set_step_status(2, "等待步骤 1", "waiting")
        self._set_step_status(3, "等待步骤 2", "waiting")
        self._set_controls_running(True)
        self._log("一键三步流水线已启动。")
        if timing_result and timing_result["resumed"]:
            previous = self._timing_service.format_duration(
                timing_result["previous_active_ms"]
            )
            self._log(
                f"已接续同一业务批次，前 {timing_result['run_number'] - 1} 次"
                f"累计有效用时 {previous}。"
            )
        self._refresh_timing_label()
        self._start_transport_worker()

    def _start_transport_worker(self):
        self._timing_start_step(1)
        auto_mode = bool(self.mode_combo.currentData())
        worker = Worker(
            self.file_path,
            auto_mode,
            self.manual_captcha.isChecked() if not auto_mode else False,
            self.page_retry.value(),
            self.auto_continue.isChecked(),
            self._captcha_retry_count(),
            self.only_yellow.isChecked(),
            captcha_model_manager=self.captcha_model_manager,
            captcha_collection_enabled=self.captcha_collection_enabled,
            captcha_sample_collection_enabled=(
                self.captcha_sample_collection_enabled
            ),
        )
        self.current_worker = worker
        worker.log.connect(self._log)
        worker.progress.connect(lambda value: self._on_step_progress(1, value))
        worker.pause_signal.connect(lambda: self._on_pause_requested(1, "等待验证码或人工处理"))
        worker.input_signal.connect(self._transport_manual_input)
        if hasattr(worker, "retry_signal"):
            worker.retry_signal.connect(self._timing_retry)
        if hasattr(worker, "captcha_attempt_signal"):
            worker.captcha_attempt_signal.connect(self._report_captcha_attempt)
        worker.finished.connect(lambda result, obj=worker: self._transport_finished(obj, result))
        worker.start()

    def _transport_finished(self, worker, result):
        if worker is not self.current_worker:
            return
        self._retire_worker(worker)
        self._reload_preview(force=True)
        if self.stopping:
            self._timing_finish_step(1, "stopped", "用户手动停止")
            self._finish_stopped(1)
        elif result == "失败":
            self._timing_finish_step(1, "failed", "步骤 1 执行失败")
            self._set_step_status(1, "执行失败", "failed")
            self._finish_pipeline(
                False,
                "步骤 1 失败，流水线已停止。",
                outcome="failed",
            )
        else:
            self._timing_finish_step(1, "succeeded")
            self._set_step_status(1, "已完成", "success")
            self.current_step = 2
            self._set_step_status(2, "正在执行", "running")
            self._log("步骤 1 完成，自动开始步骤 2：营运企业回填。")
            QTimer.singleShot(0, self._start_backfill_worker)

    def _start_backfill_worker(self):
        self._timing_start_step(2)
        auto_mode = bool(self.mode_combo.currentData())
        worker = BusinessBackfillWorker(
            self.file_path,
            self.auto_continue.isChecked(),
            auto_mode,
            self._captcha_retry_count(),
            self.manual_captcha.isChecked() if not auto_mode else False,
            captcha_model_manager=self.captcha_model_manager,
            captcha_collection_enabled=self.captcha_collection_enabled,
            captcha_sample_collection_enabled=(
                self.captcha_sample_collection_enabled
            ),
        )
        self.current_worker = worker
        worker.log.connect(self._log)
        worker.progress.connect(lambda value: self._on_step_progress(2, value))
        worker.pause_signal.connect(lambda: self._on_pause_requested(2, "等待验证码或人工处理"))
        worker.input_signal.connect(self._business_manual_input)
        if hasattr(worker, "retry_signal"):
            worker.retry_signal.connect(self._timing_retry)
        if hasattr(worker, "captcha_attempt_signal"):
            worker.captcha_attempt_signal.connect(self._report_captcha_attempt)
        worker.browser_loading_started.connect(
            self._on_backfill_browser_loading_started
        )
        worker.browser_loading_finished.connect(
            self._on_backfill_browser_loading_finished
        )
        worker.finished.connect(lambda result, obj=worker: self._backfill_finished(obj, result))
        self._on_backfill_browser_loading_started()
        worker.start()

    def _on_backfill_browser_loading_started(self):
        if (
            not self.pipeline_running
            or self.stopping
            or self.current_step != 2
        ):
            return
        was_loading = self._backfill_browser_loading
        self._backfill_browser_loading = True
        self._timing_pause("等待步骤 2 浏览器加载")
        if not self.continue_btn.isEnabled():
            self._set_step_status(2, "等待浏览器加载", "paused")
        if not was_loading:
            self._log("步骤 2 正在加载浏览器，此段等待不计入有效用时。")

    def _on_backfill_browser_loading_finished(self):
        if not self._backfill_browser_loading:
            return
        self._backfill_browser_loading = False
        if (
            not self.pipeline_running
            or self.stopping
            or self.current_step != 2
        ):
            return
        if not self.continue_btn.isEnabled():
            self._timing_resume("步骤 2 浏览器加载完成")
            self._set_step_status(2, "正在执行", "running")
        self._log("步骤 2 浏览器加载完成，开始计入有效用时。")

    def _backfill_finished(self, worker, result):
        if worker is not self.current_worker:
            return
        self._backfill_browser_loading = False
        self._retire_worker(worker)
        self._reload_preview(force=True)
        if self.stopping:
            self._timing_finish_step(2, "stopped", "用户手动停止")
            self._finish_stopped(2)
        elif result == "失败":
            self._timing_finish_step(2, "failed", "步骤 2 执行失败")
            self._set_step_status(2, "执行失败", "failed")
            self._finish_pipeline(
                False,
                "步骤 2 失败，流水线已停止。",
                outcome="failed",
            )
        else:
            self._timing_finish_step(2, "succeeded")
            self._set_step_status(2, "已完成", "success")
            self.current_step = 3
            self._set_step_status(3, "正在启动浏览器", "running")
            self._log("步骤 2 完成，自动开始步骤 3：爱企查信息补齐。")
            QTimer.singleShot(0, self._start_aiqicha_worker)

    def _start_aiqicha_worker(self):
        self._timing_start_step(3)
        if not self._reload_preview(force=True):
            self._set_step_status(3, "读取表格失败", "failed")
            self._finish_pipeline(False, "步骤 3 无法重新读取业务表格。")
            return
        if COMPANY_COL_NAME not in self.df.columns:
            self._set_step_status(3, "缺少企业列", "failed")
            self._finish_pipeline(False, f"步骤 3 缺少“{COMPANY_COL_NAME}”列。")
            return
        for column in TARGET_COLUMNS:
            if column not in self.df.columns:
                self.df[column] = ""
        self.model.set_dataframe(self.df)

        worker = QueryWorker(
            self.df.copy(),
            COMPANY_COL_NAME,
            browser_profile_directory=self._aiqicha_profile_directory(),
        )
        self.current_worker = worker
        worker.log_signal.connect(self._log)
        worker.progress_signal.connect(lambda current, total: self._on_aiqicha_progress(current, total))
        worker.row_done_signal.connect(self._on_aiqicha_row)
        worker.login_required_signal.connect(self._on_login_required)
        worker.pause_signal.connect(lambda: self._on_pause_requested(3, "等待完成爱企查验证"))
        worker.resume_signal.connect(self._on_aiqicha_resumed)
        if hasattr(worker, "retry_signal"):
            worker.retry_signal.connect(self._timing_retry)
        worker.finished_signal.connect(lambda success, obj=worker: self._aiqicha_finished(obj, success))
        worker.start()

    def _on_login_required(self):
        self.awaiting_login = True
        self._timing_pause("等待用户登录爱企查")
        self.continue_btn.setEnabled(True)
        self.pause_btn.setEnabled(False)
        self._set_step_status(3, "请登录爱企查后点击继续", "paused")
        self._log("请在已打开的浏览器中登录爱企查，然后点击“继续执行”。")

    def _on_aiqicha_resumed(self):
        self.awaiting_login = False
        self._timing_resume("爱企查登录或验证已完成")
        self.continue_btn.setEnabled(False)
        self.pause_btn.setEnabled(True)
        self._set_step_status(3, "正在查询", "running")

    def _on_aiqicha_row(self, row_index, values):
        self.model.update_row(row_index, values)
        self.df = self.model.dataframe

    def _aiqicha_finished(self, worker, success):
        if worker is not self.current_worker:
            return
        self._retire_worker(worker)
        saved = self._save_aiqicha_results()
        if self.stopping:
            self._timing_finish_step(3, "stopped", "用户手动停止")
            self._finish_stopped(3)
        elif not success or not saved:
            self._timing_finish_step(3, "failed", "步骤 3 执行或保存失败")
            self._set_step_status(3, "执行失败" if not success else "保存失败", "failed")
            self._finish_pipeline(
                False,
                "步骤 3 未正常完成，请查看日志。",
                outcome="failed",
            )
        else:
            self._timing_finish_step(3, "succeeded")
            self._set_step_status(3, "已完成", "success")
            self.progress_bar.setValue(100)
            self._record_workflow_stats()
            self._finish_pipeline(
                True,
                "三个步骤已全部完成，结果已保存到原业务表格。",
                outcome="succeeded",
            )

    def _save_aiqicha_results(self):
        if self.df is None or self.df.empty:
            return True
        workbook = None
        temporary_path = ""
        try:
            workbook = openpyxl.load_workbook(self.file_path)
            sheet = workbook.active
            header = {cell.value: cell.column for cell in sheet[1] if cell.value is not None}
            next_column = max(header.values(), default=0) + 1
            for name in TARGET_COLUMNS:
                if name not in header:
                    sheet.cell(1, next_column, name)
                    header[name] = next_column
                    next_column += 1
            for row_index in range(len(self.df.index)):
                excel_row = row_index + 2
                for name in TARGET_COLUMNS:
                    value = self.df.at[row_index, name] if name in self.df.columns else ""
                    if pd.isna(value) or str(value).strip() in {"nan", "None"}:
                        continue
                    sheet.cell(excel_row, header[name], str(value).strip())
            directory = os.path.dirname(os.path.abspath(self.file_path))
            file_name = os.path.basename(self.file_path)
            stem, suffix = os.path.splitext(file_name)
            temporary_path = os.path.join(
                directory,
                f".{stem}.{uuid.uuid4().hex}.tmp{suffix}",
            )
            workbook.save(temporary_path)
            workbook.close()
            workbook = None
            os.replace(temporary_path, self.file_path)
            temporary_path = ""
            self._last_file_mtime = os.path.getmtime(self.file_path)
            self._log("爱企查结果已保存到原业务表格。")
            return True
        except PermissionError:
            self._log("❌ 保存失败：请关闭 Excel/WPS 中打开的业务表格。")
        except Exception as exc:
            self._log(f"❌ 保存爱企查结果失败：{exc}")
        finally:
            if workbook is not None:
                workbook.close()
            if temporary_path and os.path.exists(temporary_path):
                try:
                    os.unlink(temporary_path)
                except OSError:
                    pass
        return False

    @staticmethod
    def classify_workflow_row(row):
        """返回单行所属的唯一完成类型指标。"""
        company = str(row.get(COMPANY_COL_NAME, "")).strip()
        company_lower = company.lower()
        empty_company_values = {
            "", "nan", "none", "null", "无", "暂无", "已补缴",
            "查询失败", "验证码识别失败", "未查询到公司",
        }
        cert_text = str(row.get("运输证号_纯数字", "")).strip()
        has_transport = has_meaningful_value(cert_text) and any(
            character.isdigit() for character in cert_text
        )
        phone_text = str(row.get(PHONE_COL_NAME, "")).strip()
        has_phone = has_meaningful_value(phone_text) and any(
            character.isdigit() for character in phone_text
        )

        if company_lower in empty_company_values:
            return WORKFLOW_EMPTY_METRIC
        if "无运输证号" in company or not has_transport:
            return WORKFLOW_NO_TRANSPORT_METRIC
        if "无营运信息" in company:
            return WORKFLOW_NO_OPERATION_METRIC
        if "个体" in company or company == "个人":
            return WORKFLOW_INDIVIDUAL_METRIC
        if has_phone:
            return WORKFLOW_HAS_PHONE_METRIC
        return WORKFLOW_NO_PHONE_METRIC

    @classmethod
    def calculate_workflow_counts(cls, dataframe):
        """按互斥口径统计完整流程结果，六种分类之和等于总计。"""
        counts = {
            WORKFLOW_TOTAL_METRIC: 0,
            WORKFLOW_EMPTY_METRIC: 0,
            WORKFLOW_NO_TRANSPORT_METRIC: 0,
            WORKFLOW_NO_OPERATION_METRIC: 0,
            WORKFLOW_INDIVIDUAL_METRIC: 0,
            WORKFLOW_NO_PHONE_METRIC: 0,
            WORKFLOW_HAS_PHONE_METRIC: 0,
        }
        if dataframe is None:
            return counts

        counts[WORKFLOW_TOTAL_METRIC] = len(dataframe.index)
        for _, row in dataframe.iterrows():
            counts[cls.classify_workflow_row(row)] += 1
        return counts

    @classmethod
    def calculate_violation_counts(cls, dataframe):
        """仅按半角竖线拆分原因，并区分有电话与其他完成类型。"""
        totals = {}
        if dataframe is None or "原因" not in dataframe.columns:
            return totals
        for _, row in dataframe.iterrows():
            reason_text = str(row.get("原因", "")).strip()
            if reason_text.lower() in {"", "nan", "none", "null", "无", "暂无"}:
                continue
            reasons = split_violation_reasons(reason_text)
            has_phone = cls.classify_workflow_row(row) == WORKFLOW_HAS_PHONE_METRIC
            for reason in reasons:
                target = totals.setdefault(
                    reason, {"total": 0, "has_phone": 0, "other": 0}
                )
                target["total"] += 1
                target["has_phone" if has_phone else "other"] += 1
        return totals

    def _record_workflow_stats(self):
        if self._stats_recorded:
            return
        if self._untracked_mode:
            self._stats_recorded = True
            self._log(
                "游客模式：业务结果已保存，本次不记录统计、计时或待同步数据。"
            )
            return
        counts = self.calculate_workflow_counts(self.df)
        details = {
            "file_name": os.path.basename(self.file_path),
            "row_count": counts[WORKFLOW_TOTAL_METRIC],
            "violation_counts": self.calculate_violation_counts(self.df),
        }
        recorded = True
        if self._stats_recorder:
            recorded = self._stats_recorder(
                counts,
                "unified_workflow",
                details,
                self._task_id,
            )
        if recorded is not False:
            self._stats_recorded = True
        self._log(
            "本次完整流程统计："
            f"总计 {counts[WORKFLOW_TOTAL_METRIC]}，"
            f"空 {counts[WORKFLOW_EMPTY_METRIC]}，"
            f"无运输证号 {counts[WORKFLOW_NO_TRANSPORT_METRIC]}，"
            f"无营运信息 {counts[WORKFLOW_NO_OPERATION_METRIC]}，"
            f"个体经营 {counts[WORKFLOW_INDIVIDUAL_METRIC]}，"
            f"公司无电话 {counts[WORKFLOW_NO_PHONE_METRIC]}，"
            f"公司有电话 {counts[WORKFLOW_HAS_PHONE_METRIC]}。"
        )
        self._log(f"违规原因统计：共拆分出 {len(details['violation_counts'])} 种原因。")

    def _on_step_progress(self, step, value):
        value = max(0, min(100, int(value)))
        overall = int(((step - 1) * 100 + value) / 3)
        self.progress_bar.setValue(overall)
        self._set_step_status(step, f"处理中 {value}%", "running")

    def _on_aiqicha_progress(self, current, total):
        percent = int(current / total * 100) if total else 0
        self._on_step_progress(3, percent)

    def _on_pause_requested(self, step, text):
        if not self.pipeline_running or self.stopping:
            return
        self._timing_pause(text)
        self.continue_btn.setEnabled(True)
        self.pause_btn.setEnabled(False)
        self._set_step_status(step, text, "paused")

    def pause_pipeline(self):
        worker = self.current_worker
        if not worker or not worker.isRunning():
            return
        if isinstance(worker, (Worker, BusinessBackfillWorker)):
            worker.global_pause()
        elif isinstance(worker, QueryWorker):
            worker.pause()
        self._timing_pause("用户手动暂停")
        self.pause_btn.setEnabled(False)
        self.continue_btn.setEnabled(True)
        self._set_step_status(self.current_step, "已暂停", "paused")
        self._log("流水线已暂停。")

    def continue_pipeline(self):
        worker = self.current_worker
        if not worker or not worker.isRunning():
            return
        if isinstance(worker, QueryWorker) and self.awaiting_login:
            worker.confirm_login()
            self.awaiting_login = False
        else:
            worker.resume()
        if not self._backfill_browser_loading:
            self._timing_resume("用户继续执行")
        self.continue_btn.setEnabled(False)
        self.pause_btn.setEnabled(True)
        self._set_step_status(
            self.current_step,
            (
                "等待浏览器加载"
                if self._backfill_browser_loading
                else "正在执行"
            ),
            "paused" if self._backfill_browser_loading else "running",
        )
        self._log("流水线已恢复执行。")

    def stop_pipeline(self):
        if not self.pipeline_running:
            return
        self.stopping = True
        self._timing_request_stop("用户手动停止")
        self.pause_btn.setEnabled(False)
        self.continue_btn.setEnabled(False)
        self.stop_btn.setEnabled(False)
        worker = self.current_worker
        if worker and worker.isRunning():
            worker.stop()
            self._log("停止指令已发送，正在保存已完成的数据。")
        else:
            self._finish_stopped(self.current_step or 1)

    def _finish_stopped(self, step):
        if step == 3:
            self._save_aiqicha_results()
        self._timing_finish_step(step, "stopped", "用户手动停止")
        self._set_step_status(step, "已停止", "stopped")
        for later_step in range(step + 1, 4):
            self._set_step_status(later_step, "未执行", "waiting")
        self._finish_pipeline(
            False,
            "流水线已停止，已完成的数据已保留。",
            outcome="stopped",
        )

    def _finish_pipeline(self, success, message, outcome=None):
        outcome = outcome or ("succeeded" if success else "failed")
        self._timing_finish_run(outcome, message)
        self.current_worker = None
        self.awaiting_login = False
        self._backfill_browser_loading = False
        self._set_controls_running(False)
        self._reload_preview(force=True)
        self._log(("✅ " if success else "⚠️ ") + message)
        self.stopping = False
        self.current_step = 0
        self._cleanup_retired_workers()

    def _retire_worker(self, worker):
        self.current_worker = None
        self._retired_workers.append(worker)

    def _cleanup_retired_workers(self):
        alive = []
        for worker in self._retired_workers:
            if worker.isRunning():
                alive.append(worker)
            else:
                worker.deleteLater()
        self._retired_workers = alive

    def _transport_manual_input(self, plate):
        worker = self.current_worker
        if not isinstance(worker, Worker):
            return
        self._timing_pause("等待人工录入运输证号")
        text, accepted = QInputDialog.getText(
            self, "人工录入运输证号", f"车辆：{plate}\n请输入运输证号："
        )
        worker.input_result = text if accepted else ""
        worker.resume()
        self._timing_resume("人工录入运输证号已完成")
        self.continue_btn.setEnabled(False)
        self.pause_btn.setEnabled(True)

    def _business_manual_input(self, plate):
        worker = self.current_worker
        if not isinstance(worker, BusinessBackfillWorker):
            return
        self._timing_pause("等待人工录入企业信息")
        text, accepted = QInputDialog.getText(
            self, "人工录入企业", f"车辆：{plate}\n请输入公司/所有人名称："
        )
        worker.input_result = text if accepted else ""
        worker.resume()
        self._timing_resume("人工录入企业信息已完成")
        self.continue_btn.setEnabled(False)
        self.pause_btn.setEnabled(True)

    def shutdown(self, timeout_ms=8000):
        self.preview_timer.stop()
        tencent_import_worker = self.tencent_import_worker
        if tencent_import_worker and tencent_import_worker.isRunning():
            tencent_import_worker.stop()
            if not tencent_import_worker.wait(timeout_ms):
                return False
        if self.tencent_import_dialog is not None:
            self.tencent_import_dialog.finish()
        browser_check_worker = self.browser_check_worker
        if browser_check_worker and browser_check_worker.isRunning():
            if not browser_check_worker.wait(timeout_ms):
                return False
        worker = self.current_worker
        if worker and worker.isRunning():
            self.stopping = True
            worker.stop()
            if not worker.wait(timeout_ms):
                return False
        for retired_worker in self._retired_workers:
            if retired_worker.isRunning() and not retired_worker.wait(timeout_ms):
                return False
        if self.current_step == 3:
            self._save_aiqicha_results()
        if self._timing_service is not None and self._timing_service.is_active:
            self._timing_request_stop("客户端退出或注销")
            if self.current_step:
                self._timing_finish_step(
                    self.current_step,
                    "stopped",
                    "客户端退出或注销",
                )
            self._timing_finish_run("stopped", "客户端退出或注销")
        self.timing_timer.stop()
        self.timing_heartbeat_timer.stop()
        self.current_worker = None
        self.pipeline_running = False
        self._cleanup_retired_workers()
        return True

    def _shutdown_threads(self):
        threads = []
        for thread in (
            self.browser_check_worker,
            self.tencent_import_worker,
            self.current_worker,
            *self._retired_workers,
        ):
            if thread is not None and thread not in threads:
                threads.append(thread)
        return threads

    def has_running_shutdown_threads(self):
        return any(
            thread.isRunning()
            for thread in self._shutdown_threads()
        )

    def force_shutdown(self, terminate_wait_ms=1_500):
        """
        保存已完成数据和计时状态后，终止仍未响应停止的工作线程。

        返回 False 表示数据保护步骤失败，此时不会调用 QThread.terminate()。
        """
        self.last_force_shutdown_error = ""
        self.preview_timer.stop()
        self.stopping = True
        threads = self._shutdown_threads()

        for thread in threads:
            stop = getattr(thread, "stop", None)
            if callable(stop):
                try:
                    stop()
                except Exception:
                    pass

        if self.current_step == 3 and not self._save_aiqicha_results():
            self.last_force_shutdown_error = (
                "爱企查已完成结果未能写入业务表格。请关闭 Excel/WPS "
                "占用后重试，程序尚未强制退出。"
            )
            return False

        reason = "用户保存已完成数据后强制结束任务"
        service = self._timing_service
        if service is not None and service.is_active:
            try:
                service.request_stop(reason)
                if (
                    self.current_step
                    and self.current_step not in self._timing_finished_steps
                ):
                    service.finish_step("stopped", reason)
                    self._timing_finished_steps.add(self.current_step)
                snapshot = service.finish_run("stopped", reason)
                self._refresh_timing_label(snapshot)
            except Exception as exc:
                self.last_force_shutdown_error = (
                    f"业务计时状态保存失败：{exc}。程序尚未强制退出。"
                )
                self._log(f"❌ {self.last_force_shutdown_error}")
                return False

        self.timing_timer.stop()
        self.timing_heartbeat_timer.stop()
        stubborn_threads = []
        for thread in threads:
            if thread.isRunning():
                try:
                    thread.requestInterruption()
                except Exception:
                    pass
                try:
                    thread.terminate()
                except Exception:
                    pass
                try:
                    thread.wait(terminate_wait_ms)
                except Exception:
                    pass
            if thread.isRunning():
                stubborn_threads.append(thread)
            else:
                thread.deleteLater()

        self.browser_check_worker = (
            self.browser_check_worker
            if self.browser_check_worker in stubborn_threads
            else None
        )
        self.tencent_import_worker = (
            self.tencent_import_worker
            if self.tencent_import_worker in stubborn_threads
            else None
        )
        self.current_worker = (
            self.current_worker
            if self.current_worker in stubborn_threads
            else None
        )
        self._retired_workers = [
            thread
            for thread in stubborn_threads
            if thread is not self.browser_check_worker
            and thread is not self.current_worker
        ]
        self.pipeline_running = False
        self.awaiting_login = False
        self._log("已保存全部已完成数据，卡住的任务已强制结束。")
        return True
