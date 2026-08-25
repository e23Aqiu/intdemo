from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path
from typing import ClassVar

from PyQt5.QtCore import Qt, QTimer
from PyQt5.QtGui import QColor
from PyQt5.QtWidgets import (
    QAbstractItemView,
    QCheckBox,
    QComboBox,
    QFormLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QPlainTextEdit,
    QPushButton,
    QSpinBox,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
)

from ..online.api import ApiResponseError, NetworkUnavailable
from .account_page import AccountPage
from .auth_dialogs import RenameAccountDialog
from .file_dialogs import SystemFileDialog as QFileDialog
from .frameless import FramelessDialog
from .frameless import FramelessMessageBox as QMessageBox
from .loading_dialog import run_with_loading
from .table_utils import make_table_columns_resizable

_CALL_FAILED = object()
_FILTER_SELECTION_UNSET = object()


class _AccountSettingsDialog(FramelessDialog):
    def __init__(
        self,
        account=None,
        parent=None,
        *,
        current_account=None,
        roads=None,
    ):
        super().__init__(parent)
        self.editing_account = account
        self.current_account = current_account
        self.roads = list(roads or [])
        self._global_manager = current_account is None or current_account.is_admin
        self.setWindowTitle("在线账号设置")
        self.setMinimumWidth(430)
        layout = QVBoxLayout(self)
        title = QLabel("编辑在线账号" if account else "新建在线账号")
        title.setObjectName("PageTitle")
        layout.addWidget(title)
        form = QFormLayout()
        self.username = QLineEdit(account.username if account else "")
        self.username.setEnabled(account is None)
        self.username.setMaxLength(80)
        self.display_name = QLineEdit(account.name_label if account else "")
        self.display_name.setMaxLength(120)
        self.role = QComboBox()
        self.role.addItem("中心站账号", "station")
        if self._global_manager:
            self.role.addItem("测试账号", "test")
            self.role.addItem("路段管理员", "road_admin")
            self.role.addItem("管理员", "admin")
        self.scope = QComboBox()
        self.road = QComboBox()
        if self._global_manager:
            self.road.addItem("未分配（直属管理员）", None)
        for road in self.roads:
            road_id = str(road.get("id") or "").strip()
            road_name = str(road.get("name") or "").strip()
            if road_id and road_name:
                self.road.addItem(road_name, road_id)
        if current_account is not None and current_account.is_road_admin:
            current_road_id = str(current_account.road_id or "")
            if self.road.findData(current_road_id) < 0 and current_account.road_name:
                self.road.addItem(current_account.road_name, current_road_id)
            self.road.setCurrentIndex(
                max(0, self.road.findData(current_road_id))
            )
            self.road.setEnabled(False)
        self.device_limit = QSpinBox()
        self.device_limit.setRange(1, 10000)
        self.device_limit.setSuffix(" 台")
        self.device_limit.setToolTip(
            "按已登记且未撤销的设备计算；10000 台可视为不限制"
        )
        self.active = QCheckBox("允许登录")
        if account:
            account_type = (
                "test" if account.is_test else account.resolved_account_type
            )
            self.role.setCurrentIndex(self.role.findData(account_type))
            road_index = self.road.findData(
                str(account.road_id) if account.road_id else None
            )
            if road_index < 0 and account.road_id and account.road_name:
                self.road.addItem(account.road_name, str(account.road_id))
                road_index = self.road.count() - 1
            if road_index >= 0:
                self.road.setCurrentIndex(road_index)
            self.device_limit.setValue(
                int(getattr(account, "_device_limit", 10000) or 10000)
            )
            self.active.setChecked(account.is_active)
        else:
            self.device_limit.setValue(10000)
            self.active.setChecked(True)
        form.addRow("登录名", self.username)
        form.addRow("账号显示名", self.display_name)
        form.addRow("角色", self.role)
        form.addRow("所属路段", self.road)
        form.addRow("数据范围", self.scope)
        form.addRow("登录设备上限", self.device_limit)
        form.addRow("状态", self.active)
        layout.addLayout(form)
        hint = QLabel(
            "新账号初始密码为 123456，首次登录必须修改；"
            "达到设备上限后，新电脑将无法登录。"
            if account is None
            else "设备上限按未撤销设备计算；降低上限前请先在设备列表中撤销多余设备。"
        )
        hint.setObjectName("Muted")
        hint.setWordWrap(True)
        layout.addWidget(hint)
        buttons = QHBoxLayout()
        buttons.addStretch()
        cancel = QPushButton("取消")
        cancel.clicked.connect(self.reject)
        save = QPushButton("保存")
        save.setObjectName("PrimaryButton")
        save.clicked.connect(self._accept_if_valid)
        buttons.addWidget(cancel)
        buttons.addWidget(save)
        layout.addLayout(buttons)
        self.role.currentIndexChanged.connect(self._sync_hierarchy_controls)
        self.road.currentIndexChanged.connect(self._sync_hierarchy_controls)
        self.display_name.textChanged.connect(self._sync_road_manager_name)
        self._sync_hierarchy_controls()
        if account:
            scope_index = self.scope.findData(account.effective_data_scope)
            self.scope.setCurrentIndex(scope_index if scope_index >= 0 else 0)

    def _sync_hierarchy_controls(self, *_args):
        account_type = self.role.currentData()
        current_scope = self.scope.currentData()
        has_assigned_road = bool(self.road.currentData())
        self.scope.blockSignals(True)
        self.scope.clear()
        if account_type == "admin":
            scope_options = [("全部数据", "all")]
        elif account_type == "test":
            scope_options = [("仅本人数据", "own"), ("全部数据", "all")]
        elif self.current_account is not None and self.current_account.is_road_admin:
            scope_options = [("仅本人数据", "own"), ("本路段数据", "road")]
            if (
                self.editing_account is not None
                and self.editing_account.effective_data_scope == "all"
            ):
                scope_options.append(("全部数据（仅管理员可授予）", "all"))
        elif account_type == "road_admin" or has_assigned_road:
            scope_options = [
                ("仅本人数据", "own"),
                ("本路段数据", "road"),
                ("全部数据", "all"),
            ]
        else:
            scope_options = [("仅本人数据", "own"), ("全部数据", "all")]
        for label, value in scope_options:
            self.scope.addItem(label, value)
        index = self.scope.findData(current_scope)
        if self.editing_account is None and account_type == "road_admin":
            index = self.scope.findData("road")
        elif index < 0 and account_type == "road_admin":
            index = self.scope.findData("road")
        self.scope.setCurrentIndex(index if index >= 0 else 0)
        self.scope.blockSignals(False)

        if account_type == "road_admin":
            self.road.setEditable(True)
            self.road.setEditText(self.display_name.text().strip())
            self.road.setEnabled(False)
        else:
            self.road.setEditable(False)
            if self.road.currentIndex() < 0 and self.road.count():
                self.road.setCurrentIndex(0)
            self.road.setEnabled(account_type == "station" and self._global_manager)
        self.road.setToolTip(
            "测试账号直属管理员且不分配路段；"
            "路段管理员的账号显示名就是路段分类名称；"
            "中心站可选择已有路段或未分配。"
            if self._global_manager
            else "路段管理员只能管理自己所属路段。"
        )

    def _sync_road_manager_name(self, text):
        if self.role.currentData() == "road_admin":
            self.road.setEditText(str(text).strip())

    def _accept_if_valid(self):
        username = self.username.text().strip()
        if not username or not self.display_name.text().strip():
            QMessageBox.warning(self, "资料不完整", "登录名和站点显示名不能为空。")
            return
        invalid_username = any(
            character.isspace() or not character.isprintable()
            for character in username
        )
        if invalid_username:
            QMessageBox.warning(
                self,
                "登录名无效",
                "登录名不能包含空格、换行或其他控制字符。",
            )
            self.username.setFocus()
            return
        account_type = self.role.currentData()
        if account_type == "station":
            road_id, _road_name = self._road_reference()
            if self.scope.currentData() == "road" and not road_id:
                QMessageBox.warning(
                    self,
                    "所属路段无效",
                    "本路段数据范围必须先选择所属路段。",
                )
                return
        self.accept()

    def _road_reference(self):
        road_name = self.road.currentText().strip()
        index = self.road.currentIndex()
        custom_name = bool(
            self.road.isEditable()
            and (
                index < 0
                or road_name != self.road.itemText(index).strip()
            )
        )
        return (
            None if custom_name else self.road.currentData(),
            road_name,
        )

    def values(self) -> dict:
        account_type = self.role.currentData()
        resolved_type = "station" if account_type == "test" else account_type
        data_scope = self.scope.currentData()
        values = {
            "username": self.username.text().strip().lower(),
            "display_name": self.display_name.text().strip(),
            "role": "admin" if resolved_type == "admin" else "user",
            "account_type": resolved_type,
            "is_test": account_type == "test",
            "stats_scope": "all" if data_scope == "all" else "own",
            "data_scope": data_scope,
            "device_limit": self.device_limit.value(),
            "is_active": self.active.isChecked(),
        }
        if account_type == "station":
            road_id, _road_name = self._road_reference()
            if road_id:
                values["road_id"] = str(road_id)
            else:
                values["road_id"] = None
        return values


class _DevicesDialog(FramelessDialog):
    def __init__(self, page, account, parent=None):
        super().__init__(parent)
        self.page = page
        self.account = account
        self.setWindowTitle("设备管理")
        self.resize(760, 430)
        layout = QVBoxLayout(self)
        title = QLabel(f"{account.name_label} · 已登记设备")
        title.setObjectName("PageTitle")
        layout.addWidget(title)
        hint = QLabel(
            "撤销会立即终止该设备的在线会话；"
            "已离线设备最长仍可能使用到其 7 天授权过期。"
        )
        hint.setObjectName("Muted")
        hint.setWordWrap(True)
        layout.addWidget(hint)
        self.table = QTableWidget(0, 5)
        self.table.setHorizontalHeaderLabels(
            ["设备名", "客户端", "最近在线", "登记时间", "状态"]
        )
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        make_table_columns_resizable(
            self.table,
            [160, 130, 170, 170, 95],
        )
        layout.addWidget(self.table)
        buttons = QHBoxLayout()
        revoke = QPushButton("撤销所选设备")
        revoke.setObjectName("DangerButton")
        revoke.clicked.connect(self._revoke)
        close = QPushButton("关闭")
        close.clicked.connect(self.accept)
        buttons.addWidget(revoke)
        buttons.addStretch()
        buttons.addWidget(close)
        layout.addLayout(buttons)
        self.devices = []
        QTimer.singleShot(0, self.refresh)

    def refresh(self):
        result = self.page._call(
            self.page.session.api.admin_devices,
            self.account.server_account_id,
        )
        if result is _CALL_FAILED:
            return
        self.devices = result
        self.table.setRowCount(len(self.devices))
        for row, device in enumerate(self.devices):
            values = [
                device.get("name") or "-",
                device.get("client_version") or "-",
                str(device.get("last_seen_at") or "-").replace("T", " ")[:19],
                str(device.get("created_at") or "-").replace("T", " ")[:19],
                "已撤销" if device.get("revoked_at") else "有效",
            ]
            for column, value in enumerate(values):
                self.table.setItem(row, column, QTableWidgetItem(str(value)))

    def _revoke(self):
        row = self.table.currentRow()
        if row < 0 or row >= len(self.devices):
            return
        device = self.devices[row]
        if device.get("revoked_at"):
            return
        reply = QMessageBox.question(
            self,
            "确认撤销设备",
            f"确定撤销设备“{device.get('name') or device.get('device_uid')}”吗？",
            QMessageBox.Yes | QMessageBox.No,
        )
        if reply != QMessageBox.Yes:
            return
        if (
            self.page._call(
                self.page.session.api.admin_revoke_device,
                device["id"],
            )
            is _CALL_FAILED
        ):
            return
        self.refresh()


class _AuditDetailDialog(FramelessDialog):
    def __init__(self, row, action_label, parent=None):
        super().__init__(parent)
        self.setWindowTitle("审计记录详情")
        self.resize(620, 500)
        layout = QVBoxLayout(self)
        title = QLabel("审计记录详情")
        title.setObjectName("PageTitle")
        layout.addWidget(title)
        form = QFormLayout()
        form.addRow(
            "时间",
            QLabel(str(row.get("created_at") or "-").replace("T", " ")[:19]),
        )
        form.addRow("操作", QLabel(action_label))
        form.addRow(
            "目标",
            QLabel(
                f"{row.get('target_type') or '-'} · "
                f"{row.get('target_id') or '-'}"
            ),
        )
        form.addRow("操作账号", QLabel(str(row.get("actor_account_id") or "-")))
        form.addRow("来源 IP", QLabel(str(row.get("ip_address") or "-")))
        form.addRow("请求编号", QLabel(str(row.get("request_id") or "-")))
        layout.addLayout(form)
        details_label = QLabel("完整数据")
        details_label.setObjectName("SettingFieldLabel")
        layout.addWidget(details_label)
        self.details = QPlainTextEdit()
        self.details.setReadOnly(True)
        self.details.setPlainText(
            json.dumps(
                row.get("details") or {},
                ensure_ascii=False,
                indent=2,
                default=str,
            )
        )
        layout.addWidget(self.details, 1)
        buttons = QHBoxLayout()
        buttons.addStretch()
        close = QPushButton("关闭")
        close.clicked.connect(self.accept)
        buttons.addWidget(close)
        layout.addLayout(buttons)


class _AuditDialog(FramelessDialog):
    ACTION_LABELS: ClassVar = {
        "account.create": "创建账号",
        "account.update": "修改账号",
        "account.archive": "归档账号",
        "account.restore": "恢复账号",
        "account.password_reset": "重置密码",
        "account.data_reset": "重置统计",
        "account.delete": "永久删除归档账号",
        "device.revoke": "撤销设备",
        "message.delete": "删除用户消息",
        "captcha.policy.update": "修改验证码上传策略",
        "captcha.dataset.export": "导出验证码数据集",
        "captcha.dataset.import": "导入验证码数据集",
        "captcha.model.create": "创建验证码候选模型",
        "captcha.model.activate": "应用验证码模型",
        "captcha.model.use_builtin": "恢复内置验证码模型",
        "captcha.model.delete": "删除验证码模型",
    }

    def __init__(self, page, parent=None):
        super().__init__(parent)
        self.page = page
        self.setWindowTitle("管理员审计记录")
        self.resize(940, 520)
        layout = QVBoxLayout(self)
        title = QLabel("管理员审计记录")
        title.setObjectName("PageTitle")
        layout.addWidget(title)
        hint = QLabel("展示最近 200 条在线管理操作，记录由服务器保存。")
        hint.setObjectName("Muted")
        layout.addWidget(hint)
        self.table = QTableWidget(0, 6)
        self.table.setHorizontalHeaderLabels(
            ["时间", "操作", "目标", "来源 IP", "请求编号", "详情"]
        )
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        self.table.setShowGrid(False)
        self.table.verticalHeader().setVisible(False)
        self.table.cellDoubleClicked.connect(self._open_detail)
        make_table_columns_resizable(
            self.table,
            [170, 130, 160, 135, 210, 260],
        )
        layout.addWidget(self.table, 1)
        buttons = QHBoxLayout()
        refresh = QPushButton("刷新")
        refresh.clicked.connect(self.refresh)
        self.export_btn = QPushButton("导出记录")
        self.export_btn.setToolTip("导出当前已加载的审计记录")
        self.export_btn.clicked.connect(self._export)
        self.export_btn.setEnabled(False)
        close = QPushButton("关闭")
        close.clicked.connect(self.accept)
        buttons.addWidget(refresh)
        buttons.addWidget(self.export_btn)
        buttons.addStretch()
        buttons.addWidget(close)
        layout.addLayout(buttons)
        self.rows = []
        QTimer.singleShot(0, self.refresh)

    def refresh(self):
        payload = self.page._call(self.page.session.api.admin_audit, 200)
        if payload is _CALL_FAILED:
            return
        self.rows = list(payload.get("items") or [])
        self.export_btn.setEnabled(bool(self.rows))
        self.table.setRowCount(len(self.rows))
        for row_index, row in enumerate(self.rows):
            details = row.get("details") or {}
            if isinstance(details, dict):
                details_text = "，".join(
                    f"{key}={value}" for key, value in details.items()
                )
            else:
                details_text = str(details)
            values = [
                str(row.get("created_at") or "-").replace("T", " ")[:19],
                self.ACTION_LABELS.get(
                    str(row.get("action") or ""),
                    str(row.get("action") or "-"),
                ),
                str(row.get("target_id") or "-"),
                str(row.get("ip_address") or "-"),
                str(row.get("request_id") or "-"),
                details_text or "-",
            ]
            for column, value in enumerate(values):
                self.table.setItem(
                    row_index,
                    column,
                    QTableWidgetItem(value),
                )

    @staticmethod
    def _export_row_values(row):
        details = row.get("details") or {}
        if isinstance(details, dict):
            details_text = json.dumps(
                details,
                ensure_ascii=False,
                separators=(",", ":"),
                default=str,
            )
        else:
            details_text = str(details)
        return [
            str(row.get("created_at") or "-").replace("T", " ")[:19],
            _AuditDialog.ACTION_LABELS.get(
                str(row.get("action") or ""),
                str(row.get("action") or "-"),
            ),
            str(row.get("target_type") or "-"),
            str(row.get("target_id") or "-"),
            str(row.get("actor_account_id") or "-"),
            str(row.get("ip_address") or "-"),
            str(row.get("request_id") or "-"),
            details_text or "-",
        ]

    def _export(self):
        if not self.rows:
            return
        default_name = (
            f"审计记录_{datetime.now().astimezone():%Y%m%d_%H%M%S}.xlsx"
        )
        target, selected_filter = QFileDialog.getSaveFileName(
            self,
            "导出审计记录",
            default_name,
            "Excel 工作簿 (*.xlsx);;CSV 文件 (*.csv)",
        )
        if not target:
            return
        target_path = Path(target)
        is_csv = "CSV" in str(selected_filter) or target_path.suffix.lower() == ".csv"
        if is_csv and target_path.suffix.lower() != ".csv":
            target_path = target_path.with_suffix(".csv")
        elif not is_csv and target_path.suffix.lower() != ".xlsx":
            target_path = target_path.with_suffix(".xlsx")
        headers = [
            "时间",
            "操作",
            "目标类型",
            "目标编号",
            "操作账号",
            "来源 IP",
            "请求编号",
            "详情",
        ]
        values = [self._export_row_values(row) for row in self.rows]
        try:
            if is_csv:
                with target_path.open(
                    "w",
                    encoding="utf-8-sig",
                    newline="",
                ) as handle:
                    writer = csv.writer(handle)
                    writer.writerow(headers)
                    writer.writerows(values)
            else:
                from openpyxl import Workbook
                from openpyxl.styles import Alignment, Font, PatternFill

                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "审计记录"
                sheet.append(headers)
                for row_values in values:
                    sheet.append(row_values)
                header_fill = PatternFill("solid", fgColor="1D8178")
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
                widths = [21, 20, 14, 38, 38, 18, 38, 60]
                for index, width in enumerate(widths, start=1):
                    sheet.column_dimensions[chr(64 + index)].width = width
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
                for row in sheet.iter_rows(min_row=2):
                    row[-1].alignment = Alignment(wrap_text=True, vertical="top")
                workbook.save(target_path)
        except (ImportError, OSError, ValueError) as exc:
            QMessageBox.warning(self, "导出失败", str(exc))
            return
        QMessageBox.information(
            self,
            "导出完成",
            f"已导出 {len(values)} 条审计记录：\n{target_path}",
        )

    def _open_detail(self, row_index, _column):
        if 0 <= row_index < len(self.rows):
            row = self.rows[row_index]
            action = str(row.get("action") or "-")
            _AuditDetailDialog(
                row,
                self.ACTION_LABELS.get(action, action),
                self,
            ).exec_()


class OnlineAccountPage(AccountPage):
    """Offline account management UI extended with server-only controls."""

    def __init__(self, database, current_account, session, parent=None):
        self.session = session
        self.server_rows = {}
        self.all_accounts = []
        self.roads = []
        super().__init__(
            database,
            current_account,
            parent,
            defer_refresh=True,
        )

        self.page_subtitle.setText(
            "按管理员层级管理账号、所属路段、数据范围、设备、归档和审计。"
        )
        self.list_title.setText("在线账号列表")

        road_admin_card, road_admin_value = self._summary_card(
            "路段管理员",
            "#6845bd",
        )
        self.summary_values["road_admin"] = road_admin_value
        self.summary_layout.addWidget(road_admin_card, 1)

        online_card, online_value = self._summary_card(
            "当前在线设备",
            "#3478f6",
        )
        self.summary_values["online"] = online_value
        self.summary_layout.addWidget(online_card, 1)

        self.create_btn.setText("＋ 新建在线账号")
        self.rename_btn.setText("修改名称")
        self.permission_btn.setText("权限与设备数")
        self.reset_btn.setText("重置密码")
        self.toggle_btn.setText("停用/启用")
        self.delete_btn.setText("归档账号")
        self.delete_btn.setObjectName("DangerButton")

        self.devices_btn = QPushButton("设备列表")
        self.devices_btn.clicked.connect(self._devices)
        toggle_index = self.account_action_layout.indexOf(self.toggle_btn)
        self.account_action_layout.insertWidget(toggle_index, self.devices_btn)

        self.purge_btn = QPushButton("删除归档数据")
        self.purge_btn.setObjectName("DangerButton")
        self.purge_btn.clicked.connect(self._purge_archived_account)
        delete_index = self.account_action_layout.indexOf(self.delete_btn)
        self.account_action_layout.insertWidget(delete_index + 1, self.purge_btn)

        self.audit_btn = QPushButton("审计记录")
        self.audit_btn.clicked.connect(self._show_audit)
        self.account_action_layout.insertWidget(1, self.audit_btn)

        filter_label = QLabel("账号状态")
        filter_label.setObjectName("Muted")
        self.account_filter = QComboBox()
        self.account_filter.setObjectName("AccountArchiveFilter")
        self.account_filter.addItem("在用账号", "active")
        self.account_filter.addItem("已归档账号", "archived")
        self.account_filter.addItem("全部账号", "all")
        self.account_filter.setToolTip("切换在用账号和已归档账号")

        role_filter_label = QLabel("角色")
        role_filter_label.setObjectName("Muted")
        self.role_filter = QComboBox()
        self.role_filter.setObjectName("AccountRoleFilter")
        self.role_filter.addItem("全部角色", "all")
        self.role_filter.addItem("管理员", "admin")
        self.role_filter.addItem("路段管理员", "road_admin")
        self.role_filter.addItem("中心站账号", "station")
        self.role_filter.addItem("测试账号", "test")
        self.role_filter.setToolTip("按账号角色筛选")

        road_filter_label = QLabel("所属路段")
        road_filter_label.setObjectName("Muted")
        self.road_filter = QComboBox()
        self.road_filter.setObjectName("AccountRoadFilter")
        self.road_filter.setMinimumWidth(130)
        self.road_filter.setToolTip("按所属路段筛选；未分配包含测试账号")
        self._refresh_road_filter_options()

        self.list_header_layout.insertWidget(1, filter_label)
        self.list_header_layout.insertWidget(2, self.account_filter)
        self.list_header_layout.insertWidget(3, role_filter_label)
        self.list_header_layout.insertWidget(4, self.role_filter)
        self.list_header_layout.insertWidget(5, road_filter_label)
        self.list_header_layout.insertWidget(6, self.road_filter)
        self.account_filter.currentIndexChanged.connect(
            self._apply_account_filters
        )
        self.role_filter.currentIndexChanged.connect(
            self._apply_account_filters
        )
        self.road_filter.currentIndexChanged.connect(
            self._apply_account_filters
        )

        self.reset_stats_btn.setText("重置在线统计")
        self.reset_stats_btn.setToolTip(
            "清除所选账号在服务器上的全部汇总数据，并写入审计记录"
        )

        # Compatibility aliases retained for existing automation and callers.
        self.edit_btn = self.permission_btn
        self.password_btn = self.reset_btn
        self.archive_btn = self.delete_btn

        self.table.setColumnCount(13)
        self.table.setHorizontalHeaderLabels(
            [
                "用户名称",
                "登录名",
                "角色",
                "数据范围",
                "状态",
                "设备/上限",
                "在线",
                "创建时间",
                "最后登录",
                "最近系统",
                "首次改密",
                "归档",
                "所属路段",
            ]
        )
        make_table_columns_resizable(
            self.table,
            [160, 130, 105, 110, 90, 110, 85, 170, 170, 105, 105, 90, 120],
        )
        # Do not perform a blocking HTTPS request while MainWindow is still
        # being constructed. MainWindow calls refresh() when the user opens
        # the account page, after the initial event loop has started.
        self._selection_changed()

    def _refresh_summary(self):
        accounts = self.all_accounts or self.accounts
        self.summary_values["total"].setText(str(len(accounts)))
        self.summary_values["admin"].setText(
            str(sum(account.is_admin for account in accounts))
        )
        self.summary_values["user"].setText(
            str(
                sum(
                    account.is_station and not account.is_test
                    for account in accounts
                )
            )
        )
        self.summary_values["active"].setText(
            str(sum(account.is_active for account in accounts))
        )
        self.summary_values["road_admin"].setText(
            str(sum(account.is_road_admin for account in accounts))
        )

    def _call(self, function, *args):
        function_name = getattr(function, "__name__", "")
        message = (
            "加载中…"
            if function_name
            in {
                "admin_accounts",
                "admin_roads",
                "admin_devices",
                "admin_audit",
            }
            else "处理中…"
        )
        try:
            return run_with_loading(
                self,
                message,
                lambda: function(self.session.access_token(), *args),
            )
        except (ApiResponseError, NetworkUnavailable) as exc:
            QMessageBox.warning(self, "在线操作失败", self._api_error_text(exc))
            return _CALL_FAILED

    @staticmethod
    def _api_error_text(error):
        if not isinstance(error, ApiResponseError) or not isinstance(
            error.details, list
        ):
            return str(error)
        field_labels = {
            "username": "登录名",
            "display_name": "站点显示名",
            "role": "角色",
            "account_type": "账号类型",
            "is_test": "账号类型",
            "stats_scope": "数据范围",
            "data_scope": "数据范围",
            "road_id": "所属路段",
            "road_name": "所属路段",
            "device_limit": "登录设备上限",
            "is_active": "账号状态",
        }
        messages = []
        for detail in error.details[:4]:
            if not isinstance(detail, dict):
                continue
            location = list(detail.get("loc") or [])
            field = str(location[-1]) if location else "请求参数"
            label = field_labels.get(field, field)
            message = str(detail.get("msg") or "参数无效").removeprefix(
                "Value error, "
            )
            messages.append(f"{label}：{message}")
        return "\n".join(messages) if messages else str(error)

    @staticmethod
    def _legacy_server_rejects_test_marker(error):
        if (
            not isinstance(error, ApiResponseError)
            or error.status_code != 422
            or not isinstance(error.details, list)
        ):
            return False
        return any(
            isinstance(detail, dict)
            and list(detail.get("loc") or [])[-1:] == ["is_test"]
            and (
                str(detail.get("type") or "") == "extra_forbidden"
                or "extra inputs" in str(detail.get("msg") or "").lower()
            )
            for detail in error.details
        )

    def _update_account_with_compatibility(self, account_id, payload):
        def update(request_payload):
            return self.session.api.admin_update_account(
                self.session.access_token(),
                account_id,
                request_payload,
            )

        try:
            return run_with_loading(
                self,
                "处理中…",
                lambda: update(payload),
            )
        except ApiResponseError as exc:
            if not (
                "is_test" in payload
                and self._legacy_server_rejects_test_marker(exc)
            ):
                QMessageBox.warning(
                    self,
                    "在线操作失败",
                    self._api_error_text(exc),
                )
                return _CALL_FAILED
            fallback_payload = dict(payload)
            requested_test_account = bool(fallback_payload.pop("is_test"))
            if requested_test_account or not fallback_payload:
                QMessageBox.warning(
                    self,
                    "在线服务版本过旧",
                    "当前在线服务不支持修改账号类型，请先将服务端升级到 v1.1.1。",
                )
                return _CALL_FAILED
            try:
                return run_with_loading(
                    self,
                    "正在兼容旧版服务…",
                    lambda: update(fallback_payload),
                )
            except (ApiResponseError, NetworkUnavailable) as retry_error:
                QMessageBox.warning(
                    self,
                    "在线操作失败",
                    self._api_error_text(retry_error),
                )
                return _CALL_FAILED
        except NetworkUnavailable as exc:
            QMessageBox.warning(self, "在线操作失败", self._api_error_text(exc))
            return _CALL_FAILED

    @staticmethod
    def _scope_label(scope):
        return {"own": "本人", "road": "本路段", "all": "全部"}.get(
            str(scope),
            "本人",
        )

    @staticmethod
    def _account_role_filter_key(account):
        if account.is_test:
            return "test"
        return account.resolved_account_type

    def _refresh_road_filter_options(self):
        current_value = (
            self.road_filter.currentData()
            if self.road_filter.count()
            else "all"
        )
        road_options = {}
        for road in self.roads:
            road_id = str(road.get("id") or "").strip()
            road_name = str(road.get("name") or "").strip()
            if road_id and road_name:
                road_options[road_id] = road_name
        # Older compatible services may not expose the roads endpoint. The
        # account rows still carry enough information to build the filter.
        for account in self.all_accounts:
            road_id = str(account.road_id or "").strip()
            road_name = str(account.road_name or "").strip()
            if road_id and road_name:
                road_options.setdefault(road_id, road_name)

        self.road_filter.blockSignals(True)
        self.road_filter.clear()
        self.road_filter.addItem("全部路段", "all")
        self.road_filter.addItem("未分配", "unassigned")
        for road_id, road_name in sorted(
            road_options.items(),
            key=lambda item: item[1].casefold(),
        ):
            self.road_filter.addItem(road_name, road_id)
        current_index = self.road_filter.findData(current_value)
        self.road_filter.setCurrentIndex(max(current_index, 0))
        self.road_filter.blockSignals(False)

    def _matches_account_filters(self, account):
        archive_mode = self.account_filter.currentData()
        if archive_mode == "active" and account.is_archived:
            return False
        if archive_mode == "archived" and not account.is_archived:
            return False

        role_mode = self.role_filter.currentData()
        if (
            role_mode != "all"
            and self._account_role_filter_key(account) != role_mode
        ):
            return False

        road_mode = self.road_filter.currentData()
        if road_mode == "all":
            return True
        # Global administrators have no road membership; "unassigned" is
        # reserved for directly managed station/test accounts.
        if account.is_admin:
            return False
        if road_mode == "unassigned":
            return account.is_test or not account.road_id
        return not account.is_test and str(account.road_id or "") == str(
            road_mode
        )

    def _apply_account_filters(
        self,
        _index=None,
        *,
        selected_server_id=_FILTER_SELECTION_UNSET,
    ):
        if selected_server_id is _FILTER_SELECTION_UNSET:
            selected = self._selected_account()
            selected_server_id = (
                selected.server_account_id if selected else None
            )
        self.accounts = [
            account
            for account in self.all_accounts
            if self._matches_account_filters(account)
        ]
        self.table.blockSignals(True)
        self.table.clearSelection()
        self.table.setRowCount(len(self.accounts))
        for row, account in enumerate(self.accounts):
            server = self.server_rows[account.server_account_id]
            values = [
                account.name_label,
                account.username,
                account.role_label,
                self._scope_label(account.effective_data_scope),
                "启用" if account.is_active else "停用",
                (
                    f"{server.get('active_device_count', 0)}/"
                    f"{server.get('device_limit', 10000)}"
                ),
                str(server.get("online_device_count", 0)),
                str(server.get("created_at") or account.created_at)
                .replace("T", " ")[:19],
                (account.last_login or "-").replace("T", " ")[:19],
                str(server.get("last_login_system") or "-"),
                "需要修改" if account.must_change_password else "已设置",
                "已归档" if account.is_archived else "-",
                (
                    "不适用"
                    if account.is_admin
                    else account.road_name or "未分配"
                ),
            ]
            for column, value in enumerate(values):
                item = QTableWidgetItem(str(value))
                if column == 0:
                    item.setData(Qt.UserRole, account.id)
                if column == 2:
                    item.setForeground(
                        QColor("#6845bd")
                        if account.is_account_manager
                        else QColor("#26718b")
                    )
                if column == 4:
                    item.setForeground(
                        QColor("#188b57")
                        if account.is_active and not account.is_archived
                        else QColor("#d33f49")
                    )
                if column == 6 and int(server.get("online_device_count") or 0):
                    item.setForeground(QColor("#188b57"))
                if column == 10 and account.must_change_password:
                    item.setForeground(QColor("#b26a00"))
                if column == 11 and account.is_archived:
                    item.setForeground(QColor("#d33f49"))
                item.setTextAlignment(
                    Qt.AlignLeft | Qt.AlignVCenter
                    if column in (0, 1, 12)
                    else Qt.AlignCenter
                )
                self.table.setItem(row, column, item)
            if account.server_account_id == selected_server_id:
                self.table.selectRow(row)
        self.table.blockSignals(False)
        self._selection_changed()

    def refresh(self):
        roads_loader = getattr(self.session.api, "admin_roads", None)
        if callable(roads_loader):
            roads = self._call(roads_loader)
            if roads is not _CALL_FAILED and roads is not None:
                self.roads = list(roads)
        rows = self._call(self.session.api.admin_accounts)
        if rows is _CALL_FAILED:
            return
        if rows is None:
            return
        rows = list(rows)
        selected = self._selected_account()
        selected_server_id = selected.server_account_id if selected else None
        server_account_ids = {
            str(row.get("id") or row.get("account_id") or "").strip()
            for row in rows
            if str(row.get("id") or row.get("account_id") or "").strip()
        }
        cached_remote_accounts = {
            account.server_account_id: account
            for account in self.database.list_accounts()
            if account.server_account_id
        }
        self.accounts = []
        self.server_rows = {}
        reconciled_accounts = self.database.reconcile_remote_accounts(
            rows,
            prune_missing=self.current_account.is_admin,
        )
        removed_account_ids = (
            [
                account.id
                for server_account_id, account in cached_remote_accounts.items()
                if server_account_id not in server_account_ids
                and account.id != self.current_account.id
            ]
            if self.current_account.is_admin
            else []
        )
        for server_row, account in zip(rows, reconciled_accounts):
            object.__setattr__(
                account,
                "_device_limit",
                int(server_row.get("device_limit") or 1),
            )
            self.accounts.append(account)
            self.server_rows[account.server_account_id] = server_row
        self.all_accounts = list(self.accounts)
        self._refresh_road_filter_options()
        self._refresh_summary()
        self.summary_values["online"].setText(
            str(
                sum(
                    int(row.get("online_device_count") or 0)
                    for row in rows
                )
            )
        )
        self._apply_account_filters(
            selected_server_id=selected_server_id,
        )
        for account_id in removed_account_ids:
            self.account_deleted.emit(account_id)

    def _selection_changed(self):
        account = self._selected_account()
        selected = account is not None
        other = selected and account.id != self.current_account.id
        is_station = selected and account.resolved_account_type == "station"
        has_statistics = is_station and not account.is_test
        self.rename_btn.setEnabled(selected)
        self.permission_btn.setEnabled(selected)
        self.password_btn.setEnabled(other)
        self.devices_btn.setEnabled(selected)
        self.toggle_btn.setEnabled(other and not account.is_archived if account else False)
        self.archive_btn.setEnabled(other)
        self.purge_btn.setEnabled(bool(other and account and account.is_archived))
        self.export_btn.setEnabled(has_statistics)
        self.import_btn.setEnabled(has_statistics)
        self.reset_stats_btn.setEnabled(other and has_statistics)
        self.data_range_selector.setEnabled(has_statistics)
        self.archive_btn.setText(
            "恢复账号" if account and account.is_archived else "归档账号"
        )
        self.toggle_btn.setText(
            "停用账号"
            if account and account.is_active
            else "启用账号"
        )
        if account:
            server = self.server_rows.get(account.server_account_id, {})
            state = (
                "已归档"
                if account.is_archived
                else ("正常" if account.is_active else "已停用")
            )
            suffix = (
                " · 当前登录账号"
                if account.id == self.current_account.id
                else ""
            )
            hierarchy = (
                "全部路段"
                if account.is_admin
                else (
                    f"路段 {account.road_name}"
                    if account.road_name
                    else "未分配 · 直属管理员"
                )
            )
            self.selection_hint.setText(
                f"已选择：{account.name_label} · {account.role_label} · {state}"
                f" · {hierarchy}"
                f" · 数据范围 {self._scope_label(account.effective_data_scope)}"
                f" · 设备 {server.get('active_device_count', 0)}/"
                f"{server.get('device_limit', 10000)}{suffix}"
            )
        else:
            self.selection_hint.setText("请选择一个账号进行管理")
            self.toggle_btn.setText("停用/启用")
        self.import_btn.setToolTip(
            "在线版不允许只导入本机缓存；点击可查看说明"
            if has_statistics
            else "测试账号不参与统计" if is_station else "请选择中心站账号"
        )
        self.export_btn.setToolTip(
            "导出当前客户端已同步到的站点汇总数据"
            if has_statistics
            else "测试账号不参与统计" if is_station else "请选择中心站账号"
        )
        self.toggle_btn.setObjectName(
            "DangerButton"
            if other and account and account.is_active
            else ("PrimaryButton" if other and account else "")
        )
        self.toggle_btn.style().unpolish(self.toggle_btn)
        self.toggle_btn.style().polish(self.toggle_btn)

    def _create_account(self):
        dialog = _AccountSettingsDialog(
            parent=self,
            current_account=self.current_account,
            roads=self.roads,
        )
        if dialog.exec_() != dialog.Accepted:
            return
        payload = dialog.values()
        # Keep the legacy marker omission for ordinary accounts.  The v1.2
        # hierarchy fields are still sent to enforce mandatory road assignment.
        if not payload.get("is_test"):
            payload.pop("is_test", None)
        result = self._call(
            self.session.api.admin_create_account,
            payload,
        )
        if result is not _CALL_FAILED:
            QMessageBox.information(
                self,
                "创建成功",
                "账号已创建并启用，初始密码为 123456；首次登录需要修改密码。",
            )
            self.refresh()

    def _rename_account(self):
        account = self._selected_account()
        if not account:
            return
        dialog = RenameAccountDialog(account, self)
        if dialog.exec_() != dialog.Accepted:
            return
        display_name = dialog.value()
        result = self._call(
            self.session.api.admin_update_account,
            account.server_account_id,
            {"display_name": display_name},
        )
        if result is not _CALL_FAILED:
            self.account_name_changed.emit(account.id, display_name)
            self.refresh()

    def _change_permission(self):
        account = self._selected_account()
        if not account:
            return
        dialog = _AccountSettingsDialog(
            account,
            self,
            current_account=self.current_account,
            roads=self.roads,
        )
        if dialog.exec_() != dialog.Accepted:
            return
        requested = dialog.values()
        requested.pop("username", None)
        current = {
            "display_name": account.name_label,
            "role": account.role,
            "account_type": account.resolved_account_type,
            "is_test": bool(account.is_test),
            "stats_scope": account.stats_scope,
            "data_scope": account.effective_data_scope,
            "device_limit": int(
                getattr(account, "_device_limit", 10000) or 10000
            ),
            "is_active": bool(account.is_active),
        }
        if account.resolved_account_type != "admin":
            current["road_id"] = str(account.road_id or "")
        payload = {
            field: value
            for field, value in requested.items()
            if value != current.get(field)
        }
        if not payload:
            return
        result = self._update_account_with_compatibility(
            account.server_account_id,
            payload,
        )
        if result is not _CALL_FAILED:
            if "display_name" in payload:
                self.account_name_changed.emit(
                    account.id,
                    payload["display_name"],
                )
            if {
                "role",
                "account_type",
                "is_test",
                "stats_scope",
                "data_scope",
                "road_id",
                "road_name",
                "is_active",
            } & payload.keys():
                self.account_permission_changed.emit(
                    account.id,
                    str(payload.get("role", account.role)),
                )
            self.refresh()

    def _reset_password(self):
        account = self._selected_account()
        if not account:
            return
        reply = QMessageBox.question(
            self,
            "确认重置密码",
            "密码将恢复为 123456，全部在线会话会被撤销，并要求下次登录改密。",
            QMessageBox.Yes | QMessageBox.No,
        )
        if reply == QMessageBox.Yes:
            result = self._call(
                self.session.api.admin_reset_password,
                account.server_account_id,
            )
            if result is not _CALL_FAILED:
                QMessageBox.information(
                    self,
                    "重置完成",
                    "密码已恢复为 123456，用户下次登录必须修改密码。",
                )
                self.refresh()

    def _toggle_active(self):
        account = self._selected_account()
        if not account or account.id == self.current_account.id:
            return
        enabled = not account.is_active
        action = "启用" if enabled else "停用"
        reply = QMessageBox.question(
            self,
            f"确认{action}",
            f"确定要{action}账号“{account.username}”吗？",
            QMessageBox.Yes | QMessageBox.No,
        )
        if reply != QMessageBox.Yes:
            return
        result = self._call(
            self.session.api.admin_update_account,
            account.server_account_id,
            {"is_active": enabled},
        )
        if result is not _CALL_FAILED:
            self.refresh()

    def _devices(self):
        account = self._selected_account()
        if account:
            _DevicesDialog(self, account, self).exec_()
            self.refresh()

    def _delete_account(self):
        account = self._selected_account()
        if not account:
            return
        action = "恢复" if account.is_archived else "归档"
        reply = QMessageBox.question(
            self,
            f"确认{action}",
            (
                "恢复后账号仍保持停用，需要在“编辑设置”中单独启用。"
                if account.is_archived
                else "归档会禁止登录并终止在线会话，但历史统计会保留。"
            ),
            QMessageBox.Yes | QMessageBox.No,
        )
        if reply != QMessageBox.Yes:
            return
        method = (
            self.session.api.admin_restore_account
            if account.is_archived
            else self.session.api.admin_archive_account
        )
        if self._call(method, account.server_account_id) is not _CALL_FAILED:
            self.account_deleted.emit(account.id)
            self.refresh()

    def _purge_archived_account(self):
        account = self._selected_account()
        if (
            not account
            or not account.is_archived
            or account.id == self.current_account.id
        ):
            return
        reply = QMessageBox.question(
            self,
            "永久删除归档账号",
            "将永久删除该账号、业务统计、计时记录、设备和会话数据，"
            "且无法恢复。"
            + (
                "该路段下的中心站将改为“未分配”，并由管理员直接管理。"
                if account.is_road_admin
                else ""
            )
            + "\n\n"
            f"确定永久删除“{account.name_label}（{account.username}）”吗？",
            QMessageBox.Yes | QMessageBox.No,
        )
        if reply != QMessageBox.Yes:
            return
        result = self._call(
            self.session.api.admin_delete_archived_account,
            account.server_account_id,
        )
        if result is _CALL_FAILED:
            return
        self.database.purge_remote_account_cache(account.server_account_id)
        self.account_deleted.emit(account.id)
        QMessageBox.information(
            self,
            "删除完成",
            (
                "路段管理员已永久删除；所属中心站已改为未分配并由管理员直接管理。"
                if account.is_road_admin
                else "归档账号及其数据已永久删除。"
            ),
        )
        self.refresh()

    def _reset_station_statistics(self):
        account = self._selected_account()
        if not account:
            return
        reply = QMessageBox.question(
            self,
            "确认重置统计数据",
            f"确定清除“{account.name_label}”的所有在线汇总数据吗？\n"
            "此操作独立于归档，并会写入审计记录；本机原始 Excel 不受影响。",
            QMessageBox.Yes | QMessageBox.No,
        )
        if reply != QMessageBox.Yes:
            return
        result = self._call(
            self.session.api.admin_reset_data,
            account.server_account_id,
        )
        if result is not _CALL_FAILED:
            self.station_data_changed.emit(account.id)
            QMessageBox.information(
                self,
                "重置完成",
                "服务器汇总数据已清除；账号、设备及本机原始文件不受影响。",
            )
            self.refresh()

    def _import_station_data(self):
        account = self._selected_account()
        if not account or account.is_admin:
            return
        QMessageBox.information(
            self,
            "在线版导入说明",
            "为避免形成“只在管理员电脑可见”的本机数据，在线测试版暂不执行"
            "管理员代导入。\n\n"
            "当前仍可导出已同步汇总数据；各站点产生的新数据会通过同步队列"
            "自动上传服务器。",
        )

    def _show_audit(self):
        _AuditDialog(self, self).exec_()
